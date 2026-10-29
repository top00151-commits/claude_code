# -*- coding: utf-8 -*-
"""z1191l — 베트남 인보이스 HS 자동 채움 (대표 지시 2026-10-26)
   「엑셀 양식은 지금 올려준 양식으로 최종 적용해 … 직원들이 가장 편하고 실수없이 빠르게 할수있는 방법으로」

흐름 (엑셀 왕복 — 법인은 지금처럼 엑셀만 다룬다)
  ① 본사(안지연 프로)가 인보이스(최종 양식: 수책·소모품)를 올린다
     → HS 사전의 **확정(CONFIRMED)** 값으로 HS CODE·Tên hàng 칸만 채운 같은 파일을 돌려준다.
     → 법인이 채울 칸은 노랑, 확인만 할 칸은 주황으로 칠한다.
  ② 법인은 노란 칸만 채워 이음 방으로 회신한다.
  ③ 본사가 회신 파일을 올리면 사전이 자란다(기존 z967 규칙: 새 규격 추가 · 같은 HS 확인 · 다른 HS 는 '확인 대기').

⚠ 파일은 통째로 다시 저장하지 않는다 — 엑셀 도구로 열었다 저장하면 사진·도형·인쇄 설정을 잃을 수 있다.
   엑셀 안쪽 XML 에서 **그 칸들만** 바꾸고 나머지(사진 52장·병합·수식·인쇄 영역)는 바이트 그대로 둔다.
⚠ 두 최종 양식 모두 인쇄 영역이 A~I 열이고 HS·Tên hàng 은 그 밖(R~T 열) → 색이 세관용 인보이스 출력에 안 찍힌다.
⚠ 통관은 법적 신고 — 확정 값만 자동으로 채운다. 확인 대기·AI 추정·후보가 갈리는 것은 채우지 않고 노랑으로 둔다.

z1191n (대표 2026-10-27 「전부 자동이면 법인 회신이 필요없는 사항인데」·「사전 확정을 하고 싶을때 어떻게해」):
  · 셈은 칸 색으로 한다(counts) — 노란 칸이 하나라도 있으면 '법인 입력'(품명만 노란 줄 포함) · 주황이면 '확인'.
    노랑·주황이 하나도 없으면 no_reply — 법인 회신이 필요 없고 받은 파일이 최종 파일이다.
  · 본사가 ① 결과에서 HS 를 고르면(choices) 사전에 확정하고, overrides({규격 키: HS})로 원본에서 파일을 다시 만든다.
    적힌 값은 원래 덮지 않는다 — 본사가 고른 규격만 덮는다.

z1191o (대표 2026-10-27 「여기서 표준 약식 이름을 "수책통관","일반통관" 으로 변경하자..」):
  · 양식 이름 수책 → 수책통관 · 소모품 → 일반통관. 알아보는 값(form)은 「수책」·「소모품」 그대로 둔다
    (기록 표에 이미 저장된 줄도 같은 값) — 화면·결과에는 form_name() 으로 새 이름을 쓴다.

z1191p (대표 2026-10-27 「올해 수출한 내용을 수책통관, 일반통관 정리한 파일… 예전 파일을 보고 HS코드 및 금액도
        비교할수있게 미리 등록을 해 줬으면해」 → 「진행해」):
  · scan 이 줄마다 수량·단가·금액(qty·price·amount)과 인보이스 날짜(date)도 읽는다 — 31줄 머리글 F·H·I(모든 KNK 양식 같음).
  · hist_lines = 지난 수출 이력 줄 · compare = ① 결과 줄 ↔ 이력(같은 규격만 · 단가 ±HIST_PRICE_TOL% · HS 다름) — 화면에만.

z1191r (대표 2026-10-27 「등록할때 이미지도 가능하면 같이 사진,단가 부분도 볼수있게 등록을해줘」):
  · row_images = 품목 줄에 붙은 그림(그림 자리 from~to 줄이 품목 줄에 걸리면 그 줄 · 첫 그림만) — 54개 실물 그림 510 중 450이 품목 줄.
"""
import io
import re
import zipfile
from xml.sax.saxutils import escape as _xml_escape

# ── 머리글 찾기 — main.py `_hs_parse_invoice_xlsx` 와 같은 낱말(사전 성장과 같은 칸을 본다) ──
HEAD_HS = re.compile(r"HS\s*CODE|HS\s*code|^HS$", re.I)
HEAD_HS_EXACT = re.compile(r"^\s*HS(\s*CODE)?\s*$", re.I)      # 'HS CODE CHECK' 같은 보조 칸보다 이것을 먼저
HEAD_VN = re.compile(r"T[êe]n\s*h[àa]ng", re.I)
HEAD_DESC = re.compile(r"Description", re.I)
WS_SKIP = re.compile(r"PACKING|P/L|CONTENT|COVER|양식|SAMPLE", re.I)
TOTAL_ROW = re.compile(r"^\s*(SUB\s*)?TOTAL\b|^\s*합\s*계", re.I)
INVOICE_NO = re.compile(r"KNK-\d{6}-\d+")
# z1191p: 수량·단가·금액 머리글('14) Q'ty /Unit' · '15) Unit Price' · '16) Amount') · 인보이스 날짜 칸 이름('DATE:')
HEAD_QTY = re.compile(r"Q'?\s*ty\b|Quantity", re.I)
HEAD_PRICE = re.compile(r"Unit\s*Price", re.I)
HEAD_AMT = re.compile(r"\bAmount\b", re.I)
HEAD_DATE = re.compile(r"^\s*DATE\s*:?\s*$", re.I)

# 상태 — 화면·기록에 같은 글자를 쓴다
AUTO = "AUTO"          # 확정 값으로 채움(색 없음)
CHECK = "CHECK"        # 채웠지만 확인 필요(주황) — 규격 전체가 아니라 모델 글자만 같았다
CONFLICT = "CONFLICT"  # 이미 적힌 HS 가 사전 확정 값과 다름(주황 · 값은 그대로)
NEED = "NEED"          # 법인이 채울 칸(노랑)
KEEP = "KEEP"          # 이미 적혀 있고 문제없음(그대로)

YELLOW = "FFFFFF00"
ORANGE = "FFFFC000"

# z1191o: 화면에 보이는 양식 이름 — 알아보는 값은 그대로(기록 표에 저장된 값도 이것)
FORM_NAMES = {"수책": "수책통관", "소모품": "일반통관"}


def form_name(form) -> str:
    """양식 값(여러 시트면 「·」로 이은 글자) → 화면 이름. 「예전 양식」 같은 다른 값은 그대로."""
    return "·".join(FORM_NAMES.get(f, f) for f in str(form or "").split("·") if f)


def norm_key(s) -> str:
    """사전 키 = 규격 정규화(공백 제거 + 대문자) — main.py `_hs_norm_key` 와 같다."""
    return re.sub(r"\s+", "", str(s or "")).upper()


def tokens(s) -> set:
    """모델 글자 조각 — 5글자 이상이고 숫자가 든 것만(소모품 양식은 'PARTS OF CONNECTORS SP100F…' 처럼
    앞에 분류 글이 붙어 규격 전체로는 사전과 안 맞는다)."""
    out = set()
    for x in re.split(r"[\s()/,]+", str(s or "")):
        x = x.strip(":;,.")
        if len(x) >= 5 and re.search(r"\d", x):
            out.add(norm_key(x))
    return out


def digits(s) -> str:
    return re.sub(r"\D", "", re.sub(r"\.0$", "", str(s if s is not None else "")))


# z1191v (대표 2026-10-28 「0 , #N/A 이건 엑셀 데이터가 함수로 연결되어서 오는건데 잘못 된거야... 이렇게 나오는 부분들은
#   HS코드 등록에서 모두 삭제해.」 → 사전 44줄 삭제 → 「1번 재발 막기 진행해」):
#   엑셀 함수가 못 찾을 때 내는 값 — 「0」(법인 회신 260901 의 =IFERROR(INDEX([다른 파일]…),"0") · 빈 칸을 가리킨 INDEX) ·
#   #N/A 같은 오류 표시(260420 의 VLOOKUP). 품명(Tên hàng) 칸이 이 값이면 데이터가 아니다 →
#   ③ 회신 · 📚 등록 · HS 사전 올림은 그 줄을 사전에 넣지 않고 · ① 은 빈칸으로 보고 · 📌 는 그 이력 줄을 쓰지 않는다.
#   (이력 줄은 실제 인보이스 기록이라 그대로 남긴다)
_FORMULA_JUNK = re.compile(
    r"^\s*(?:[+-]?0+(?:\.0+)?|#(?:N/A|REF!|VALUE!|DIV/0!|NAME\?|NULL!|NUM!|SPILL!|CALC!|GETTING_DATA|FIELD!|BLOCKED!|CONNECT!|BUSY!|UNKNOWN!))\s*$",
    re.I)


def formula_junk(v) -> bool:
    """엑셀 함수 실패 값인가 — 「0」 · 숫자 0 · 「#N/A」 같은 오류 표시. 빈칸은 아니다(False)."""
    if v is None or isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return v == 0
    return bool(_FORMULA_JUNK.match(str(v)))


def formula_junk_sql(v) -> int:
    """SQLite 함수로 등록해 쓰는 판(📌 가 이력 줄을 고르는 SQL) — 같은 판정 · 참이면 1."""
    return 1 if formula_junk(v) else 0


# ─────────────────────────────────────────────────────────────
# 사전 색인
# ─────────────────────────────────────────────────────────────
def build_index(entries) -> dict:
    """entries: export_hs_dict 줄(dict) 목록 → {exact: {키: 줄}, tok: {조각: [줄…]}}"""
    exact, tok = {}, {}
    for e in entries:
        k = e.get("part_key") or ""
        if not k:
            continue
        exact[k] = e
        ts = tokens(e.get("model")) | ({k} if (len(k) >= 5 and re.search(r"\d", k)) else set())
        for t in ts:
            tok.setdefault(t, []).append(e)
    return {"exact": exact, "tok": tok}


def _confirmed(e) -> bool:
    return (e.get("status") or "") == "CONFIRMED" and len(digits(e.get("hs_code"))) == 8


def lookup(index, desc):
    """규격 하나 → (방법, 줄, 메모).
    방법: exact(규격 그대로 · 확정) / token(모델 글자만 · 확정 후보가 모두 같은 HS) /
          pending(사전에 있으나 확정 아님) / ambiguous(후보 HS 가 갈림) / none(사전에 없음)"""
    k = norm_key(desc)
    e = index["exact"].get(k)
    if e is not None:
        if _confirmed(e):
            return "exact", e, ""
        return "pending", e, "사전 %s(%s)" % ({"PENDING": "확인 대기", "AI": "AI 추정"}.get(e.get("status"), e.get("status") or "?"),
                                              digits(e.get("hs_code")) or "HS 없음")
    cands = {}
    for t in tokens(desc):
        for c in index["tok"].get(t, []):
            cands[c.get("part_key")] = c
    if not cands:
        return "none", None, ""
    conf = [c for c in cands.values() if _confirmed(c)]
    if not conf:
        c0 = next(iter(cands.values()))
        return "pending", c0, "같은 모델이 사전에 확정 아님(%s)" % (digits(c0.get("hs_code")) or "HS 없음")
    hs_set = {digits(c.get("hs_code")) for c in conf}
    if len(hs_set) > 1:
        return "ambiguous", None, "같은 모델 후보 HS 가 갈림: " + " · ".join(sorted(hs_set))
    best = sorted(conf, key=lambda c: (0 if (c.get("vn_name") or "").strip() else 1, -(c.get("seen_count") or 0)))[0]
    return "token", best, "모델만 일치: 사전 규격 " + (best.get("model") or best.get("part_key") or "")


# ─────────────────────────────────────────────────────────────
# 엑셀 읽기 — 시트·머리글·품목 줄
# ─────────────────────────────────────────────────────────────
def _sheet_paths(z) -> list:
    """workbook.xml + rels → [(시트 이름, 상태, xml 경로)] (문서 순서)"""
    wb = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid2t = {}
    for m in re.finditer(r"<Relationship\b([^>]*)/?>", rels):
        a = m.group(1)
        i = re.search(r'\bId="([^"]+)"', a)
        t = re.search(r'\bTarget="([^"]+)"', a)
        if i and t:
            rid2t[i.group(1)] = t.group(1)
    out = []
    for m in re.finditer(r"<sheet\b([^>]*)/?>", wb):
        a = m.group(1)
        nm = re.search(r'\bname="([^"]*)"', a)
        st = re.search(r'\bstate="([^"]*)"', a)
        rid = re.search(r'\br:id="([^"]+)"', a) or re.search(r'\bid="([^"]+)"', a)
        if not (nm and rid and rid.group(1) in rid2t):
            continue
        tgt = rid2t[rid.group(1)]
        path = tgt.lstrip("/") if tgt.startswith("/") else "xl/" + tgt
        name = nm.group(1).replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"').replace("&apos;", "'")
        out.append((name, (st.group(1) if st else "visible"), path))
    return out


def _num(v):
    """엑셀 칸 → 숫자(수식은 저장된 계산 값 · '1,234.5' 같은 글자 숫자도) · 아니면 None."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _inv_date(top: dict) -> str:
    """머리 칸(1~15줄) 'DATE:' 바로 오른쪽 값 → 'YYYY-MM-DD'(날짜·엑셀 날짜 숫자·'2026-09-21…' 글자) · 모르면 ''."""
    import datetime as _dt
    for (r, c), v in sorted(top.items()):
        if not (isinstance(v, str) and HEAD_DATE.match(v)):
            continue
        d = top.get((r, c + 1))
        if isinstance(d, _dt.datetime):
            return d.date().isoformat()
        if isinstance(d, _dt.date):
            return d.isoformat()
        if isinstance(d, (int, float)) and not isinstance(d, bool) and 30000 < d < 80000:
            from openpyxl.utils.datetime import from_excel
            return from_excel(d).date().isoformat()
        m = re.match(r"\s*(20\d\d)[-./](\d{1,2})[-./](\d{1,2})", str(d or ""))
        if m:
            return "%s-%02d-%02d" % (m.group(1), int(m.group(2)), int(m.group(3)))
    return ""


def scan(data: bytes) -> dict:
    """인보이스 엑셀 → {sheets:[{title, path, header_row, col_desc, col_hs, col_vn, form, items:[…]}], invoice_no}
    품목 줄 = 머리글 아래에서 Description 칸에 글이 있는 줄(합계 줄 제외).
    z1191p: 줄마다 qty·price·amount(수식은 저장된 계산 값 · 없으면 None) · 시트에 col_qty·col_price·col_amount ·
            결과에 date(머리 칸 'DATE:' 옆 인보이스 날짜 'YYYY-MM-DD' · 없으면 '')."""
    import openpyxl
    z = zipfile.ZipFile(io.BytesIO(data))
    paths = _sheet_paths(z)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        def rows_of(title, min_row, max_row, max_col):
            if title not in wb.sheetnames:
                return None
            return wb[title].iter_rows(min_row=min_row, max_row=max_row, max_col=max_col, values_only=True)
        return _scan_core(paths, rows_of)
    finally:
        wb.close()


def scan_fast(data: bytes) -> dict:
    """z1191u: scan 과 **같은 규칙**(_scan_core) — 칸 값만 엑셀 XML 에서 바로 읽는다(서식 표를 안 읽어 빠르다).
    다른 점 하나: 날짜 칸이 날짜가 아니라 엑셀 날짜 숫자로 온다(_inv_date 는 둘 다 읽는다).
    최종 양식 틀은 서식이 3만 개 넘게 쌓여 openpyxl 로 한 번 읽는 데 2~3초 — 옛 양식 → 최종 양식 바꾸기에서 쓴다."""
    z = zipfile.ZipFile(io.BytesIO(data))
    paths = _sheet_paths(z)
    ss = shared_strings(z)
    cache = {}

    def rows_of(title, min_row, max_row, max_col):
        p = next((pp for t, _, pp in paths if t == title), None)
        if p is None or p not in z.namelist():
            return None
        if p not in cache:
            cache[p] = sheet_cells(z, p, ss)
        cells = cache[p]
        top = max((r for r, _ in cells), default=0)
        hi = top if max_row is None else min(max_row, top)
        return (tuple((cells.get((r, c)) or (None, None))[0] for c in range(1, max_col + 1)) for r in range(min_row, hi + 1))
    return _scan_core(paths, rows_of)


def shared_strings(z) -> list:
    """z1191u: 글자 표(sharedStrings) → 목록(서식 있는 글 조각은 이어 붙임 · 발음 표시는 뺌)."""
    if "xl/sharedStrings.xml" not in z.namelist():
        return []
    x = z.read("xl/sharedStrings.xml").decode("utf-8")
    out = []
    for m in re.finditer(r"<si\b[^>]*?(?:/>|>(.*?)</si>)", x, re.S):
        body = re.sub(r"<rPh\b.*?</rPh>", "", m.group(1) or "", flags=re.S)
        out.append(_text("".join(re.findall(r"<t\b[^>]*>(.*?)</t>", body, re.S))))
    return out


_ENT = re.compile(r"&(#[xX][0-9A-Fa-f]+|#[0-9]+|lt|gt|quot|apos|amp);")
_ENT_NAMED = {"lt": "<", "gt": ">", "quot": '"', "apos": "'", "amp": "&"}


def _ent(m):
    t = m.group(1)
    if t[0] != "#":
        return _ENT_NAMED[t]
    try:
        return chr(int(t[2:], 16) if t[1] in "xX" else int(t[1:]))
    except (ValueError, OverflowError):
        return m.group(0)


def _unesc(s: str) -> str:
    """XML 글자 풀기 — 한 번에(&#38;lt; 는 글자 「&lt;」) · 숫자 표기(&#44288; · &#x1F4A1;)도 푼다
    (openpyxl 같은 도구는 한글·베트남어를 숫자 표기로 적는다 — 10-28 실측: 「관리코드」 = &#44288;&#47532;&#53076;&#46300;)."""
    return _ENT.sub(_ent, s) if "&" in s else s


def _text(s: str) -> str:
    """XML 글자 → 칸 글자: 줄바꿈은 XML 규칙대로 \\n 하나로 · &lt; 같은 표기 풀기 · 엑셀 _x000D_ 표기 풀기(openpyxl 과 같게)."""
    s = _unesc(s.replace("\r\n", "\n").replace("\r", "\n"))
    if "_x" in s:
        s = re.sub(r"_x([0-9A-Fa-f]{4})_", lambda m: chr(int(m.group(1), 16)), s)
    return s


def sheet_cells(z, path, ss) -> dict:
    """z1191u: 시트 XML → {(줄, 칸): (값, 수식 글자 또는 None)} — 값은 저장된 값(수식이면 계산해 둔 값 · 날짜는 엑셀 숫자).
    공유 수식은 주인 칸 수식을 그 칸 자리로 옮겨 적는다."""
    return sheet_cells_xml(z.read(path).decode("utf-8"), ss)


def sheet_cells_xml(x: str, ss) -> dict:
    """sheet_cells 와 같음 — 시트 XML 글자를 바로 받는다(변환기가 만든 시트를 저장 전에 읽을 때)."""
    out, shared = {}, {}
    for cm in re.finditer(r'<c\b[^>]*?\br="([A-Z]+)(\d+)"[^>]*?(?:/>|>(.*?)</c>)', x, re.S):
        col, row = _col_num(cm.group(1)), int(cm.group(2))
        cell = cm.group(0)
        head = cell[:cell.index(">") + 1]
        tm = re.search(r'\bt="(\w+)"', head)
        t = tm.group(1) if tm else "n"
        body = cm.group(3) or ""
        val = None
        if t == "inlineStr":
            val = _text("".join(re.findall(r"<t\b[^>]*>(.*?)</t>", body, re.S)))
        else:
            vm = re.search(r"<v>(.*?)</v>", body, re.S)
            if vm is not None:
                raw = vm.group(1)
                if t == "s":
                    val = ss[int(raw)] if raw.strip().isdigit() and int(raw) < len(ss) else None
                elif t == "b":
                    val = raw.strip() == "1"
                elif t in ("str", "e"):
                    val = _text(raw)
                else:
                    try:
                        val = int(raw) if re.fullmatch(r"-?\d+", raw.strip()) else float(raw)
                    except ValueError:
                        val = _unesc(raw)
        ftext = None
        fm = re.search(r"<f\b([^>]*?)(?:/>|>(.*?)</f>)", body, re.S)
        if fm:
            if 't="shared"' in fm.group(1):
                si = re.search(r'\bsi="(\d+)"', fm.group(1))
                si = si.group(1) if si else ""
                if fm.group(2):
                    shared[si] = (col, row, _unesc(fm.group(2)))
                    ftext = _unesc(fm.group(2))
                elif si in shared:
                    c0, r0, tx = shared[si]
                    ftext = translate_formula(tx, row - r0, col - c0)
            elif fm.group(2) is not None:
                ftext = _unesc(fm.group(2))
        if val is not None or ftext is not None:
            out[(row, col)] = (val, ftext)
    return out


_FREF = re.compile(r"((?:'[^']+'|[A-Za-z_][\w.]*)!)?(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\w(])")


def translate_formula(f: str, dr: int, dc: int = 0) -> str:
    """z1191u: 채우기처럼 옮기기 — $ 없는 줄·칸 주소만 dr·dc 만큼. 다른 시트 주소도 엑셀 채우기처럼 옮기고
    (공유 수식 =INVOICE!D34 → 아래 칸 =INVOICE!D35) 따옴표 글자 안은 그대로 둔다."""
    out, pos, in_str = [], 0, False
    for m in _FREF.finditer(f):
        seg = f[pos:m.start()]
        in_str ^= (seg.count('"') % 2 == 1)
        out.append(seg)
        pos = m.end()
        s = m.start()
        if in_str or (s > 0 and (f[s - 1].isalnum() or f[s - 1] in "_.")):
            out.append(m.group(0))
            continue
        col, row = m.group(3), int(m.group(5))
        if not m.group(2) and dc:
            col = _col_letter(_col_num(col) + dc)
        if not m.group(4):
            row += dr
        out.append("%s%s%s%s%d" % (m.group(1) or "", m.group(2), col, m.group(4), row))
    out.append(f[pos:])
    return "".join(out)


def _scan_core(paths, rows_of) -> dict:
    """scan · scan_fast 의 본체 — rows_of(시트, 첫 줄, 끝 줄 또는 None, 끝 칸) 은 줄마다 값 튜플(빈 칸 None)을 낸다."""
    res = {"sheets": [], "invoice_no": "", "date": ""}
    for title, state, path in paths:
        if state != "visible" or WS_SKIP.search(title or ""):
            continue
        first_rows = rows_of(title, 1, 45, 30)
        if first_rows is None:
            continue
        head, top = [], {}
        for r, row in enumerate(first_rows, start=1):
            for c, v in enumerate(row, start=1):
                if v is None:
                    continue
                t = str(v).strip()
                head.append((r, c, t))
                if r <= 15:
                    top[(r, c)] = v
                if not res["invoice_no"] and r <= 15:
                    mm = INVOICE_NO.search(t)
                    if mm:
                        res["invoice_no"] = mm.group(0)
        hs_cells = [(r, c, t) for r, c, t in head if HEAD_HS.search(t) and len(t) < 20]
        if not hs_cells:
            continue
        exact = [x for x in hs_cells if HEAD_HS_EXACT.match(x[2])]
        hr, hc, _ = (exact or hs_cells)[0]
        vn = next(((r, c) for r, c, t in head if HEAD_VN.search(t)), None)
        ds = next(((r, c) for r, c, t in head if HEAD_DESC.search(t)), None)
        if not ds:
            continue
        if not res["date"]:
            res["date"] = _inv_date(top)
        words = {t.upper() for r, c, t in head if r == hr}
        if {"CATEGORY", "SUPPLIER"} <= words:
            form = "수책"
        elif "IMAGE" in words and "C/O" in words:
            form = "소모품"
        else:
            form = "예전 양식"
        # z1191p: 수량·단가·금액 칸 — Description·HS 머리글 줄 근처의 A~L 열만
        #   (예전 양식은 31줄에 품명·수량·단가·금액 · HS 머리글은 32~33줄 — 오른쪽의 'Duty AMOUNT' 같은 칸은 안 본다)
        near = range(min(ds[0], hr) - 1, max(ds[0], hr) + 1)

        def _col(rx):
            return next((c for r, c, t in head if r in near and c <= 12 and len(t) < 30 and rx.search(t)), None)
        cq, cp, ca = _col(HEAD_QTY), _col(HEAD_PRICE), _col(HEAD_AMT)

        def _at(row, col):
            return _num(row[col - 1]) if (col and len(row) >= col) else None
        start = max(hr, ds[0], vn[0] if vn else hr) + 1
        items = []
        mc = max(hc, vn[1] if vn else hc, ds[1], 4, cq or 0, cp or 0, ca or 0)
        for r, row in enumerate(rows_of(title, start, None, mc) or (), start=start):
            # 합계(TOTAL) 줄에서 품목 읽기를 멈춘다 — 그 아래 서명·비고 글을 품목으로 오인해 칠하지 않게(최종 양식은 A열 'TOTAL')
            if any(isinstance(v, str) and TOTAL_ROW.match(v) for v in row[:4]):
                break
            d = row[ds[1] - 1] if len(row) >= ds[1] else None
            d = "" if d is None else str(d).strip()
            if not d:
                continue
            h = row[hc - 1] if len(row) >= hc else None
            v = row[vn[1] - 1] if (vn and len(row) >= vn[1]) else None
            items.append({"row": r, "desc": d, "hs_raw": "" if h is None else str(h).strip(),
                          "vn_raw": "" if v is None else str(v).strip(),
                          "qty": _at(row, cq), "price": _at(row, cp), "amount": _at(row, ca)})
        res["sheets"].append({"title": title, "path": path, "header_row": hr, "col_desc": ds[1], "col_hs": hc,
                              "col_vn": vn[1] if vn else None, "form": form, "items": items,
                              "col_qty": cq, "col_price": cp, "col_amount": ca})
    return res


# ─────────────────────────────────────────────────────────────
# 채움 계획 — 줄마다 무엇을 쓰고 무엇을 칠할지
# ─────────────────────────────────────────────────────────────
def plan(sc, index, overrides=None) -> list:
    """줄마다 {sheet, row, key, desc, status, write_hs, write_vn, color_hs, color_vn, why, dict_hs,
              hs_written, dict_seen, dict_last, dict_model, dict_status, choices}
    z1191n: overrides = {규격 키: HS} — 본사가 ① 결과에서 고른 값. 그 규격 줄은 적힌 값과 달라도 그 값으로 쓰고 칠하지 않는다.
            choices = 본사가 그 자리에서 고를 수 있는 HS [{hs, kind: 'file'(적힌 값) | 'dict'(사전 값)}] — 없으면 []."""
    acts = []
    ov_map = overrides or {}

    def _junk_done(a, vn_raw):
        """z1191v: 품명 칸에 적힌 엑셀 함수 값 — 사전 품명으로 덮었으면 그대로, 아니면 칸을 비우고 노랑(법인 입력).
        (수식이 든 칸이면 apply 가 값을 안 쓰고 '직접 확인'으로 바꾼다 — 수식은 건드리지 않는다)"""
        if a["write_vn"]:
            tail = "품명 칸 「%s」(엑셀 함수 값) → 사전 품명으로 채움" % vn_raw[:12]
        else:
            a["write_vn"] = ""
            a["color_vn"] = a["color_vn"] or YELLOW
            if a["status"] == KEEP:
                a["status"] = NEED
            tail = "품명 칸 「%s」은 엑셀 함수 값이라 비움(법인 입력)" % vn_raw[:12]
        a["why"] = (a["why"] + " · " if a["why"] else "") + tail

    for sh in sc["sheets"]:
        has_vn = sh["col_vn"] is not None
        for it in sh["items"]:
            vn_raw = (it["vn_raw"] or "").strip()
            vn_junk = has_vn and formula_junk(vn_raw)        # z1191v: 적힌 「0」·「#N/A」 = 빈칸으로 본다
            hs0, vn0 = digits(it["hs_raw"]), ("" if vn_junk else vn_raw)
            key = norm_key(it["desc"])
            way, e, memo = lookup(index, it["desc"])
            dhs = digits(e.get("hs_code")) if e else ""
            dvn = (e.get("vn_name") or "").strip() if e else ""
            if formula_junk(dvn):
                dvn = ""                                     # z1191v: 사전에 남은 함수 값도 품명으로 쓰지 않는다
            a = {"sheet": sh["title"], "row": it["row"], "key": key, "desc": it["desc"], "status": KEEP,
                 "write_hs": None, "write_vn": None, "color_hs": None, "color_vn": None,
                 "why": "", "dict_hs": dhs if way in ("exact", "token", "pending") else "",
                 # z1191n: 본사가 고를 때 보는 근거 — 사전 값이 몇 번 나왔고 마지막이 언제인지 · 다른 규격(모델만 일치)이면 그 규격
                 "hs_written": str(it["hs_raw"] or "").strip(),
                 "dict_seen": int(e.get("seen_count") or 0) if e else 0,
                 "dict_last": str(e.get("last_seen") or "") if e else "",
                 "dict_model": (e.get("model") or e.get("part_key") or "") if (e and e.get("part_key") != key) else "",
                 "dict_status": str(e.get("status") or "") if e else "",
                 "vn_written": vn0, "dict_vn": dvn,
                 "choices": []}
            ov = ov_map.get(key)
            if ov:
                # z1191n: 본사가 고른 값 — 적힌 값과 다르면 그 값으로 바꿔 쓴다(적힌 값은 원래 덮지 않는다 — 본사가 고른 규격만)
                a["why"] = "본사 확정 %s" % ov
                if hs0 != ov:
                    a["write_hs"] = ov
                if has_vn and not vn0:
                    if way == "exact" and dhs == ov and dvn:
                        a["write_vn"] = dvn
                    else:
                        a.update(color_vn=YELLOW, why=a["why"] + " · 품명은 사전에 없음")
                a["status"] = AUTO if (a["write_hs"] is not None or a["write_vn"] is not None) else KEEP
                if vn_junk:
                    _junk_done(a, vn_raw)
                acts.append(a)
                continue
            if it["hs_raw"] and len(hs0) not in (6, 8):
                a.update(status=CONFLICT, color_hs=ORANGE, why="HS 칸 형식 확인(8자리 숫자가 아님): %s" % it["hs_raw"][:20])
                if way in ("exact", "token") and len(dhs) == 8:
                    a["choices"] = [{"hs": dhs, "kind": "dict"}]
                if has_vn and not vn0:
                    a["color_vn"] = YELLOW
            elif hs0:
                # 이미 적힌 HS(본사가 미리 채움 · 장비 6자리 포함) — 값은 절대 덮지 않는다
                if way in ("exact", "token") and len(hs0) == 8 and dhs != hs0:
                    a.update(status=CONFLICT, color_hs=ORANGE,
                             why="적힌 HS %s ↔ 사전 확정 %s" % (hs0, dhs) + (" (%s)" % memo if way == "token" else ""),
                             choices=[{"hs": hs0, "kind": "file"}, {"hs": dhs, "kind": "dict"}])
                elif way in ("exact", "token") and dhs == hs0 and has_vn and not vn0 and dvn:
                    a.update(status=AUTO, write_vn=dvn, why="HS 는 이미 같음 · 품명만 채움")
                elif has_vn and not vn0:
                    a.update(status=NEED, color_vn=YELLOW, why="베트남어 품명 비어 있음")
            elif way == "exact":
                a.update(status=AUTO, write_hs=dhs, why="사전 확정")
                if has_vn and not vn0:
                    if dvn:
                        a["write_vn"] = dvn
                    else:
                        a.update(color_vn=YELLOW, why="사전 확정 · 품명은 사전에 없음")
            elif way == "token":
                a.update(status=CHECK, write_hs=dhs, color_hs=ORANGE, why=memo, choices=[{"hs": dhs, "kind": "dict"}])
                if has_vn and not vn0:
                    if dvn:
                        a.update(write_vn=dvn, color_vn=ORANGE)
                    else:
                        a["color_vn"] = YELLOW
            else:
                a.update(status=NEED, color_hs=YELLOW,
                         why={"none": "사전에 없음(새 품목)"}.get(way, memo or "사전에 없음"))
                # z1191n: 사전에 이 규격이 '확인 대기'·'AI 추정'으로 있으면 본사가 그 값을 바로 확정할 수 있다
                if way == "pending" and e is not None and e.get("part_key") == key and len(dhs) == 8:
                    a["choices"] = [{"hs": dhs, "kind": "dict"}]
                if has_vn and not vn0:
                    a["color_vn"] = YELLOW
            if vn_junk:
                _junk_done(a, vn_raw)
            acts.append(a)
    return acts


def summary(acts) -> dict:
    n = {k: 0 for k in (AUTO, CHECK, CONFLICT, NEED, KEEP)}
    for a in acts:
        n[a["status"]] += 1
    n["total"] = len(acts)
    return n


def category(a) -> str:
    """z1191n: 화면·기록에서 세는 칸 — 노란 칸이 하나라도 있으면 '법인 입력', 주황(또는 확인 상태)이면 '확인'.
    상태만 세면 '사전 확정 · 품명은 사전에 없음' 줄(HS 자동 · 품명 노랑)이 '자동'으로 세져 노란 칸이 숫자·요청 글에서 빠졌다
    (운영 사전 확정 1,989 중 63줄이 품명 없음 — 10-27 실측)."""
    cs = (a.get("color_hs"), a.get("color_vn"))
    if YELLOW in cs or a["status"] == NEED:
        return NEED
    if ORANGE in cs or a["status"] in (CHECK, CONFLICT):
        return CHECK
    return AUTO if a["status"] == AUTO else KEEP


def counts(acts) -> dict:
    """z1191n: {total, auto, check, need, keep, no_reply} — no_reply = 노랑·주황이 하나도 없음(법인 회신 필요 없음 · 받은 파일이 최종)."""
    n = {"total": len(acts), "auto": 0, "check": 0, "need": 0, "keep": 0}
    for a in acts:
        n[{AUTO: "auto", CHECK: "check", NEED: "need", KEEP: "keep"}[category(a)]] += 1
    n["no_reply"] = (n["need"] == 0 and n["check"] == 0)
    return n


# ─────────────────────────────────────────────────────────────
# XML 에서 칸만 바꾸기
# ─────────────────────────────────────────────────────────────
def _col_letter(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _col_num(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n


_CELL_RE = r'<c\b[^>]*?\br="%s"[^>]*?(?:/>|>.*?</c>)'


def _add_styles(styles: str, need: set):
    """need = {(원래 스타일 번호, 색)} → (새 styles.xml, {(원래, 색): 새 번호})
    채우기(fill) 2개를 더하고, 쓰이는 원래 서식(xf)을 복제해 그 채우기만 바꾼다(테두리·글꼴·맞춤은 그대로)."""
    if not need:
        return styles, {}
    fm = re.search(r"<fills\b[^>]*>(.*?)</fills>", styles, re.S)
    if not fm:
        raise ValueError("styles.xml 에 fills 가 없습니다")
    n_fill = len(re.findall(r"<fill\b", fm.group(1)))
    colors = sorted({c for _, c in need})
    fid = {c: n_fill + i for i, c in enumerate(colors)}
    add = "".join('<fill><patternFill patternType="solid"><fgColor rgb="%s"/><bgColor indexed="64"/></patternFill></fill>' % c
                  for c in colors)
    fills_new = re.sub(r'\bcount="\d+"', 'count="%d"' % (n_fill + len(colors)), fm.group(0)[:fm.group(0).index(">") + 1], 1) \
        + fm.group(1) + add + "</fills>"
    styles = styles[:fm.start()] + fills_new + styles[fm.end():]
    cm = re.search(r"<cellXfs\b[^>]*>(.*?)</cellXfs>", styles, re.S)
    if not cm:
        raise ValueError("styles.xml 에 cellXfs 가 없습니다")
    xfs = re.findall(r"<xf\b[^>]*?(?:/>|>.*?</xf>)", cm.group(1), re.S)
    mapping, extra = {}, []
    for (s0, color) in sorted(need):
        base = xfs[s0] if 0 <= s0 < len(xfs) else xfs[0]
        tag_end = base.index(">")
        head, rest = base[:tag_end], base[tag_end:]
        if head.endswith("/"):
            head, rest = head[:-1], "/" + rest
        head = re.sub(r'\s(fillId|applyFill)="[^"]*"', "", head)
        head += ' fillId="%d" applyFill="1"' % fid[color]
        extra.append(head + rest)
        mapping[(s0, color)] = len(xfs) + len(extra) - 1
    open_tag = cm.group(0)[:cm.group(0).index(">") + 1]
    open_tag = re.sub(r'\bcount="\d+"', 'count="%d"' % (len(xfs) + len(extra)), open_tag, 1)
    cell_new = open_tag + cm.group(1) + "".join(extra) + "</cellXfs>"
    styles = styles[:cm.start()] + cell_new + styles[cm.end():]
    return styles, mapping


def _cell_style(cell_xml: str) -> int:
    m = re.search(r'\bs="(\d+)"', cell_xml[:cell_xml.index(">") + 1])
    return int(m.group(1)) if m else 0


def _edit_sheet(xml: str, edits: list, style_of_new, skipped=None) -> tuple:
    """edits: [(줄, 칸 번호, 값 또는 None, 색 또는 None)] → (새 xml, 필요한 스타일 집합, 칸별 원래 스타일)
    1차: 원래 스타일만 모은다(style_of_new=None). 2차: 새 스타일 번호로 실제로 바꾼다.
    ⚠수식이 든 칸에는 값을 쓰지 않는다(색만) — skipped 에 칸 주소를 남긴다."""
    need, orig = set(), {}
    by_row = {}
    for r, c, val, color in edits:
        by_row.setdefault(r, []).append((c, val, color))
    out, pos = [], 0
    for rm in re.finditer(r'<row\b[^>]*?\br="(\d+)"[^>]*?(?:/>|>(.*?)</row>)', xml, re.S):
        r = int(rm.group(1))
        if r not in by_row:
            continue
        row_xml = rm.group(0)
        self_closing = rm.group(2) is None and row_xml.rstrip().endswith("/>")
        if self_closing:
            row_open = row_xml[:-2] + ">"
            body = ""
        else:
            row_open = row_xml[:row_xml.index(">") + 1]
            body = rm.group(2) or ""
        for c, val, color in sorted(by_row[r]):
            ref = "%s%d" % (_col_letter(c), r)
            cm = re.search(_CELL_RE % ref, body, re.S)
            if cm:
                s0 = _cell_style(cm.group(0))
            else:
                # 칸이 없으면 같은 줄 왼쪽 칸이 아니라 '바로 위 품목 줄 같은 칸' 서식을 빌린다 — 없으면 0
                s0 = orig.get(("up", c), 0)
            orig[(r, c)] = s0
            orig[("up", c)] = s0
            if color:
                need.add((s0, color))
            if style_of_new is None:
                continue
            s_new = style_of_new.get((s0, color), s0) if color else s0
            if val is not None and cm and "<f" in cm.group(0):
                if skipped is not None:
                    skipped.append(ref)
                val = None
            if val == "":
                new_cell = '<c r="%s" s="%d"/>' % (ref, s_new)       # z1191v: 값 비우기(적힌 함수 값 「0」 등) — 서식·색만 남긴다
            elif val is not None:
                new_cell = '<c r="%s" s="%d" t="inlineStr"><is><t xml:space="preserve">%s</t></is></c>' % (
                    ref, s_new, _xml_escape(val))
            elif cm:
                head = cm.group(0)[:cm.group(0).index(">") + 1]
                if re.search(r'\bs="\d+"', head):
                    head2 = re.sub(r'\bs="\d+"', 's="%d"' % s_new, head, 1)
                else:
                    head2 = head[:-2] + ' s="%d"/>' % s_new if head.endswith("/>") else head[:-1] + ' s="%d">' % s_new
                new_cell = head2 + cm.group(0)[len(head):]
            else:
                new_cell = '<c r="%s" s="%d"/>' % (ref, s_new)
            if cm:
                body = body[:cm.start()] + new_cell + body[cm.end():]
            else:
                ins = len(body)
                for cc in re.finditer(r'<c\b[^>]*?\br="([A-Z]+)\d+"', body):
                    if _col_num(cc.group(1)) > c:
                        ins = cc.start()
                        break
                body = body[:ins] + new_cell + body[ins:]
                sp = re.search(r'\bspans="(\d+):(\d+)"', row_open)
                if sp and c > int(sp.group(2)):
                    row_open = row_open.replace(sp.group(0), 'spans="%s:%d"' % (sp.group(1), c))
        if style_of_new is not None:
            out.append(xml[pos:rm.start()])
            out.append(row_open + body + "</row>")
            pos = rm.end()
    if style_of_new is None:
        return xml, need, orig
    out.append(xml[pos:])
    xml2 = "".join(out)
    # 칸이 기존 범위 밖이면 dimension 을 넓힌다(엑셀 도구가 범위 밖 칸을 못 읽는 일 방지)
    dm = re.search(r'<dimension ref="([A-Z]+)(\d+):([A-Z]+)(\d+)"/>', xml2)
    if dm and edits:
        mc = max(c for _, c, _, _ in edits)
        if mc > _col_num(dm.group(3)):
            xml2 = xml2.replace(dm.group(0), '<dimension ref="%s%s:%s%s"/>' % (dm.group(1), dm.group(2), _col_letter(mc), dm.group(4)))
    return xml2, need, orig


def apply(data: bytes, sc: dict, acts: list, skipped=None) -> bytes:
    """계획대로 칸 값·색만 바꾼 새 엑셀(bytes). 바뀌는 항목 = 해당 시트 xml · styles.xml · workbook.xml(다시 계산 표시).
    skipped(목록)를 주면 수식이 있어 값을 안 쓴 칸 주소를 담는다."""
    z = zipfile.ZipFile(io.BytesIO(data))
    sheet_of = {s["title"]: s for s in sc["sheets"]}
    per_sheet = {}
    for a in acts:
        sh = sheet_of[a["sheet"]]
        e = per_sheet.setdefault(sh["path"], [])
        if a["write_hs"] is not None or a["color_hs"]:
            e.append((a["row"], sh["col_hs"], a["write_hs"], a["color_hs"]))
        if sh["col_vn"] and (a["write_vn"] is not None or a["color_vn"]):
            e.append((a["row"], sh["col_vn"], a["write_vn"], a["color_vn"]))
    per_sheet = {p: e for p, e in per_sheet.items() if e}
    if not per_sheet:
        return data
    need = set()
    sheet_xml = {}
    for p, edits in per_sheet.items():
        x = z.read(p).decode("utf-8")
        sheet_xml[p] = x
        _, n, _ = _edit_sheet(x, edits, None)
        need |= n
    styles = z.read("xl/styles.xml").decode("utf-8")
    styles2, mapping = _add_styles(styles, need)
    changed = {"xl/styles.xml": styles2.encode("utf-8")}
    for p, edits in per_sheet.items():
        x2, _, _ = _edit_sheet(sheet_xml[p], edits, mapping, skipped)
        changed[p] = x2.encode("utf-8")
    wbx = z.read("xl/workbook.xml").decode("utf-8")
    cp = re.search(r"<calcPr\b[^>]*?/>", wbx)
    if cp and "fullCalcOnLoad" not in cp.group(0):
        wbx = wbx.replace(cp.group(0), cp.group(0)[:-2] + ' fullCalcOnLoad="1"/>')
        changed["xl/workbook.xml"] = wbx.encode("utf-8")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zo:
        for info in z.infolist():
            buf = changed.get(info.filename)
            if buf is None:
                buf = z.read(info.filename)
            zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            zi.compress_type = zipfile.ZIP_DEFLATED if info.filename in changed else info.compress_type
            zi.external_attr = info.external_attr
            zi.create_system = info.create_system
            zo.writestr(zi, buf)
    return out.getvalue()


def invoice_no_quick(data: bytes) -> str:
    """회신 파일의 인보이스 번호(KNK-YYMMDD-n) — 글자 표(sharedStrings)·시트 안 글자에서 처음 나오는 것. 없으면 ''."""
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        return ""
    names = [n for n in ("xl/sharedStrings.xml",) if n in z.namelist()]
    names += [p for _, _, p in _sheet_paths(z) if p in z.namelist()]
    for n in names:
        m = INVOICE_NO.search(z.read(n).decode("utf-8", "replace"))
        if m:
            return m.group(0)
    return ""


def file_stem(name: str) -> str:
    """파일 이름 → 비교용 앞부분: 확장자·회신 꼬리(-re · _RE_DOAN · -LAN · (2) …)를 뗀 소문자."""
    s = re.sub(r"\.[A-Za-z0-9]{3,4}$", "", str(name or "")).strip().lower()
    prev = None
    while prev != s:
        prev = s
        s = re.sub(r"[\s_\-]*(re|rev\d*|doan|lan|final|ok|v\d+|\(\d{1,2}\))$", "", s)   # (2) 같은 짧은 번호만 — (261022) 날짜는 남긴다
    return s


def fill(data: bytes, entries, overrides=None) -> dict:
    """한 번에: 읽기 → 계획 → 칸 바꾸기. 반환 {ok, out, scan, acts, summary, counts, error}
    z1191n: overrides = 본사가 고른 값({규격 키: HS}) · counts = 칸 색으로 센 숫자(no_reply 포함)."""
    try:
        sc = scan(data)
    except zipfile.BadZipFile:
        return {"ok": False, "error": "엑셀(.xlsx) 파일이 아닙니다 — 예전 형식(.xls)이면 .xlsx 로 저장해서 올려 주세요"}
    if not sc["sheets"]:
        return {"ok": False, "error": "인보이스 시트를 찾지 못했습니다 — 'HS CODE'·'Description' 머리글이 있는 시트가 필요합니다"}
    if not any(s["items"] for s in sc["sheets"]):
        return {"ok": False, "error": "품목 줄이 없습니다 — Description 칸에 규격이 적힌 줄이 필요합니다"}
    idx = build_index(entries)
    acts = plan(sc, idx, overrides)
    skipped = []
    out = apply(data, sc, acts, skipped)
    if skipped:
        sheet_of = {s["title"]: s for s in sc["sheets"]}
        sk = set(skipped)
        for a in acts:
            sh = sheet_of[a["sheet"]]
            refs = {"%s%d" % (_col_letter(sh["col_hs"]), a["row"])}
            if sh["col_vn"]:
                refs.add("%s%d" % (_col_letter(sh["col_vn"]), a["row"]))
            if refs & sk:
                a.update(status=CONFLICT, why="수식이 든 칸이라 값을 쓰지 않았습니다 — 직접 확인", choices=[])
    return {"ok": True, "out": out, "scan": sc, "acts": acts, "summary": summary(acts), "counts": counts(acts), "skipped": skipped}


# ─────────────────────────────────────────────────────────────
# z1191p 지난 수출 이력 — 대표 2026-10-27 「예전 파일을 보고 HS코드 및 금액도 비교할수있게 미리 등록」 → 「진행해」
#   (제안: 비교 기준 ±20% · 사전에 없는 규격은 사전에 추가 · 앞으로 ③ 회신도 이력에 쌓기)
# ─────────────────────────────────────────────────────────────
KIND_BY_FORM = {"수책": "수책통관", "소모품": "일반통관"}
HIST_KINDS = ("수책통관", "일반통관")
HIST_PRICE_TOL = 20.0    # 단가 비교 기준 ±% — 같은 품목 1,580쌍 중 96%가 ±20% 안(환율로 조금씩 흔들린다 · 10-27 실측)


def kind_of(sc) -> str:
    """양식 → 통관 구분(시트가 모두 같은 양식일 때만) · 예전 양식은 모름('')."""
    ks = {KIND_BY_FORM.get(s["form"], "") for s in sc.get("sheets", [])}
    return ks.pop() if len(ks) == 1 else ""


def date_of(sc) -> str:
    """인보이스 날짜 — 머리 칸 'DATE:' 옆 값, 없으면 번호(KNK-YYMMDD-n)에서 · 모르면 ''."""
    if sc.get("date"):
        return sc["date"]
    m = re.match(r"KNK-(\d{2})(\d{2})(\d{2})-", sc.get("invoice_no") or "")
    if m and 1 <= int(m.group(2)) <= 12 and 1 <= int(m.group(3)) <= 31:
        return "20%s-%s-%s" % m.groups()
    return ""


def hist_lines(sc) -> list:
    """품목 줄 → 이력 줄 [{sheet,row,key,desc,qty,price,amount,hs,vn}].
    HS 는 8자리만(7자리는 앞에 0 — 엑셀이 숫자로 바꿔 앞 0 을 잃은 것 · 사전 학습과 같은 규칙) · 6자리 장비 HS·빈 칸은 ''.
    금액 칸이 비어 있으면(수식인데 저장된 계산 값이 없는 파일) 수량×단가 — 금액 칸 수식이 그 계산이다."""
    out = []
    for sh in sc.get("sheets", []):
        for it in sh["items"]:
            hs = digits(it["hs_raw"])
            if len(hs) == 7:
                hs = "0" + hs
            q, p, a = it.get("qty"), it.get("price"), it.get("amount")
            if a is None and q is not None and p is not None:
                a = round(q * p, 6)
            out.append({"sheet": sh["title"], "row": it["row"], "key": norm_key(it["desc"]), "desc": it["desc"],
                        "qty": q, "price": p, "amount": a, "hs": hs if len(hs) == 8 else "",
                        "vn": (it.get("vn_raw") or "").strip()})
    return out


def _rels(z, path) -> dict:
    """파트의 관계 → {rId: 묶음 안 전체 경로}('../media/x.png' 같은 상대 경로를 풀어서)."""
    d, b = path.rsplit("/", 1) if "/" in path else ("", path)
    rp = (d + "/" if d else "") + "_rels/" + b + ".rels"
    if rp not in z.namelist():
        return {}
    out = {}
    for m in re.finditer(r"<Relationship\b([^>]*)/?>", z.read(rp).decode("utf-8", "replace")):
        a = m.group(1)
        i = re.search(r'\bId="([^"]+)"', a)
        t = re.search(r'\bTarget="([^"]+)"', a)
        if not (i and t) or re.search(r'\bTargetMode="External"', a):
            continue
        tgt = t.group(1)
        parts, st = (tgt.lstrip("/") if tgt.startswith("/") else (d + "/" + tgt)).split("/"), []
        for p in parts:
            if p == "..":
                if st:
                    st.pop()
            elif p and p != ".":
                st.append(p)
        out[i.group(1)] = "/".join(st)
    return out


def row_images(data: bytes, sc) -> dict:
    """z1191r: 품목 줄에 붙은 그림 → {(시트, 줄): (확장자, 바이트)}.
    그림 자리(drawing 의 from~to 줄)가 품목 줄에 걸리면 그 줄 — 한 줄에 여러 장이면 첫 장만 · 품목 줄 밖(서명·도장)은 안 본다.
    (예전 일반통관 양식은 B열 · 표준 양식은 IMAGE 칸 — 칸은 보지 않고 줄로만 잇는다)"""
    z = zipfile.ZipFile(io.BytesIO(data))
    names = set(z.namelist())
    out = {}
    for sh in sc.get("sheets", []):
        items = {it["row"] for it in sh["items"]}
        if not items:
            continue
        for dp in [t for t in _rels(z, sh["path"]).values() if "/drawings/" in t and t.endswith(".xml") and t in names]:
            dx = z.read(dp).decode("utf-8", "replace")
            dr = _rels(z, dp)
            for am in re.finditer(r"<((?:\w+:)?)(twoCellAnchor|oneCellAnchor)\b.*?</\1\2>", dx, re.S):
                a = am.group(0)
                fr = re.search(r"<(?:\w+:)?from>.*?<(?:\w+:)?row>(\d+)</", a, re.S)
                if not fr:
                    continue
                to = re.search(r"<(?:\w+:)?to>.*?<(?:\w+:)?row>(\d+)</", a, re.S)
                r0 = int(fr.group(1)) + 1
                r1 = max(r0, int(to.group(1)) + 1) if to else r0
                row = next((r for r in range(r0, r1 + 1) if r in items), None)
                if row is None or (sh["title"], row) in out:
                    continue
                for e in re.findall(r'<(?:\w+:)?blip\b[^>]*?\br:embed="([^"]+)"', a):
                    mp = dr.get(e)
                    if mp and mp in names:
                        out[(sh["title"], row)] = (mp.rsplit(".", 1)[-1].lower(), z.read(mp))
                        break
    return out


def compare(acts, prices, past, own_invoice="", tol=HIST_PRICE_TOL) -> dict:
    """① 결과 줄 ↔ 지난 수출 이력 — 규격이 **그대로** 같은 줄만 본다(모델 글자만 같은 다른 품목의 단가를 섞지 않는다).
    acts = plan 결과 · prices = {(시트, 줄): 이번 단가} · past = {규격 키: [{invoice_no, inv_date, kind, hs, price, id}]}
    같은 인보이스(번호)의 이력은 빼고 본다(이 인보이스가 이미 이력에 들어가 있어도 자기와 비교하지 않게).
    반환 {matched, checked, price:[단가가 ±tol% 넘게 다른 줄], hs:[이번 HS ≠ 마지막 수출 HS], ev:{(시트, 줄): 근거 글},
          img:{(시트, 줄): 마지막 수출 사진}} — z1191r: 단가·HS 줄에도 그 규격의 마지막 사진(img)
    ⚠화면에만 쓴다 — 법인에 보내는 파일·이음 방 글에는 넣지 않는다(단가)."""
    out = {"matched": 0, "checked": 0, "price": [], "hs": [], "ev": {}, "img": {}}
    for a in acts:
        rows = [p for p in past.get(a["key"], []) if p.get("invoice_no") != own_invoice]
        if not rows:
            continue
        out["matched"] += 1
        rows.sort(key=lambda p: (p.get("inv_date") or "", p.get("id") or 0))
        invs = {p.get("invoice_no") for p in rows}
        hinv = {}     # HS 마다 몇 번(인보이스 수) — 한 인보이스에 같은 규격이 여러 줄이어도 한 번(「지난 수출 N번」과 같은 단위)
        for p in rows:
            if p.get("hs"):
                hinv.setdefault(p["hs"], set()).add(p.get("invoice_no"))
        hcnt = {h: len(s) for h, s in hinv.items()}
        lh = next((p for p in reversed(rows) if p.get("hs")), None)
        last = rows[-1]
        img = next((p["img"] for p in reversed(rows) if p.get("img")), "")
        if img:
            out["img"][(a["sheet"], a["row"])] = img
        ev = "지난 수출 %d번" % len(invs)
        if hcnt:
            ev += " · HS " + " · ".join("%s×%d" % (h, n) for h, n in sorted(hcnt.items(), key=lambda x: (-x[1], x[0])))
        out["ev"][(a["sheet"], a["row"])] = ev + " · 마지막 %s(%s)" % (last.get("invoice_no") or "", (last.get("inv_date") or "")[:10])
        wr = digits(a.get("hs_written"))
        now_hs = a.get("write_hs") or (wr if len(wr) == 8 else "")
        if now_hs and lh and now_hs != lh["hs"]:
            out["hs"].append({"sheet": a["sheet"], "row": a["row"], "desc": a["desc"], "now": now_hs, "last": lh["hs"],
                              "last_inv": lh.get("invoice_no") or "", "last_date": (lh.get("inv_date") or "")[:10],
                              "counts": hcnt, "n": len(invs), "img": img})
        cur = prices.get((a["sheet"], a["row"]))
        lp = next((p for p in reversed(rows) if (p.get("price") or 0) > 0), None)
        if cur and cur > 0 and lp:
            out["checked"] += 1
            diff = (cur - lp["price"]) / lp["price"] * 100.0
            if abs(diff) > tol:
                ps = [p["price"] for p in rows if (p.get("price") or 0) > 0]
                out["price"].append({"sheet": a["sheet"], "row": a["row"], "desc": a["desc"], "now": round(cur, 4),
                                     "last": round(lp["price"], 4), "diff": round(diff, 1),
                                     "last_inv": lp.get("invoice_no") or "", "last_date": (lp.get("inv_date") or "")[:10],
                                     "min": round(min(ps), 4), "max": round(max(ps), 4), "n": len(ps), "img": img})
    return out
