#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
BOM 부품 사진 이어붙이기 — 공용 한 벌 (대표 지시 2026-09-10)
==============================================================
대표 지시: "인벤터 엑셀 파일에는 부품들 사진이 있어. 그걸 우리 BOM양식에는 지금은
빠지게 되어 있는데, 우리 양식 마지막 열에 사진도 같이 넣어 주면 좋겠어.
영업팀이나 다른 팀들은 참고할 수 있게."
  · 대표 선택 (나) = 간이판에서 끝내지 않고 **통일판까지** 이어간다.
  · 대표 확정 2026-09-10 = "사진은 아주 작게. 필요한 사용자가 **엑셀칸을 키워서**
    확인할 수 있게."

⭐ 그래서 사진을 「칸에 매어」 붙인다 (twoCellAnchor)
  사진을 고정 크기로 붙이면 칸을 키워도 **사진은 그대로**라 대표 지시를 못 지킨다.
  칸 한 개에 매어 두면 줄 높이·열 너비를 늘리는 만큼 사진도 함께 커진다.
  인벤터 원본이 쓰는 방식과 같다(실측 확인).
  → 평소에는 아주 작게 있다가, 볼 사람만 칸을 늘려 크게 본다. 표는 길어지지 않는다.

왜 한 벌로 두는가
  사진은 ①인벤터→간이판 ②간이판→통일판 ③통일판 개정 — 세 곳에서 다뤄진다.
  ⭐같은 개념을 곳마다 따로 그리면 한쪽만 규칙에서 빠진다(2026-09-07 아바타 사건).
  그래서 읽기·넣기·밀기를 여기 한 곳에만 둔다.

실측으로 확인한 사실 (2026-09-10 · 인벤터 실물 AA00~AD00)
  · 사진은 전부 B열에 **한 줄당 1장** twoCellAnchor 로 붙어 있다.
  · 구매품 60줄 **전부** 사진이 있었다(100%). 장당 평균 5.3KB.
  · openpyxl 3.1.5 는 사진을 읽고(ws._images) 저장까지 지킨다.
  🔴 ws.insert_rows() 는 **사진을 밀지 않는다** — 병합 칸과 같은 함정.
     줄을 넣는 쪽은 반드시 shift_photos() 를 같이 불러야 한다.
"""
import io

from openpyxl.drawing.image import Image as _XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker as _Marker
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor as _TwoCell
from openpyxl.utils import get_column_letter as _gl

# 사진 칸 기본 크기 — 대표 확정 "아주 작게".
# 통일판 자료줄 기본 높이가 24.95pt 이므로, 그 안에 들어가면 **표가 하나도 길어지지 않는다.**
PHOTO_PT = 24.0                   # 줄 높이(pt) 최소값 = 사진이 차지할 세로
_PT_PER_PX = 0.75                 # 96dpi 기준 1px = 0.75pt
_CH_PER_PX = 7.0                  # 엑셀 열 너비 1 ≈ 7px

HEADER_TEXT = "PHOTO" + chr(10) + "사진"


def row_height_for(pt=None):
    """사진 줄의 최소 높이(pt). 원래 줄이 이미 더 높으면 그대로 둔다."""
    return float(pt or PHOTO_PT)


def col_width_for(pt=None):
    """사진 열 너비 — 칸이 **정사각형**이 되게 줄 높이에 맞춘다.

    네모난 칸이라야 사진이 찌그러지지 않고, 사용자가 칸을 키울 때도 모양이 유지된다.
    """
    px = float(pt or PHOTO_PT) / _PT_PER_PX
    return round(px / _CH_PER_PX, 2)


def _to_bytes(img):
    """openpyxl 이 읽은 그림 → 원본 바이트. 스트림이면 되감아 읽는다."""
    ref = getattr(img, "ref", None)
    if ref is None:
        return None
    if isinstance(ref, (bytes, bytearray)):
        return bytes(ref)
    if hasattr(ref, "read"):
        try:
            ref.seek(0)
        except Exception:
            pass
        return ref.read()
    try:
        with open(str(ref), "rb") as fh:
            return fh.read()
    except Exception:
        return None


def read_photos(ws):
    """{엑셀 줄번호(1부터): 그림 바이트} 와 못 읽은 장수.

    ⛔ 조용히 건너뛰지 않는다 — 못 읽은 수를 함께 돌려준다.
    한 줄에 여러 장이 붙어 있으면 첫 장만 쓴다(인벤터는 한 줄당 1장).
    """
    got, bad = {}, 0
    for img in list(getattr(ws, "_images", []) or []):
        frm = getattr(getattr(img, "anchor", None), "_from", None)
        if frm is None:
            bad += 1
            continue
        row = int(frm.row) + 1
        if row in got:
            continue
        data = _to_bytes(img)
        if not data:
            bad += 1
            continue
        got[row] = data
    return got, bad


def put_photo(ws, row, col, data, pt=None):
    """그 칸에 사진을 **칸에 매어** 넣는다. 넣었으면 True.

    ⭐ 칸 한 개(row,col)~(row+1,col+1) 에 매므로 사용자가 줄 높이나 열 너비를
       늘리면 사진도 그만큼 커진다 (대표 지시 2026-09-10).
    """
    if not data:
        return False
    try:
        im = _XLImage(io.BytesIO(data))
    except Exception:
        return False

    # 칸 한 개에 매단다 — 고정 크기(OneCellAnchor)로 두면 칸을 키워도 안 커진다.
    im.anchor = _TwoCell(
        editAs="twoCell",
        _from=_Marker(col=col - 1, colOff=0, row=row - 1, rowOff=0),
        to=_Marker(col=col, colOff=0, row=row, rowOff=0),
    )
    ws.add_image(im)

    want_h = row_height_for(pt)
    cur_h = ws.row_dimensions[row].height
    if cur_h is None or cur_h < want_h:
        ws.row_dimensions[row].height = want_h

    letter = _gl(col)
    want_w = col_width_for(pt)
    cur_w = ws.column_dimensions[letter].width
    if cur_w is None or cur_w < want_w:
        ws.column_dimensions[letter].width = want_w
    return True


def shift_photos(ws, at_row, n):
    """at_row(1부터) 이상에 붙은 사진을 n줄 아래로 민다.

    🔴 openpyxl 의 insert_rows 는 사진을 옮겨 주지 않는다. 이걸 안 부르면
       줄을 넣은 뒤 **사진만 제자리에 남아 엉뚱한 부품에 붙는다.**
    """
    if n <= 0:
        return 0
    moved = 0
    for img in list(getattr(ws, "_images", []) or []):
        a = getattr(img, "anchor", None)
        frm = getattr(a, "_from", None)
        if frm is None or (int(frm.row) + 1) < at_row:
            continue
        frm.row = int(frm.row) + n
        to = getattr(a, "to", None)            # 칸에 매인 사진은 끝점도 같이
        if to is not None and hasattr(to, "row"):
            to.row = int(to.row) + n
        moved += 1
    return moved


def copy_photos(src_ws, dst_ws, row_map, col, pt=None):
    """원본 줄 -> 대상 줄 대응표대로 사진을 옮긴다. (옮긴 수, 못 읽은 수)"""
    photos, bad = read_photos(src_ws)
    done = 0
    for src_row, dst_row in row_map.items():
        data = photos.get(src_row)
        if data and put_photo(dst_ws, dst_row, col, data, pt):
            done += 1
    return done, bad


def write_header(ws, row, col, ref_col=None, pt=None):
    """사진 열 머리글을 옆 칸과 같은 모양으로 적는다."""
    from copy import copy as _copy
    c = ws.cell(row=row, column=col)
    if ref_col:
        try:
            c._style = _copy(ws.cell(row=row, column=ref_col)._style)
        except Exception:
            pass
    c.value = HEADER_TEXT
    ws.column_dimensions[_gl(col)].width = col_width_for(pt)
    return c
