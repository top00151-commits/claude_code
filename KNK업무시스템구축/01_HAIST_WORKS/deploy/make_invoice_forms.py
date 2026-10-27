# -*- coding: utf-8 -*-
"""z1191m — 베트남 인보이스 「표준 양식」 만들기 (대표 지시 2026-10-27
   「본사에서 올릴 인보이스 표준양식(수책,소모품)을 여기서 받을수있게해줘.. 기본적은 공통 정보는 사전에 작성된 상태로..」)

대표가 준 최종 양식(완성본)에서 **인보이스마다 바뀌는 칸만** 비운 빈 양식을 만든다 → app/seed/invoice_form_*.xlsx
  남김: 머리 공통 정보(보내는 곳·받는 곳·담당·T/T·SAME AS ABOVE·INCOTERMS·인천→하노이) · 칸 제목 ·
        금액 수식(수량×단가) · 합계 줄 · 서명·도장 그림 · 줄 번호(소모품 C열) · 서식·병합·인쇄 설정·줄 수
  비움: 인보이스 번호 · 날짜 · 선적일 · 인보이스·패킹리스트 품목 줄 · 품목 사진
  덤:   지난 인보이스 글자가 엑셀 안 글자 목록(sharedStrings)에 숨어 남지 않게 쓰는 글자만 남긴다 ·
        계산 순서표(calcChain)는 빼서 엑셀이 다시 만들게 한다 · 열 때 다시 계산
  ⚠공통·가변 구분 근거 = 대표 양식 2개 + 이음 방 본사 인보이스 11개 칸 대조(10-27 실측 · 13개 모두 같은 칸만 공통)
  ⚠파일을 통째로 다시 저장하지 않는다(XML 칸만) — 그림·서식을 잃지 않게(app/hs_fill.py 와 같은 방식)

쓰는 법:  python deploy/make_invoice_forms.py <수책 완성본.xlsx> <소모품 완성본.xlsx> [출력 폴더=app/seed]
"""
import io
import os
import re
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from app import hs_fill as H  # noqa: E402

TOTAL = re.compile(r"^\s*TOTAL\s*$", re.I)


def _find_rows(data: bytes) -> dict:
    """시트마다 머리글 줄·첫 품목 줄·합계 줄 + 인보이스 번호·날짜·선적일 칸 주소"""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    out = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 400, 400), max_col=22, values_only=True))
        head = next((i + 1 for i, r in enumerate(rows) if any(isinstance(v, str) and "Description" in v for v in r)), None)
        if not head:
            continue
        tot = next((i + 1 for i, r in enumerate(rows) if i + 1 > head and r and isinstance(r[0], str) and TOTAL.match(r[0])), None)
        first = next((i + 1 for i, r in enumerate(rows) if i + 1 > head and len(r) > 3 and r[3] not in (None, "", " ")), None)
        info = {"head": head, "first": first, "total": tot, "clear": []}
        for i, r in enumerate(rows[:30]):
            for j, v in enumerate(r):
                if not isinstance(v, str):
                    continue
                if H.INVOICE_NO.search(v):
                    info["clear"].append((i + 1, j + 1))                       # 인보이스 번호
                if v.strip().upper().startswith("DATE") and j + 1 < len(r) and not str(r[j + 1] or "").startswith("="):
                    info["clear"].append((i + 1, j + 2))                       # 날짜(라벨 오른쪽 · 수식이면 안 비움)
                if "Sailing On or About" in v and i + 1 < len(rows) and not str(rows[i + 1][j] or "").startswith("="):
                    info["clear"].append((i + 2, j + 1))                       # 선적일(라벨 아래 · 수식이면 안 비움)
        out[ws.title] = info
    wb.close()
    return out


def make(data: bytes, keep_number_col=None) -> bytes:
    z = zipfile.ZipFile(io.BytesIO(data))
    rows = _find_rows(data)
    paths = {name: path for name, _, path in H._sheet_paths(z)}
    changed = {}
    for name, info in rows.items():
        p = paths[name]
        x = z.read(p).decode("utf-8")
        is_inv = not re.search(r"PACK", name, re.I)
        lo, hi = info["first"], (info["total"] or info["first"]) - 1

        def fix_cell(m, r):
            cell = m.group(0)
            ref = re.search(r'\br="([A-Z]+)(\d+)"', cell)
            col = ref.group(1)
            s = re.search(r'\bs="(\d+)"', cell[:cell.index(">") + 1])
            st = (' s="%s"' % s.group(1)) if s else ""
            f = re.search(r"<f\b[^>]*?(?:/>|>.*?</f>)", cell, re.S)
            if is_inv and f and (re.search(r"[A-Z]+%d\b" % r, f.group(0)) or 't="shared"' in f.group(0)):
                return '<c r="%s%d"%s>%s</c>' % (col, r, st, f.group(0))                 # 금액 수식은 남김(값은 다시 계산)
            if is_inv and keep_number_col and col == keep_number_col and not f:
                return cell                                                              # 줄 번호
            return '<c r="%s%d"%s/>' % (col, r, st)

        def fix_row(rm):
            r = int(rm.group(1))
            body = rm.group(2)
            if body is None:
                return rm.group(0)
            in_items = lo and lo <= r <= hi
            targets = {c for rr, c in info["clear"] if rr == r}
            if not in_items and not targets:
                return rm.group(0)
            def one(cm):
                if in_items:
                    return fix_cell(cm, r)
                ref = re.search(r'\br="([A-Z]+)\d+"', cm.group(0)).group(1)
                if H._col_num(ref) in targets:
                    st = re.search(r'\bs="(\d+)"', cm.group(0)[:cm.group(0).index(">") + 1])
                    return '<c r="%s%d"%s/>' % (ref, r, (' s="%s"' % st.group(1)) if st else "")
                return cm.group(0)
            body2 = re.sub(r"<c\b[^>]*?(?:/>|>.*?</c>)", one, body, flags=re.S)
            head = rm.group(0)[:rm.group(0).index(">") + 1]
            return head + body2 + "</row>"

        x2 = re.sub(r'<row\b[^>]*?\br="(\d+)"[^>]*?(?:/>|>(.*?)</row>)', fix_row, x, flags=re.S)
        changed[p] = x2
        # 품목 줄에 붙은 사진(그림판 앵커) 빼기 — 머리·서명 쪽 그림은 둔다
        rel = p.replace("worksheets/", "worksheets/_rels/") + ".rels"
        if rel in z.namelist() and lo:
            rx = z.read(rel).decode("utf-8")
            dm = re.search(r'Target="\.\./drawings/(drawing\d+\.xml)"', rx)
            if dm:
                dp = "xl/drawings/" + dm.group(1)
                dx = z.read(dp).decode("utf-8")
                gone = []

                def drop(am):
                    a = am.group(0)
                    fr = re.search(r"<xdr:from>.*?<xdr:row>(\d+)</xdr:row>", a, re.S)
                    if fr and "<xdr:pic" in a and lo - 1 <= int(fr.group(1)) <= hi - 1:
                        gone.extend(re.findall(r'r:embed="([^"]+)"', a))
                        return ""
                    return a
                dx2 = re.sub(r"<xdr:(twoCellAnchor|oneCellAnchor)\b.*?</xdr:\1>", drop, dx, flags=re.S)
                if gone:
                    changed[dp] = dx2
                    drel = "xl/drawings/_rels/" + dm.group(1) + ".rels"
                    drx = z.read(drel).decode("utf-8")
                    still = set(re.findall(r'r:embed="([^"]+)"', dx2))
                    drx2 = re.sub(r'<Relationship\b[^>]*\bId="([^"]+)"[^>]*/>',
                                  lambda mm: "" if (mm.group(1) in gone and mm.group(1) not in still) else mm.group(0), drx)
                    changed[drel] = drx2
    # 쓰는 글자만 남기기(sharedStrings 다시 번호 매김)
    if "xl/sharedStrings.xml" in z.namelist():
        ss = z.read("xl/sharedStrings.xml").decode("utf-8")
        sis = re.findall(r"<si/>|<si\b.*?</si>", ss, re.S)   # 빈 <si/> 를 먼저 — 뒤 항목과 붙어 잡히지 않게
        um = re.search(r'\buniqueCount="(\d+)"', ss)
        if um and int(um.group(1)) != len(sis):
            raise SystemExit("글자 목록 개수가 맞지 않음: %s ≠ %d" % (um.group(1), len(sis)))
        sheets = {p: changed.get(p) or z.read(p).decode("utf-8") for p in paths.values()}
        used, total = [], 0
        for p, x in sheets.items():
            for m in re.finditer(r'<c\b[^>]*\bt="s"[^>]*>\s*<v>(\d+)</v>', x):
                i = int(m.group(1))
                total += 1
                if i not in used:
                    used.append(i)
        used.sort()
        remap = {old: new for new, old in enumerate(used)}
        for p, x in sheets.items():
            x2 = re.sub(r'(<c\b[^>]*\bt="s"[^>]*>\s*<v>)(\d+)(</v>)', lambda m: m.group(1) + str(remap[int(m.group(2))]) + m.group(3), x)
            changed[p] = x2
        i0 = ss.index("<sst")
        i1 = ss.index(">", i0) + 1
        head = ss[i0:i1]                                   # <sst …> 여는 태그만(앞 XML 선언은 pre 에)
        head = re.sub(r'\bcount="\d+"', 'count="%d"' % total, head)
        head = re.sub(r'\buniqueCount="\d+"', 'uniqueCount="%d"' % len(used), head)
        changed["xl/sharedStrings.xml"] = ss[:i0] + head + "".join(sis[i] for i in used) + "</sst>"
    # 계산 순서표 빼기 · 열 때 다시 계산
    drop_parts = set()
    if "xl/calcChain.xml" in z.namelist():
        drop_parts.add("xl/calcChain.xml")
        wr = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        changed["xl/_rels/workbook.xml.rels"] = re.sub(r'<Relationship\b[^>]*calcChain[^>]*/>', "", wr)
        ct = z.read("[Content_Types].xml").decode("utf-8")
        changed["[Content_Types].xml"] = re.sub(r'<Override\b[^>]*calcChain[^>]*/>', "", ct)
    wbx = z.read("xl/workbook.xml").decode("utf-8")
    cp = re.search(r"<calcPr\b[^>]*?/>", wbx)
    if cp and "fullCalcOnLoad" not in cp.group(0):
        changed["xl/workbook.xml"] = wbx.replace(cp.group(0), cp.group(0)[:-2] + ' fullCalcOnLoad="1"/>')
    # 더는 아무도 안 가리키는 그림 파일 빼기
    rel_text = "".join((changed.get(n) or z.read(n).decode("utf-8", "replace")) for n in z.namelist() if n.endswith(".rels"))
    for n in z.namelist():
        if n.startswith("xl/media/") and n.rsplit("/", 1)[-1] not in rel_text:
            drop_parts.add(n)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zo:
        for info in z.infolist():
            if info.filename in drop_parts:
                continue
            buf = changed.get(info.filename)
            buf = buf.encode("utf-8") if isinstance(buf, str) else (buf if buf is not None else z.read(info.filename))
            zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            zo.writestr(zi, buf)
    return out.getvalue()


if __name__ == "__main__":
    src_s, src_m = sys.argv[1], sys.argv[2]
    dst = sys.argv[3] if len(sys.argv) > 3 else os.path.join(os.path.dirname(HERE), "app", "seed")
    for src, name, numcol in ((src_s, "invoice_form_suchaek.xlsx", None), (src_m, "invoice_form_somo.xlsx", "C")):
        b = make(open(src, "rb").read(), keep_number_col=numcol)
        open(os.path.join(dst, name), "wb").write(b)
        print("만듦", name, len(b), "바이트")
