# -*- coding: utf-8 -*-
"""z1191u 시험 — 옛 인보이스 양식 → 최종 양식(수책통관 · 일반통관)으로 바꿔 채우기
   대표 2026-10-28 「본사에서 업로드하는 양식, 베트남에서 업로드하는 양식을 지금 첨부한 파일로 반영해서 만들어줘..
   기존에 올리는 양식들이 이전꺼라 조금씩 다른데.. 확인해서 작성한 내용들이 포함된 문서는 항상 최종 양식으로 변경되어
   작성이 되도록 적용해줘...」 → 이해 보고 → 「진행해」
  §1 원문: ① 경로가 채우기 전에 바꾼다 · 못 바꾸면 올린 모양 그대로 · 올린 파일 _src 보관 · conv_note · 화면 고르기 · 바꿈 요약
  §2 바꾸기(모듈): 옛 수책 · 옛 일반(사진 · 줄 늘림) · 이미 최종 · 인보이스 두 장 · 칸 모양 다름 · 저장값 새로 셈 · 수식 옮기기
  §3 ① 경로: 바꿔 채움 · 고른 양식 · 못 바꿈 안내 · 본사 확정 뒤에도 최종 양식 · 기록 표
  §4 실물(수출자료 폴더가 있을 때만 — 단가가 들어 있어 git 밖): 58개
   ⚠ 가짜 옛 양식은 openpyxl 로 만든다(수식 저장값 없음 — 엑셀이 아닌 도구로 저장한 파일도 금액 대조가 되는지 함께 본다)
"""
import ast
import contextlib
import datetime as dt
import glob
import io
import json
import os
import re
import sys
import tempfile
import zipfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
os.environ["KNK_SSO_SERVICE_KEY"] = "z1191u-test-only-not-a-real-key"   # 시험용 가짜 값

import openpyxl                                            # noqa: E402
from openpyxl.drawing.image import Image as XLImage        # noqa: E402

import app.database as D                                   # noqa: E402
from app import hs_fill as H                               # noqa: E402
from app import hs_form_convert as C                       # noqa: E402

OK = FAIL = 0


def check(label, cond, got=None):
    global OK, FAIL
    if cond:
        OK += 1
        print("  [OK] " + label)
    else:
        FAIL += 1
        print("  [FAIL] " + label + ("   ← %r" % (got,) if got is not None else ""))


def quiet(fn, *a, **k):
    with contextlib.redirect_stdout(io.StringIO()):
        return fn(*a, **k)


def serial(d):
    return (d - dt.datetime(1899, 12, 30)).days


# ─────────────────────────────────────────────────────────────
# 가짜 옛 양식
# ─────────────────────────────────────────────────────────────
INV_DATE, SAIL = dt.datetime(2026, 10, 28), dt.datetime(2026, 10, 30)


def head(ws, no, incoterms):
    """옛 양식 머리 — 받는 곳 주소 PROVINCE · TAX ID 없음(최종 양식은 CITY · TAX ID)"""
    for ref, v in (("A1", " COMMERCIAL  INVOICE"), ("A2", "1) Shipper / Exporter"), ("F2", "8) No. & Date of Invoice"),
                   ("B4", "KNK CO.,LTD"), ("G4", no), ("H4", "DATE:"), ("I4", INV_DATE),
                   ("B5", "2, Gisan-ro 87 beon-gil, Byeongjeom-gu, "), ("B6", "Gyeonggi-do, Republic of Korea."),
                   ("F7", "9) No. & Date of L/C"), ("B8", "ATTN : Ji-Yeon, An (+82-10-2739-2660)"),
                   ("A10", "2) For Account & Risk of Messrs"), ("F10", "T/T"), ("B12", "KNK VINA Co., Ltd."),
                   ("B13", "LOT CN5B-2,QUE VO III INDUSTRIAL PARK,"), ("B14", "QUE VO WARD, BAC NINH PROVINCE, VIETNAM."),
                   ("F13", "10) Issuing Bank"), ("A19", "3) Notify Party"), ("F19", "11) Remarks"), ("B21", "SAME AS ABOVE"),
                   ("G21", incoterms), ("A24", "4) Port of Loading"), ("D24", "5) Final Destination"), ("A25", "Incheon, Korea"),
                   ("D25", "Hanoi, Vietnam"), ("A28", "6) Carrier"), ("D28", "7) Sailing On or About"), ("D29", SAIL),
                   ("A31", "12) Marks and Number of PKGS"), ("D31", "13) Description of Goods"), ("F31", " 14) Q'ty /Unit"),
                   ("H31", "15) Unit Price"), ("I31", "16) Amount")):
        ws[ref] = v


def pl_sheet(wb, items, box=None):
    """패킹리스트 — 머리글 35 · 품목 37~ · 중량·부피는 첫 줄에만(최종 양식은 세로 합친 칸)"""
    pl = wb.create_sheet("PACKING LIST")
    for ref, v in (("A1", "PACKING LIST"), ("A2", "1) Shipper / Exporter"), ("F2", "8) No. & Date of Invoice"),
                   ("G4", "=++INVOICE!G4"), ("I4", "DATE :"), ("A35", "12) Marks and Number of PKGS"),
                   ("D35", "13) Description of Goods"), ("F35", " 14) Q'ty /Unit"), ("H35", "Net -"), ("I35", "Gross-"),
                   ("J35", "Measurement"), ("H36", "Weight"), ("I36", "Weight")):
        pl[ref] = v
    if box:
        pl["L35"], pl["M35"], pl["N35"] = "NO.", "SIZE", "G/W"
    for i, (desc, qty) in enumerate(items):
        r = 37 + i
        pl.cell(r, 4, desc)
        pl.cell(r, 6, qty)
        pl.cell(r, 7, "EA")
    pl["H37"], pl["I37"], pl["J37"] = 20, 25, 0.5
    if box:
        pl["L37"], pl["M37"], pl["N37"] = box
    t = 37 + len(items) + 1
    pl.cell(t, 1, "TOTAL")
    pl.cell(t, 6, "=SUM(F37:G%d)" % (t - 1))
    for c in "HIJ":
        pl["%s%d" % (c, t)] = "=SUM(%s37:%s%d)" % (c, c, t - 1)
    return pl


SU_ITEMS = [  # (품명, 수량, 단가, 금액(None=수식 F*H), HS, Tên hàng)
    ("VALVE SY5120-5DZ", 3, 12.5, None, "11111111", "Van điện từ"),
    ("SENSOR E2E-X5ME1", 2, 30, None, None, None),
    ("BEARING 6204ZZ", 10, 1.2, 12, None, None),
    ("CYLINDER CDQ2B32-20D", 1, 85, None, None, None),
    ("CABLE 3M", 5, 2, None, None, None),
]


def old_suchaek():
    """옛 수책 — 관리코드(B) · 장비 묶음(K24) · ORIGIN·SUPPLIER·CATEGORY·HS·Tên hàng·발송일·Duty·VAT · J 「장비명(N대)」"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INVOICE"
    head(ws, "KNK-261190-1", "INCOTERMS : F.O.B")
    for ref, v in (("K24", "관리코드"), ("L24", "MODEL"), ("M24", "품명"), ("N24", "수량"),
                   ("K25", "006M2601"), ("L25", "CTC-200"), ("M25", "CTC 부착기"), ("N25", 2)):
        ws[ref] = v
    for c, t in zip("KLMNOPQR", ("ORIGIN", "SUPPLIER", "CATEGORY", "HS CODE", "Tên hàng", "발송일", "Duty", "VAT")):
        ws["%s31" % c] = t
    for i, (desc, q, p, a, hs, vn) in enumerate(SU_ITEMS):
        r = 34 + i
        ws.cell(r, 2, "006M2601")
        ws.cell(r, 3, i + 1)
        ws.cell(r, 4, desc)
        ws.cell(r, 6, q)
        ws.cell(r, 7, "EA")
        ws.cell(r, 8, p)
        ws.cell(r, 9, a if a is not None else "=F%d*H%d" % (r, r))
        ws.cell(r, 10, "CTC 부착기(2대)")
        ws.cell(r, 11, "JAPAN")
        ws.cell(r, 12, "SMC")
        ws.cell(r, 13, "공압")
        if hs:
            ws.cell(r, 14, hs)
        if vn:
            ws.cell(r, 15, vn)
        ws.cell(r, 16, SAIL)
        ws.cell(r, 17, 0.05)
        ws.cell(r, 18, 0.1)
    t = 34 + len(SU_ITEMS) + 1
    ws.cell(t, 1, "TOTAL")
    ws.cell(t, 6, "=++SUM(F33:G%d)" % (t - 1))
    ws.cell(t, 9, "=SUM(I33:I%d)" % (t - 1))
    pl_sheet(wb, [(d, q) for d, q, *_ in SU_ITEMS], box=(1, "1100*1100*1000", 25))
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def png(color):
    from PIL import Image
    im = Image.new("RGB", (40, 30), color)
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


N_SOMO = 60


def old_somo():
    """옛 일반 — J 회사(제목 없음) · Origin · Manufacturer · invoice % · C/O · SIZE · HS · Tên hàng · 품목 사진 2장 · 60줄(틀 52줄보다 많음)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INVOICE"
    head(ws, "KNK-261190-2", "INCOTERMS : D.D.U (TEST)")
    for c, t in zip("LMNOPQR", ("Origin", "Manufacturer", "invoice %", "C/O", "SIZE", "HS CODE", "Tên hàng")):
        ws["%s31" % c] = t
    for i in range(N_SOMO):
        r = 34 + i
        ws.cell(r, 3, i + 1)
        ws.cell(r, 4, "PART-%03d TYPE A" % (i + 1))
        ws.cell(r, 6, i % 7 + 1)
        ws.cell(r, 7, "EA")
        ws.cell(r, 8, 1.5 + i)
        ws.cell(r, 9, "=F%d*H%d" % (r, r))
        ws.cell(r, 10, "DSV")
        ws.cell(r, 12, "KOREA")
        ws.cell(r, 13, "KNK")
        ws.cell(r, 14, 1)
        ws.cell(r, 15, "Y")
        ws.cell(r, 16, "10*20")
    for r, col in ((34, "FF0000"), (36, "0000FF")):
        img = XLImage(io.BytesIO(png("#" + col)))
        ws.add_image(img, "K%d" % r)
    t = 34 + N_SOMO + 1
    ws.cell(t, 1, "TOTAL")
    ws.cell(t, 9, "=SUM(I33:I%d)" % (t - 1))
    pl_sheet(wb, [("PART-%03d TYPE A" % (i + 1), i % 7 + 1) for i in range(N_SOMO)])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def old_two_invoices():
    """인보이스 두 장(260121 처럼 시트마다 번호가 다름)"""
    wb = openpyxl.Workbook()
    for k, name in enumerate(("001M2601", "002M2601")):
        ws = wb.active if k == 0 else wb.create_sheet()
        ws.title = name
        head(ws, "KNK-261190-%d" % (k + 3), "INCOTERMS : C.I.P")
        for c, t in zip("KLMNO", ("ORIGIN", "SUPPLIER", "CATEGORY", "HS CODE", "Tên hàng")):
            ws["%s31" % c] = t
        for i in range(2):
            r = 34 + i
            ws.cell(r, 4, "PART %s-%d" % (name, i))
            ws.cell(r, 6, 1)
            ws.cell(r, 8, 10)
            ws.cell(r, 9, 10)
        ws.cell(37, 1, "TOTAL")
    pl_sheet(wb, [("PART %s-%d" % (n, i), 1) for n in ("001M2601", "002M2601") for i in range(2)])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def plain_table():
    """KNK 인보이스 모양이 아닌 표 — 머리글만 맞다(z1191o 시험의 「예전 양식」과 같은 것)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INVOICE"
    ws.append(["No", "Description", "Q'ty", "HS CODE", "Tên hàng"])
    ws.append([1, "DCP32DN 15A", 2, None, None])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


SU, SO, TWO, PLAIN = old_suchaek(), old_somo(), old_two_invoices(), plain_table()
SEED = os.path.join(ROOT, "app", "seed")

print("=" * 78)
print("§1 원문")
print("=" * 78)
SRC = open(os.path.join(ROOT, "app", "main.py"), encoding="utf-8").read()
FN = {n.name: ast.get_source_segment(SRC, n) for n in ast.parse(SRC).body
      if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))
      and n.name in ("export_hs_fill_run", "_hs_fill_log_ensure", "_hs_fill_payload", "_hs_fill_confirm_do", "_hs_conv_note",
                     "export_hs_fill_page")}
T = open(os.path.join(ROOT, "app", "templates", "export_hs_fill.html"), encoding="utf-8").read()
run = FN["export_hs_fill_run"]
check("1 ① 경로 — 채우기 전에 바꾼다(convert → 바꾼 파일로 fill) · 못 바꾸거나 품목 수가 다르면 올린 파일(raw)로 fill",
      "_hfc.convert, raw, want" in run and run.index("_hfc.convert, raw, want") < run.index('_hf.fill, conv["out"], entries')
      < run.index("_hf.fill, raw, entries") and "n_fill != n_src" in run)
check("2 바꿨으면 바꾼 파일을 원본 자리(_in) · 올린 파일은 _src 로(30일 정리는 .xlsx 전부) · 기록에 conv_note",
      "fo.write(data)" in run and 'tok + "_src.xlsx"' in run and "filled_by,conv_note" in run
      and '"conv_note"' in FN["_hs_fill_log_ensure"])
check("3 바꿀 양식 = 화면 고름(suchaek → 수책 · somo → 소모품 · 그 밖은 자동)",
      'want = {"suchaek": "수책", "somo": "소모품"}.get(str(form.get("form_kind") or ""), "auto")' in run)
check("4 못 바꾸면 안내 — 「최종 양식으로 바꾸지 못해 올린 모양 그대로 채웠습니다(까닭)」 + 옛 안내(최종 양식으로 만들어 주세요)",
      "최종 양식으로 바꾸지 못해 올린 모양 그대로 채웠습니다(%s)" in FN["_hs_fill_payload"]
      and "앞으로는 최종 양식(수책통관·일반통관)으로 만들어 주세요" in FN["_hs_fill_payload"])
check("5 본사 확정(다시 만들기) 뒤에도 바꿈 요약 — 기록의 conv_note 를 결과에 · 다시 만들기는 _in(바꾼 파일)에서",
      'extra["conv"] = cvn' in FN["_hs_fill_confirm_do"] and 'lg.get("orig_stored")' in FN["_hs_fill_confirm_do"])
check("6 화면 — 고르기 3개(자동 · 수책통관 · 일반통관) · 보낼 때 form_kind · 결과 카드 바꿈 요약 · 기록 표 표시",
      T.count('type="radio" name="hfFormKind"') == 3 and "fd.append('form_kind', fk ? fk.value : 'auto')" in T
      and "function convBlock(c)" in T and "h += convBlock(j.conv);" in T and "🔁 옛 양식에서 바꿈" in T
      and "옛 양식 그대로(못 바꿈)" in T and 'lg["conv"] = _json.loads(lg.get("conv_note") or "null")' in FN["export_hs_fill_page"])

print()
print("=" * 78)
print("§2 바꾸기(모듈)")
print("=" * 78)
r1 = C.convert(SU)
check("7 옛 수책 → 바꿈 · 맞음 · 자동 수책(관리코드 · 장비 묶음 · SUPPLIER·CATEGORY)",
      r1.get("converted") and r1.get("ok") and r1.get("kind") == "수책" and r1.get("auto") and "수책통관 표시" in r1.get("why", ""),
      {k: r1.get(k) for k in ("converted", "ok", "kind", "why", "bad", "error")})
out1 = r1.get("out") or b""
bk1 = C._book(out1)
T1 = C._template("수책")
inv1, pl1 = T1["sheets"]["INV"], T1["sheets"]["PL"]
c1 = bk1["cells"](inv1["path"])
p1 = bk1["cells"](pl1["path"])
check("8 새 파일 = 수책통관 최종 양식(머리글 칸이 틀과 같음) · 품목 5줄을 그대로 읽음(scan = scan_fast)",
      C.layout_kind(out1) == "수책" and [it["desc"] for it in H.scan_fast(out1)["sheets"][0]["items"]] == [x[0] for x in SU_ITEMS]
      and H.scan(out1)["sheets"][0]["items"] == H.scan_fast(out1)["sheets"][0]["items"])
row34 = {H._col_letter(c): C._v(c1, 34, c) for c in range(2, 22) if C._v(c1, 34, c) is not None}
check("9 줄 칸 옮김 — 관리코드(L) · 장비 MODEL·품명·Q'TY(M·N·O · 장비 묶음 값) · ORIGIN(P) · SUPPLIER(Q) · CATEGORY(R) · HS(S) · Tên hàng(T) · 발송일(U)",
      row34.get("D") == "VALVE SY5120-5DZ" and row34.get("F") == 3 and row34.get("H") == 12.5 and row34.get("L") == "006M2601"
      and row34.get("M") == "CTC-200" and row34.get("N") == "CTC 부착기" and row34.get("O") == 2 and row34.get("P") == "JAPAN"
      and row34.get("Q") == "SMC" and row34.get("R") == "공압" and str(row34.get("S")) == "11111111"
      and row34.get("T") == "Van điện từ" and row34.get("U") == serial(SAIL), row34)
dr1 = r1["report"]["dropped"]
check("10 옛 양식에만 있는 Duty · VAT 는 뺀다(줄마다 5개씩 · 까닭) · 금액 수식 줄은 틀 수식 · 값 줄은 그 값",
      dr1.get("Duty", {}).get("n") == 5 and dr1.get("VAT", {}).get("n") == 5 and "옛 양식에만" in dr1["Duty"]["why"]
      and C._f(c1, 34, 9) == "F34*H34" and C._v(c1, 36, 9) == 12 and not C._f(c1, 36, 9), (list(dr1), C._f(c1, 34, 9), c1.get((36, 9))))
check("11 머리 — 번호·날짜·선적일은 원본 · INCOTERMS(F.O.B)는 원본대로(틀 C.I.P 아님) · 받는 곳 주소는 최종 양식(CITY · TAX ID)",
      C._v(c1, 4, 7) == "KNK-261190-1" and C._v(c1, 4, 9) == serial(INV_DATE) and C._v(c1, 29, 4) == serial(SAIL)
      and C._v(c1, 21, 7) == "INCOTERMS : F.O.B" and C._v(c1, 14, 2) == "QUE VO WARD, BAC NINH CITY, VIETNAM."
      and "TAX ID" in str(C._v(c1, 12, 4))
      and any(m["block"].startswith("11) Remarks") for m in r1["report"]["header_moved"])
      and any(d["block"].startswith("2) For Account") for d in r1["header_diffs"]),
      (C._v(c1, 21, 7), C._v(c1, 14, 2), r1["report"]["header_moved"]))
tot = inv1["total"]
check("12 저장값 새로 셈 — 금액 F×H(37.5) · TOTAL 금액 204.5 · 수량 21 (틀에 남은 대표 원본 261015 합계 아님)",
      C._v(c1, 34, 9) == 37.5 and abs(C._v(c1, tot, 9) - 204.5) < 1e-9 and C._v(c1, tot, 6) == 21,
      (C._v(c1, 34, 9), C._v(c1, tot, 9), C._v(c1, tot, 6)))
check("13 패킹 — 품목 5줄 · 중량·부피는 세로 합친 칸 위 칸(38줄) · 박스 표(NO.·SIZE·G/W→WEIGHT) · 머리 번호 저장값 = 이 인보이스(KNK-261015-1 아님)",
      [C._v(p1, 38 + i, 4) for i in range(5)] == [x[0] for x in SU_ITEMS] and C._v(p1, 38, 8) == 20 and C._v(p1, 38, 9) == 25
      and C._v(p1, 38, 10) == 0.5 and C._v(p1, 38, 12) == 1 and C._v(p1, 38, 13) == "1100*1100*1000" and C._v(p1, 38, 14) == 25
      and C._v(p1, 4, 7) == "KNK-261190-1" and C._f(p1, 4, 7) == "++INVOICE!G4", (C._v(p1, 4, 7), [C._v(p1, 38, c) for c in (8, 9, 10, 12, 13, 14)]))
wbx = zipfile.ZipFile(io.BytesIO(out1)).read("xl/workbook.xml").decode("utf-8")
check("14 엑셀은 열 때 다시 셈(fullCalcOnLoad) · 저장값 못 센 수식 0",
      re.search(r'<calcPr\b[^>]*fullCalcOnLoad="1"', wbx) is not None and r1["report"]["cached"]["cleared"] == 0, r1["report"].get("cached"))

r2 = C.convert(SO)
out2 = r2.get("out") or b""
T2 = C._template("소모품")
inv2 = T2["sheets"]["INV"]
c2 = C._book(out2)["cells"](inv2["path"]) if out2 else {}
check("15 옛 일반 → 바꿈 · 맞음 · 자동 일반(invoice % · C/O · Manufacturer · 품목 사진) · 60줄 > 틀 52줄 → 8줄 늘림(패킹도)",
      r2.get("converted") and r2.get("ok") and r2.get("kind") == "소모품" and r2.get("auto")
      and r2["report"]["expanded"] == N_SOMO - (inv2["last"] - inv2["first"] + 1) == 8 and r2["report"]["pl_expanded"] == 8,
      {k: r2.get(k) for k in ("converted", "ok", "kind", "why", "bad", "error")})
last2 = inv2["first"] + N_SOMO - 1
tf = C._f(c2, inv2["total"] + 8, 9) or ""
check("16 늘린 줄 — 합계 수식이 마지막 품목 줄까지 · 늘린 줄 금액 수식 · 번호(C) · 일반 J 회사 → COMPANY(K)",
      re.search(r"SUM\(I33:I(\d+)\)", tf) and int(re.search(r"SUM\(I33:I(\d+)\)", tf).group(1)) >= last2
      and C._f(c2, last2, 9) in ("F%d*H%d" % (last2, last2), "++F%d*H%d" % (last2, last2)) and C._v(c2, last2, 3) == N_SOMO
      and C._v(c2, last2, 11) == "DSV", (tf, C._f(c2, last2, 9), C._v(c2, last2, 3), C._v(c2, last2, 11)))
imgs = H.row_images(out2, H.scan_fast(out2)) if out2 else {}
check("17 품목 사진 2장 → IMAGE 칸 · 같은 줄(34 · 36)",
      r2.get("ok") and r2["report"]["photos"] == 2 and sorted(row for (_t, row) in imgs) == [34, 36], (r2.get("report", {}).get("photos"), sorted(imgs)))


def anchor_rows(data, sheet_path):
    z = zipfile.ZipFile(io.BytesIO(data))
    dp = next((t for t in H._rels(z, sheet_path).values() if "/drawings/" in t and t.endswith(".xml")), None)
    return sorted(int(x) for x in re.findall(r"<xdr:from><xdr:col>\d+</xdr:col><xdr:colOff>-?\d+</xdr:colOff><xdr:row>(\d+)</xdr:row>",
                                              z.read(dp).decode("utf-8"))) if dp else []


tpl_rows = [r for r in anchor_rows(T2["bytes"], inv2["path"]) if r >= inv2["last"]]
new_rows = anchor_rows(out2, inv2["path"])
check("18 서명·도장 그림은 늘린 만큼(8줄) 아래로",
      tpl_rows and all((r + 8) in new_rows for r in tpl_rows), (tpl_rows, new_rows[-6:]))
if r2.get("ok"):
    total2 = sum((i % 7 + 1) * (1.5 + i) for i in range(N_SOMO))
    check("19 저장값 — 일반 TOTAL 금액 = 60줄 수량×단가 합(늘린 줄 포함)",
          abs(C._v(c2, inv2["total"] + 8, 9) - total2) < 1e-6, (C._v(c2, inv2["total"] + 8, 9), total2))
else:
    check("19 저장값 — 일반 TOTAL 금액", False, r2.get("bad"))
r3 = C.convert(TWO)
check("20 인보이스 두 장(번호 다름) — 바꾸지 않음(합치면 두 장이 한 장) · 까닭에 두 번호",
      not r3.get("converted") and r3.get("multi") and "KNK-261190-3" in r3.get("error", "") and "KNK-261190-4" in r3.get("error", ""), r3)
r4 = C.convert(PLAIN)
check("21 KNK 인보이스 칸 모양이 아닌 표 — 바꾸지 않음(엉뚱한 칸에 들어가지 않게)",
      not r4.get("converted") and r4.get("layout") and "칸 모양" in r4.get("error", ""), r4)
r5 = C.convert(open(os.path.join(SEED, "invoice_form_suchaek.xlsx"), "rb").read())
r6 = C.convert(open(os.path.join(SEED, "invoice_form_somo.xlsx"), "rb").read())
check("22 이미 최종 양식(표준 양식 2개) — 손대지 않음(final)",
      r5.get("final") and r5.get("kind") == "수책" and r6.get("final") and r6.get("kind") == "소모품" and "out" not in r5)
wb_x = openpyxl.load_workbook(os.path.join(SEED, "invoice_form_suchaek.xlsx"))
wb_x["INVOICE"]["V31"] = "비고"                              # 직원이 최종 양식에 칸을 하나 덧붙였다
_bx = io.BytesIO()
wb_x.save(_bx)
r9 = C.convert(_bx.getvalue())
check("22-1 최종 양식에 칸을 덧붙인 파일도 최종 양식 — 다시 바꾸지 않는다(바꾸면 덧붙인 칸이 빠진다)",
      r9.get("final") and r9.get("kind") == "수책" and not r9.get("converted"), r9)
r7 = C.convert(SU, "소모품")
check("23 본사가 고르면 그 양식(옛 수책 → 일반통관) · 자동 아님",
      r7.get("converted") and r7.get("ok") and r7.get("kind") == "소모품" and not r7.get("auto") and r7.get("why") == "본사가 고름",
      {k: r7.get(k) for k in ("ok", "kind", "auto", "why", "bad")})
wb_nd = openpyxl.load_workbook(io.BytesIO(SU))
wb_nd["INVOICE"]["H4"] = None                               # 「DATE:」 글자 없음 — 번호 칸을 글자로 못 찾는다(10-28 회귀 z1191p 에서 찾음)
_b = io.BytesIO()
wb_nd.save(_b)
r8 = C.convert(_b.getvalue())
c8 = C._book(r8["out"])["cells"](inv1["path"]) if r8.get("out") else {}
check("23-1 「DATE:」 글자가 없어도 인보이스 번호는 옮긴다(머리에서 찾은 KNK-…) · 다시 읽어 대조에 번호 같음도 본다",
      r8.get("ok") and C._v(c8, 4, 7) == "KNK-261190-1" and H.scan_fast(r8["out"])["invoice_no"] == "KNK-261190-1"
      and "인보이스 번호: 원본" in open(C.__file__, encoding="utf-8").read(), (r8.get("ok"), r8.get("bad"), C._v(c8, 4, 7)))
check("24 수식 옮기기 — 다른 시트 주소도 엑셀 채우기처럼 · 따옴표 글자 안은 그대로 · $ 는 고정",
      H.translate_formula("INVOICE!D34", 1) == "INVOICE!D35" and H.translate_formula('IF(A1="B2",1,0)', 1) == 'IF(A2="B2",1,0)'
      and H.translate_formula("$A1+A$1+$A$1+A1", 2, 2) == "$A3+C$1+$A$1+C3"
      and H.translate_formula("'My Sheet'!$D$34+D3", 1) == "'My Sheet'!$D$34+D4")
cells = {("S", 1, 1): 2, ("S", 1, 2): 3.5, ("S", 2, 1): "abc", ("P", 4, 7): "KNK-1"}


def get(sh, r, c):
    return cells.get((sh, r, c))


def calc(f):
    try:
        return C._calc_formula(f, get, "S")
    except C._NoCalc:
        return "X"


check("24-1 글자 풀기 — 숫자 표기(openpyxl 은 한글·베트남어를 &#44288; 로 적는다)도 한 번에 · 옛 양식 원본도 scan = scan_fast",
      H._unesc("&#44288;&#47532;&#53076;&#46300;") == "관리코드" and H._unesc("Van &#273;i&#7879;n t&#7915;") == "Van điện từ"
      and H._unesc("&#38;lt;") == "&lt;" and H._unesc("a&lt;b&amp;c") == "a<b&c" and H._unesc("&#x1F4A1;") == "💡"
      and all(H.scan(d)["sheets"] == H.scan_fast(d)["sheets"] for d in (SU, SO)))
check("25 저장값 셈 — SUM · 곱 · 앞 + · 다른 시트 참조 · 빈 칸 0 · 글자 셈·모르는 함수는 못 셈(저장값 지움)",
      calc("A1*B1") == 7.0 and calc("++SUM(A1:B3)") == 5.5 and calc("++P!G4") == "KNK-1" and calc("A3") == 0
      and calc("A2*2") == "X" and calc("IF(A1>1,1,0)") == "X" and calc('A1&"x"') == "X" and calc("A1/A3") == "X")

print()
print("=" * 78)
print("§3 ① 경로")
print("=" * 78)
TMP = tempfile.mkdtemp(prefix="knk_z1191u_")
_REAL = D.DB_PATH
D.DB_PATH = os.path.join(TMP, "t.db")
assert D.DB_PATH != _REAL
with contextlib.redirect_stdout(io.StringIO()):
    D.init_db()
with D.db_session() as conn:
    have = {r[1] for r in conn.execute("PRAGMA table_info(users)")}
    for col, typ in [("employee_no", "TEXT"), ("entity", "TEXT"), ("name_vi", "TEXT"), ("rank", "TEXT"), ("team_id", "INTEGER"),
                     ("password_version", "INTEGER DEFAULT 1")]:
        if col not in have:
            conn.execute("ALTER TABLE users ADD COLUMN %s %s" % (col, typ))
    conn.execute("INSERT INTO users (name, role, employee_no, login_id, password, is_active, password_version, "
                 "can_view_sales, can_use_sales, can_view_logistics, can_use_logistics) VALUES ('가상U901','member','U901','U901','x',1,1,1,1,0,0)")
    U = dict(conn.execute("SELECT * FROM users WHERE employee_no='U901'").fetchone())
import app.main as M                                        # noqa: E402
from fastapi.testclient import TestClient                  # noqa: E402

M._HS_FILL_DIR = os.path.join(TMP, "hs_fill")
M._HS_HIST_IMG_DIR = os.path.join(TMP, "hs_hist_img")
M.get_user = lambda request, _u=U: _u
cl = TestClient(M.app, follow_redirects=False)
with D.db_session() as c:
    M._hs_dict_ensure(c)
    M._hs_hist_ensure(c)
    for t in ("export_hs_dict", "export_hist_line", "export_hist_inv"):
        c.execute("DELETE FROM %s" % t)
    for desc, hs, vn in (("VALVE SY5120-5DZ", "22222222", "Van điện từ"), ("SENSOR E2E-X5ME1", "85365099", "Cảm biến"),
                         ("BEARING 6204ZZ", "84821000", "Vòng bi")):
        c.execute("INSERT INTO export_hs_dict(part_key,model,vn_name,hs_code,status,seen_count,note,updated_at) "
                  "VALUES(?,?,?,?,'CONFIRMED',1,'',datetime('now','localtime'))", (H.norm_key(desc), desc, vn, hs))


def up(data, name, kind=None):
    d = {"form_kind": kind} if kind else {}
    r = quiet(cl.post, "/export/hs-fill", files={"file": (name, data)}, data=d)
    return r.status_code, r.json()


def log_row(name):
    with D.db_session() as c:
        r = c.execute("SELECT * FROM export_hs_fill_log WHERE file_name=? ORDER BY id DESC LIMIT 1", (name,)).fetchone()
        return dict(r) if r else {}


s1, j1 = up(SU, "Invoice&packinglist_VN(261190-1).xlsx")
cv1 = j1.get("conv") or {}
check("26 옛 수책 ① → 바꿈 요약(수책통관 · 자동 · 품목 5 · 원본대로 둔 머리 11) Remarks · 뺀 Duty·VAT) · 카드 양식 = 수책통관",
      s1 == 200 and j1.get("ok") and cv1.get("ok") and cv1.get("kind") == "수책통관" and cv1.get("auto") and cv1.get("items") == 5
      and any(b.startswith("11) Remarks") for b in cv1.get("header_moved", []))
      and {d["label"] for d in cv1.get("dropped", [])} >= {"Duty", "VAT"} and j1.get("forms") == ["수책통관"]
      and "예전 양식" not in (j1.get("hint") or ""), (s1, j1.get("forms"), cv1, j1.get("hint"), j1.get("error")))
lg1 = log_row("Invoice&packinglist_VN(261190-1).xlsx")
tok = (lg1.get("stored") or "x.xlsx")[:-5]
fin = open(os.path.join(M._HS_FILL_DIR, lg1.get("orig_stored") or "none"), "rb").read() if lg1.get("orig_stored") else b""
srcp = os.path.join(M._HS_FILL_DIR, tok + "_src.xlsx")
check("27 보관 — _in = 바꾼 최종 양식 · _src = 올린 파일 그대로 · 기록 conv_note(ok)",
      C.layout_kind(fin) == "수책" and os.path.exists(srcp) and open(srcp, "rb").read() == SU
      and json.loads(lg1.get("conv_note") or "{}").get("ok") is True, (lg1.get("orig_stored"), os.path.exists(srcp), lg1.get("conv_note")))
dl = quiet(cl.get, "/export/hs-fill/%d/download" % lg1.get("id", 0))
sc_dl = H.scan_fast(dl.content) if dl.status_code == 200 else {"sheets": []}
hs_of = {it["desc"]: it.get("hs_raw") for s in sc_dl["sheets"] for it in s["items"]}
check("28 내려받는 파일 = 최종 양식 + 사전 확정 값 채움(SENSOR·BEARING) · 적힌 값(VALVE 11111111)은 그대로(주황 확인)",
      dl.status_code == 200 and C.layout_kind(dl.content) == "수책" and hs_of.get("SENSOR E2E-X5ME1") == "85365099"
      and hs_of.get("BEARING 6204ZZ") == "84821000" and hs_of.get("VALVE SY5120-5DZ") == "11111111"
      and j1["summary"]["auto"] == 2 and j1["summary"]["check"] == 1, (dl.status_code, hs_of, j1.get("summary")))
att = next((a for a in j1.get("attention", []) if a["desc"] == "VALVE SY5120-5DZ"), None)
rc = quiet(cl.post, "/export/hs-fill/%d/confirm" % lg1.get("id", 0),
           json={"sheet": att["sheet"] if att else "", "row": att["row"] if att else 0, "hs": "22222222"}) if att else None
jc = rc.json() if rc is not None else {}
again = open(os.path.join(M._HS_FILL_DIR, lg1["stored"]), "rb").read() if lg1.get("stored") else b""
hs2 = {it["desc"]: it.get("hs_raw") for s in H.scan_fast(again)["sheets"] for it in s["items"]} if again else {}
check("29 본사 확정 뒤 — 다시 만든 파일도 최종 양식(_in 에서) · 고른 값으로 칸 · 결과 카드에 바꿈 요약 그대로",
      rc is not None and rc.status_code == 200 and jc.get("ok") and (jc.get("conv") or {}).get("ok")
      and (jc.get("conv") or {}).get("kind") == "수책통관" and C.layout_kind(again) == "수책" and hs2.get("VALVE SY5120-5DZ") == "22222222",
      (rc.status_code if rc is not None else None, jc.get("error"), jc.get("conv"), hs2.get("VALVE SY5120-5DZ")))
s2, j2 = up(SU, "Invoice&packinglist_VN(261190-1b).xlsx", "somo")
check("30 화면에서 일반통관을 고르면 옛 수책도 일반통관으로(본사가 고름)",
      s2 == 200 and (j2.get("conv") or {}).get("kind") == "일반통관" and not (j2.get("conv") or {}).get("auto")
      and j2.get("forms") == ["일반통관"], (j2.get("conv"), j2.get("forms")))
s3, j3 = up(SO, "Invoice&packinglist_VN(261190-2).xlsx")
check("31 옛 일반 ① → 일반통관 · 8줄 늘림 · 사진 2 · 품목 60줄 모두 채움 대상",
      s3 == 200 and (j3.get("conv") or {}).get("kind") == "일반통관" and j3["conv"].get("expanded") == 8
      and j3["conv"].get("photos") == 2 and j3["summary"]["total"] == N_SOMO, (j3.get("conv"), j3.get("summary"), j3.get("error")))
s4, j4 = up(TWO, "Invoice&packinglist_VN(261190-34).xlsx")
lg4 = log_row("Invoice&packinglist_VN(261190-34).xlsx")
in4 = open(os.path.join(M._HS_FILL_DIR, lg4["orig_stored"]), "rb").read() if lg4.get("orig_stored") else b""
check("32 인보이스 두 장 → 올린 모양 그대로 채움 · 안내(바꾸지 못해 · 까닭 · 최종 양식으로 만들어 주세요) · _in = 올린 파일 · _src 없음",
      s4 == 200 and j4.get("ok") and (j4.get("conv") or {}).get("ok") is False and (j4.get("conv") or {}).get("multi")
      and "최종 양식으로 바꾸지 못해 올린 모양 그대로 채웠습니다" in (j4.get("hint") or "") and "인보이스 시트가 2개" in (j4.get("hint") or "")
      and "최종 양식(수책통관·일반통관)" in (j4.get("hint") or "") and in4 == TWO
      and not os.path.exists(os.path.join(M._HS_FILL_DIR, lg4.get("stored", "x.xlsx")[:-5] + "_src.xlsx")),
      (s4, j4.get("conv"), j4.get("hint")))
s5, j5 = up(PLAIN, "invoice_plain_261190.xlsx")
check("33 칸 모양이 다른 표 → 올린 모양 그대로(전과 같음) · 안내에 까닭 · 기록 conv_note(ok False)",
      s5 == 200 and j5.get("ok") and j5.get("forms") == ["예전 양식"] and "칸 모양" in (j5.get("hint") or "")
      and json.loads(log_row("invoice_plain_261190.xlsx").get("conv_note") or "{}").get("ok") is False, (j5.get("forms"), j5.get("hint")))
page = quiet(cl.get, "/export/hs-fill").text
check("34 기록 표 — 바꾼 줄 「🔁 옛 양식에서 바꿈」 · 못 바꾼 줄 「옛 양식 그대로(못 바꿈)」 · 화면 고르기 칸",
      page.count("🔁 옛 양식에서 바꿈") >= 3 and page.count("옛 양식 그대로(못 바꿈)") >= 2 and page.count('type="radio" name="hfFormKind"') == 3)
with D.db_session() as c:
    M._hs_fill_log_ensure(c)
    M._hs_fill_log_ensure(c)                                # 두 번 불러도 'duplicate column' 없이
    cols = [r[1] for r in c.execute("PRAGMA table_info(export_hs_fill_log)")]
check("35 기록 표 칸 conv_note(없을 때만 넣음 — 두 번 불러도 한 칸)", cols.count("conv_note") == 1, cols)

print()
print("=" * 78)
print("§4 실물(수출자료 · 이 PC 에만)")
print("=" * 78)
REAL = os.path.join(ROOT, "수출자료")
files = sorted(glob.glob(os.path.join(REAL, "**", "*.xlsx"), recursive=True))
if not files:
    print("  (수출자료 폴더 없음 — 건너뜀)")
else:
    st = {"final": 0, "ok": 0, "bad": [], "multi": 0, "err": []}
    for p in files:
        res = C.convert(open(p, "rb").read())
        if res.get("final"):
            st["final"] += 1
        elif res.get("converted") and res.get("ok"):
            st["ok"] += 1
        elif res.get("multi"):
            st["multi"] += 1
        elif res.get("converted"):
            st["bad"].append((os.path.basename(p), res.get("bad", [])[:2]))
        else:
            st["err"].append((os.path.basename(p), res.get("error")))
    print("  파일 %d · 이미 최종 %d · 바꿔 맞음 %d · 인보이스 여러 장(안 바꿈) %d" % (len(files), st["final"], st["ok"], st["multi"]))
    check("36 실물 — 바꾼 것은 모두 다시 읽어 대조가 맞음 · 오류 0 (안 바꾼 것은 인보이스 여러 장뿐)",
          not st["bad"] and not st["err"] and st["ok"] >= 1, (st["bad"][:3], st["err"][:3]))

print()
print("결과: %d개 중 실패 %d" % (OK + FAIL, FAIL))
sys.exit(1 if FAIL else 0)
