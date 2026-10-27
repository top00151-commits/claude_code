# -*- coding: utf-8 -*-
"""z1191m 시험 — 인보이스 표준 양식(수책·소모품) 내려받기
   대표 지시 2026-10-27 「본사에서 올릴 인보이스 표준양식(수책,소모품)을 여기서 받을수있게해줘.. 기본적은 공통 정보는 사전에 작성된 상태로..」
  §1 양식 파일(app/seed): 공통 정보 있음 · 번호·날짜·선적일 빔 · 품목 줄 빔(두 시트) · 금액·합계·패킹리스트 머리 수식 그대로 ·
     인쇄 영역 · 서명 그림 · 숨은 글자 없음 · 자동 채움이 알아봄
  §2 경로(임시 DB · TestClient): 두 양식 200(올바른 이름) · 모르는 종류 404 · 권한 없으면 303 · 화면에 단추 둘
  §3 사용 흉내: 양식에 품목을 적어 ① 자동 채움 → 줄 수·채운 파일 · 생성 도구로 다시 만들면 같은 내용(대표 원본이 있을 때)
"""
import contextlib
import io
import os
import re
import sys
import tempfile
import zipfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
os.environ["KNK_SSO_SERVICE_KEY"] = "z1191m-test-only-not-a-real-key"   # 시험용 가짜 값

import openpyxl                                            # noqa: E402
import app.database as D                                   # noqa: E402
from app import hs_fill as H                               # noqa: E402

OK = FAIL = 0


def check(label, cond, got=None):
    global OK, FAIL
    if cond:
        OK += 1
        print("  [OK] " + label)
    else:
        FAIL += 1
        print("  [FAIL] " + label + ("   ← %r" % (got,) if got is not None else ""))


SEED = os.path.join(ROOT, "app", "seed")
FORMS = {"suchaek": ("invoice_form_suchaek.xlsx", 186, 190, "S", "T"), "somo": ("invoice_form_somo.xlsx", 85, 89, "R", "S")}
COMMON = {"B4": "KNK CO.,LTD", "B12": "KNK VINA Co., Ltd.", "F10": "T/T", "B21": "SAME AS ABOVE", "G21": "INCOTERMS : C.I.P",
          "A25": "Incheon, Korea", "D25": "Hanoi, Vietnam", "A1": " COMMERCIAL  INVOICE"}

print("=" * 78)
print("§1 양식 파일")
print("=" * 78)
for kind, (fn, last_inv, last_pl, chs, cvn) in FORMS.items():
    p = os.path.join(SEED, fn)
    data = open(p, "rb").read()
    wb = openpyxl.load_workbook(io.BytesIO(data))
    inv, pl = wb["INVOICE"], wb["PACKING LIST"]
    miss = {k: inv[k].value for k, v in COMMON.items() if inv[k].value != v}
    check("1 %s: 공통 정보 미리 적힘(KNK·KNK VINA·T/T·SAME AS ABOVE·INCOTERMS·인천→하노이)" % kind,
          not miss and "Ji-Yeon" in str(inv["B8"].value) and "NGUYEN DINH TAM" in str(inv["B17"].value), miss)
    check("2 %s: 인보이스 번호·날짜·선적일 빔(G4·I4·D29)" % kind, inv["G4"].value is None and inv["I4"].value is None and inv["D29"].value is None,
          (inv["G4"].value, inv["I4"].value, inv["D29"].value))
    left = [c.coordinate for row in inv.iter_rows(min_row=34, max_row=last_inv) for c in row
            if c.value not in (None, "") and not (c.column_letter == "I" and str(c.value).startswith("="))
            and not (kind == "somo" and c.column_letter == "C")]
    left_pl = [c.coordinate for row in pl.iter_rows(min_row=38, max_row=last_pl) for c in row if c.value not in (None, "")]
    check("3 %s: 품목 줄 빔 — INVOICE %d줄(금액 수식만%s) · PACKING LIST %d줄" % (kind, last_inv - 33, " · 줄 번호" if kind == "somo" else "", last_pl - 37),
          not left and not left_pl, (left[:5], left_pl[:5]))
    tot_row = next(r for r in range(last_inv + 1, last_inv + 5) if inv.cell(r, 1).value == "TOTAL")
    check("4 %s: 금액 수식(I34)·합계 줄·패킹리스트 머리(=INVOICE!G4) 수식 그대로" % kind,
          str(inv["I34"].value).replace("+", "").startswith("=F34*H34") and str(inv.cell(tot_row, 9).value).startswith("=SUM(I33:")
          and pl["G4"].value == "=++INVOICE!G4", (inv["I34"].value, inv.cell(tot_row, 9).value, pl["G4"].value))
    check("5 %s: 인쇄 영역 A~I · 자동 채움 칸(%s·%s) 머리글 그대로" % (kind, chs, cvn),
          "$A$1:$I$" in (inv.print_area or "") and inv[chs + "31"].value == "HS CODE" and inv[cvn + "31"].value == "Tên hàng", inv.print_area)
    z = zipfile.ZipFile(io.BytesIO(data))
    ss = z.read("xl/sharedStrings.xml").decode("utf-8")
    leak = [w for w in ("DCP32DN", "PARTS OF CONNECTORS", "Cầu dao", "KNK-261015-1", "KNK-260921-1", "1,100*1,100") if w in ss]
    media = [n for n in z.namelist() if n.startswith("xl/media/")]
    check("6 %s: 숨은 글자 없음 · 그림은 서명·도장 쪽만(%d) · 계산 순서표 없음 · 열 때 다시 계산" % (kind, len(media)),
          not leak and 1 <= len(media) <= 2 and "xl/calcChain.xml" not in z.namelist() and "fullCalcOnLoad" in z.read("xl/workbook.xml").decode(),
          (leak, media))
    r0 = H.fill(data, [])
    check("7 %s: 자동 채움이 양식을 알아봄 — 비어 있으면 '품목 줄이 없습니다'" % kind, r0["ok"] is False and "품목 줄이 없습니다" in r0["error"], r0)
    wb.close()

print()
print("=" * 78)
print("§2 경로")
print("=" * 78)
TMP = tempfile.mkdtemp(prefix="knk_z1191m_")
_REAL = D.DB_PATH
D.DB_PATH = os.path.join(TMP, "t.db")
assert D.DB_PATH != _REAL
with contextlib.redirect_stdout(io.StringIO()):
    D.init_db()
with D.db_session() as conn:
    have = {r[1] for r in conn.execute("PRAGMA table_info(users)")}
    for col, typ in [("employee_no", "TEXT"), ("entity", "TEXT"), ("name_vi", "TEXT"), ("rank", "TEXT"), ("team_id", "INTEGER"),
                     ("password_version", "INTEGER DEFAULT 1")]:
        if col not in have:
            conn.execute("ALTER TABLE users ADD COLUMN %s %s" % (col, typ))
    U = {}
    for sb, fl in (("M901", (1, 1, 0, 0)), ("M902", (0, 0, 0, 0))):
        conn.execute("INSERT INTO users (name, role, employee_no, login_id, password, is_active, password_version, "
                     "can_view_sales, can_use_sales, can_view_logistics, can_use_logistics) VALUES (?,?,?,?,'x',1,1,?,?,?,?)",
                     ("가상" + sb, "member", sb, sb) + fl)
        U[sb] = dict(conn.execute("SELECT * FROM users WHERE employee_no=?", (sb,)).fetchone())
import app.main as M                                        # noqa: E402
from fastapi.testclient import TestClient                  # noqa: E402

M._HS_FILL_DIR = os.path.join(TMP, "hs_fill")
cl = TestClient(M.app, follow_redirects=False)


def as_user(sb):
    M.get_user = lambda request, _u=U[sb]: _u


def quiet(fn, *a, **k):
    with contextlib.redirect_stdout(io.StringIO()):
        return fn(*a, **k)


as_user("M901")
got = {}
for kind, (fn, *_r) in FORMS.items():
    r = quiet(cl.get, "/export/hs-fill/form/" + kind)
    got[kind] = r
cd_ok = all(r.status_code == 200 and "attachment" in r.headers.get("content-disposition", "") for r in got.values())
check("8 표준 양식 둘 200 · 받는 이름 「…(수책)_표준양식.xlsx」·「…(소모품)_표준양식.xlsx」 · 내용 = seed 파일",
      cd_ok and "%28%EC%88%98%EC%B1%85%29_%ED%91%9C%EC%A4%80%EC%96%91%EC%8B%9D.xlsx" in got["suchaek"].headers.get("content-disposition", "")
      and got["somo"].content == open(os.path.join(SEED, FORMS["somo"][0]), "rb").read(),
      {k: (v.status_code, v.headers.get("content-disposition")) for k, v in got.items()})
r404 = quiet(cl.get, "/export/hs-fill/form/../../etc")
r404b = quiet(cl.get, "/export/hs-fill/form/xyz")
check("9 모르는 종류 → 404(경로 장난 안 됨)", r404b.status_code == 404 and r404.status_code in (404, 405), (r404.status_code, r404b.status_code))
as_user("M902")
rno = quiet(cl.get, "/export/hs-fill/form/suchaek")
check("10 수출 권한 없으면 303 /home", rno.status_code == 303 and rno.headers.get("location") == "/home", rno.status_code)
as_user("M901")
pg = quiet(cl.get, "/export/hs-fill")
check("11 화면에 「표준 양식 내려받기」 단추 둘", pg.status_code == 200 and 'href="/export/hs-fill/form/suchaek"' in pg.text
      and 'href="/export/hs-fill/form/somo"' in pg.text, pg.status_code)

print()
print("=" * 78)
print("§3 사용 흉내")
print("=" * 78)
with D.db_session() as c:
    M._hs_dict_ensure(c)
    c.execute("DELETE FROM export_hs_dict")
    c.execute("INSERT INTO export_hs_dict(part_key,model,hs_code,vn_name,status,seen_count) VALUES('DCP32DN15A','DCP32DN 15A','85362011','Cầu dao điện','CONFIRMED',2)")
wb = openpyxl.load_workbook(io.BytesIO(got["suchaek"].content))
ws = wb["INVOICE"]
ws["G4"], ws["D29"] = "KNK-261099-1", "2026-10-30"
for i, (d, q, pr) in enumerate((("DCP32DN 15A", 11, 15.16), ("NEW-PART-1", 2, 3.0), ("NEW-PART-2", 5, 1.25))):
    ws["D%d" % (34 + i)], ws["F%d" % (34 + i)], ws["H%d" % (34 + i)] = d, q, pr
b = io.BytesIO()
wb.save(b)
rr = quiet(cl.post, "/export/hs-fill", files={"file": ("Invoice&packinglist_VN(261099).xlsx", b.getvalue())})
j = rr.json()
check("12 양식에 품목 3줄을 적어 ① 자동 채움 → 3줄(자동 1 · 법인 입력 2) · 인보이스 번호 · 수책 양식으로 알아봄",
      rr.status_code == 200 and j.get("ok") and j["summary"]["total"] == 3 and j["summary"]["auto"] == 1 and j["summary"]["need"] == 2
      and j["invoice_no"] == "KNK-261099-1" and j["forms"] == ["수책"], j.get("summary") or j)
SRC_DIR = os.environ.get("KNK_FORM_SRC_DIR") or os.path.dirname(os.path.dirname(os.path.dirname(ROOT)))   # 대표 원본 위치(이 PC: JR 폴더)
SRC = [os.path.join(SRC_DIR, "Invoice&packinglist_VN(%s).xlsx" % k) for k in ("수책", "소모품")]
SRC = [p for p in SRC if os.path.exists(p)]
if len(SRC) == 2:
    sys.path.insert(0, os.path.join(ROOT, "deploy"))
    import make_invoice_forms as MK                        # noqa: E402
    same = True
    for src, kind, numcol in ((SRC[0], "suchaek", None), (SRC[1], "somo", "C")):
        a = zipfile.ZipFile(io.BytesIO(MK.make(open(src, "rb").read(), keep_number_col=numcol)))
        bz = zipfile.ZipFile(os.path.join(SEED, FORMS[kind][0]))
        same = same and sorted(a.namelist()) == sorted(bz.namelist()) and all(a.read(n) == bz.read(n) for n in a.namelist())
    check("13 생성 도구로 대표 원본에서 다시 만들면 seed 양식과 내용이 같음(손으로 고친 것 없음)", same)
else:
    print("  (13 건너뜀 — 대표 원본 파일이 이 PC 경로에 없음)")

print()
print("결과: %d개 중 실패 %d" % (OK + FAIL, FAIL))
sys.exit(1 if FAIL else 0)
