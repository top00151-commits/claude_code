#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""미분류 원본 「본인 되받기」 제한형 — 시험 (대표 확정 2026-09-10)

대표 확정 문구
  "서버 기록으로 직접 업로더임이 확인되는 사용자가 해당 원본만 그대로 돌려받도록 하세요.
   산출물·보고문·남의 자료는 열지 않고, 기존 일반 열람권한은 유지합니다."

지시안 §3 의 최소 시험 5가지를 그대로 옮기고, 헛돌지 않게 역검사를 덧댔다.
⛔ 운영 DB 를 열지 않는다. 실명·연락처·거래내용 없음(가상 「시험○○」).
"""
import json
import os
import sys
import tempfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

import app.database as D                                   # noqa: E402

TMP = tempfile.mkdtemp(prefix="knk_selforig_")
_REAL = D.DB_PATH
D.DB_PATH = os.path.join(TMP, "t.db")
assert D.DB_PATH != _REAL
D.init_db()

import app.main as appmain                                 # noqa: E402
from fastapi.testclient import TestClient                  # noqa: E402
from openpyxl import Workbook                              # noqa: E402

appmain._BT_STORE = os.path.join(TMP, "store")
os.makedirs(appmain._BT_STORE, exist_ok=True)
CUR = {"u": None}
appmain.get_user = lambda request: CUR["u"]
cl = TestClient(appmain.app, follow_redirects=False)

# 팀4 = BOM 만들기 O · 구매 실행 X · 단가 O · 구매처 X  (운영 실측 조합)
A = {"id": 901, "name": "시험설계A", "role": "member", "team_id": 4,
     "can_use_logistics": 0, "can_view_logistics": 1}
B = {"id": 902, "name": "시험설계B", "role": "member", "team_id": 4,
     "can_use_logistics": 0, "can_view_logistics": 1}
BUYER = {"id": 903, "name": "시험구매", "role": "member", "team_id": 10,
         "can_use_logistics": 1, "can_view_logistics": 1}

with D.db_session() as conn:
    for uu in (A, B, BUYER):
        conn.execute("INSERT INTO users(id, name, login_id, password, role) VALUES(?,?,?,'x',?)",
                     (uu["id"], uu["name"], "P%d" % uu["id"], uu["role"]))
    conn.execute("DELETE FROM field_access_policy")
    for cat, teams in (("purchase_price", (1, 2, 4, 10, 11, 15)),
                       ("supplier_info", (1, 10, 11, 15))):
        for t in teams:
            conn.execute("INSERT INTO field_access_policy(category, team_id) VALUES(?,?)", (cat, t))

FAIL, CNT = [], 0


def check(name, ok, detail=""):
    global CNT
    CNT += 1
    print(("  ✅" if ok else "  ❌") + f" {CNT:02d} {name}" + ("" if ok else f" — {detail}"))
    if not ok:
        FAIL.append(name)


def mk_unknown(path):
    """우리 PART LIST 양식이 **아닌** 파일 = 분류할 수 없는 원본."""
    wb = Workbook()
    ws = wb.active
    ws.cell(1, 1, "설계 메모")
    ws.cell(2, 1, "이 파일은 회사 양식이 아니다")
    ws.cell(3, 1, "TEST-UNKNOWN-001")
    wb.save(path)
    return path


def mk_ours(path, vendor="가업체", price=1000):
    """우리 통일판 양식 = 분류되는 원본(단가·협력사 들어 있음)."""
    wb = Workbook()
    ws = wb.active
    ws.cell(7, 2, "NO.")
    ws.cell(7, 5, "PRODUCT NAME 제품명")
    ws.cell(7, 8, "FINAL VENDOR 외주사")
    ws.cell(7, 20, "UNIT PRICE 단가")
    ws.cell(9, 5, "시험품")
    ws.cell(9, 8, vendor)
    ws.cell(9, 20, price)
    wb.save(path)
    return path


def seed(owner_id, in_path, out_grade=(None, None), report="보고 — 품목 3줄"):
    """실행 기록 한 줄을 직접 심는다. owner_id = 서버가 기록한 업로더."""
    outp = os.path.join(appmain._BT_STORE, "out_%d.xlsx" % owner_id)
    Workbook().save(outp)
    with D.db_session() as c:
        cur = c.execute(
            "INSERT INTO bom_tool_runs(mgmt_code, step, title, inputs, output_name, "
            "output_path, report, created_by, has_price, has_vendor) "
            "VALUES(?,?,?,?,?,?,?,?,?,?)",
            ("005M2606", "draft", "시험", json.dumps([{"name": os.path.basename(in_path),
                                                      "path": in_path}]),
             "out.xlsx", outp, report, owner_id, out_grade[0], out_grade[1]))
        return cur.lastrowid


print("=" * 72)
print("  미분류 원본 「본인 되받기」 제한형 — 시험")
print("=" * 72)

unk_a = mk_unknown(os.path.join(appmain._BT_STORE, "A_원본_미분류.xlsx"))
run_a = seed(A["id"], unk_a)

# ① 본인이 올린 미분류 원본은 본인이 그대로 받는다
CUR["u"] = A
r1 = cl.get(f"/bom/tools/download/{run_a}?f=in0")
check("① 설계A 가 직접 올린 미분류 원본을 A 가 그대로 받는다",
      r1.status_code == 200, f"{r1.status_code} {r1.headers.get('location','')[:60]}")

# ② 남의 원본은 안 열린다
CUR["u"] = B
r2 = cl.get(f"/bom/tools/download/{run_a}?f=in0")
check("② 설계B 가 A 의 원본을 요청하면 본인 예외가 안 걸린다",
      r2.status_code == 303 and "err=" in r2.headers.get("location", ""), str(r2.status_code))

# ③ 화면 표시 이름을 바꿔도 실제 업로더는 안 바뀐다
with D.db_session() as c:
    c.execute("UPDATE bom_tool_runs SET title=? WHERE id=?", ("시험설계A 작성", run_a))
CUR["u"] = B
r3 = cl.get(f"/bom/tools/download/{run_a}?f=in0")
check("③ 표시 이름을 A 로 바꿔도 B 는 못 받는다 (created_by 만 본다)",
      r3.status_code == 303 and "err=" in r3.headers.get("location", ""), str(r3.status_code))

# ④ 본인 예외로 산출물·보고문이 열리지 않는다
CUR["u"] = A
r4 = cl.get(f"/bom/tools/download/{run_a}?f=out")
check("④-1 본인 예외로 **산출물**이 열리지 않는다",
      r4.status_code == 303 and "err=" in r4.headers.get("location", ""), str(r4.status_code))
page = cl.get("/bom/tools")
check("④-2 본인 예외로 **보고문**이 열리지 않는다 (목록에서 자물쇠)",
      ("보고 — 품목 3줄" not in page.text), "보고문이 화면에 보인다")
check("④-3 본인 예외가 **구매 실행권한**을 주지 않는다",
      appmain._bom_can_purchase(A) is False)

# ⑤ 기존 일반 열람권한은 그대로
#   ⚠ 실행기록 **자체**의 등급이 불명(NULL)이면 줄 관문이 구매담당도 막는다 —
#      2026-09-08 부터의 기존 동작이고 이번 예외와 무관하다. 그래서 현실과 같게
#      산출물이 정상 판정된 실행(0,0)에 미분류 원본이 붙은 경우로 시험한다.
run_g = seed(A["id"], unk_a, out_grade=(0, 0))
CUR["u"] = BUYER
r5 = cl.get(f"/bom/tools/download/{run_g}?f=in0")
check("⑤ 두 권한을 갖춘 구매 담당은 기존대로 받는다", r5.status_code == 200, str(r5.status_code))
CUR["u"] = A
r5b = cl.get(f"/bom/tools/download/{run_g}?f=in0")
check("⑤-2 같은 실행에서 업로더 본인도 자기 원본을 받는다",
      r5b.status_code == 200, str(r5b.status_code))
CUR["u"] = B
r5c = cl.get(f"/bom/tools/download/{run_g}?f=in0")
check("⑤-3 같은 실행이라도 남(설계B)은 여전히 막힌다",
      r5c.status_code == 303 and "err=" in r5c.headers.get("location", ""), str(r5c.status_code))
CUR["u"] = BUYER
r5d = cl.get(f"/bom/tools/download/{run_a}?f=in0")
check("⑤-4 (기존 동작 보존) 실행기록 등급이 불명이면 구매담당도 막힌다 — 이번 변경과 무관",
      r5d.status_code == 303, str(r5d.status_code))

# ⑥ 분류되는(단가·협력사 든) 원본은 본인이라도 기존 규칙을 따른다
ours = mk_ours(os.path.join(appmain._BT_STORE, "A_원본_우리양식.xlsx"))
run_o = seed(A["id"], ours)
CUR["u"] = A
r6 = cl.get(f"/bom/tools/download/{run_o}?f=in0")
check("⑥ 본인이 올렸어도 **분류되는** 원본(구매처 포함)은 기존 규칙대로 막힌다",
      r6.status_code == 303 and "err=" in r6.headers.get("location", ""),
      f"{r6.status_code} — 팀4 는 구매처 열람권한이 없다")

# ⑦ 역검사 — 업로더 기록이 없으면(과거 자료) 본인 예외가 안 걸린다
run_n = seed(A["id"], unk_a)
with D.db_session() as c:
    c.execute("UPDATE bom_tool_runs SET created_by=NULL WHERE id=?", (run_n,))
CUR["u"] = A
r7 = cl.get(f"/bom/tools/download/{run_n}?f=in0")
check("⑦ 역검사: 업로더 기록이 없는 과거 자료는 예외가 안 걸린다",
      r7.status_code == 303 and "err=" in r7.headers.get("location", ""), str(r7.status_code))

# ⑧ 역검사 — 예외가 없었다면 ①도 막혔어야 한다 (시험이 헛돌지 않음)
_keep = appmain._bt_grade_file
try:
    appmain._bt_grade_file = lambda p: (1, 1)     # 미분류가 아니게 만든다
    CUR["u"] = A
    r8 = cl.get(f"/bom/tools/download/{run_a}?f=in0")
finally:
    appmain._bt_grade_file = _keep
check("⑧ 역검사: 미분류가 아니면 같은 요청이 실제로 막힌다",
      r8.status_code == 303 and "err=" in r8.headers.get("location", ""), str(r8.status_code))

# ⑨ 같은 관리번호의 **다른 실행**에 든 남의 원본은 안 열린다
unk_b = mk_unknown(os.path.join(appmain._BT_STORE, "B_원본_미분류.xlsx"))
run_b = seed(B["id"], unk_b)
CUR["u"] = A
r9 = cl.get(f"/bom/tools/download/{run_b}?f=in0")
check("⑨ 같은 관리번호라도 남이 올린 실행의 원본은 안 열린다",
      r9.status_code == 303 and "err=" in r9.headers.get("location", ""), str(r9.status_code))

# ⑩ 받은 파일이 **가공 없이 그대로**인지
CUR["u"] = A
r10 = cl.get(f"/bom/tools/download/{run_a}?f=in0")
with open(unk_a, "rb") as fh:
    same = fh.read() == r10.content
check("⑩ 받은 파일이 올린 원본과 **바이트까지 같다** (가공 없음)", same,
      f"{len(r10.content)} bytes")

print("-" * 72)
print(f"  시험 {CNT}건 · 실패 {len(FAIL)}건" + ("" if not FAIL else " → " + ", ".join(FAIL)))
sys.exit(1 if FAIL else 0)
