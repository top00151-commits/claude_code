# -*- coding: utf-8 -*-
"""z1191v 시험 — 품명 칸의 엑셀 함수 실패 값(「0」 · #N/A …)은 데이터가 아니다
   대표 2026-10-28 「0 , #N/A 이건 엑셀 데이터가 함수로 연결되어서 오는건데 잘못 된거야... 이렇게 나오는 부분들은 HS코드 등록에서
   모두 삭제해.」 → 운영 사전 44줄 삭제(15:45) → 「1번 재발 막기 진행해」
  §1 원문: 판정 한 곳(hs_fill.formula_junk) · ③ · 사전 올림 · 📚 · 📌 · 같은 날 두 HS · 화면 글
  §2 ① 채우기: 적힌 함수 값 → 사전 품명으로 덮음 / 비우고 노랑 · 사전에 남은 함수 값은 안 씀 · 수식 칸은 건드리지 않음(직접 확인)
  §3 ③ 회신 · HS 사전 올림: 함수 값 줄은 사전에 안 넣고 수만 · 이력 줄은 그대로(실제 기록)
  §4 📚 등록: 함수 값 줄만 있는 규격은 사전에 안 넣음
  §5 📌: 품명이 함수 값인 이력 줄은 「가장 최근 값」으로 안 씀 · 같은 날 두 HS 셈에서도 뺌
"""
import ast
import contextlib
import io
import os
import re
import sys
import tempfile
import zipfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
os.environ["KNK_SSO_SERVICE_KEY"] = "z1191v-test-only-not-a-real-key"   # 시험용 가짜 값

import openpyxl                                            # noqa: E402

import app.database as D                                   # noqa: E402
from app import hs_fill as H                               # noqa: E402

OK = FAIL = 0


def check(label, cond, got=None):
    global OK, FAIL
    if cond:
        OK += 1
        print("  [OK] " + label)
    else:
        FAIL += 1
        print("  [FAIL] " + label + ("   ← %r" % (got,) if got is not None else ""))


def quiet(fn, *a, **k):
    with contextlib.redirect_stdout(io.StringIO()):
        return fn(*a, **k)


def make_form(items, ino, formula_cells=()):
    """수책통관 최종 양식 칸 모양(HS=S · Tên hàng=T) — items=[(규격, HS, 품명)] · formula_cells=[(줄 번째, 저장값)]:
    그 줄 Tên hàng 칸을 엑셀 함수(IFERROR(…,"0"))로 · 저장값을 XML 로 넣는다(엑셀로 저장한 파일처럼)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INVOICE"
    ws["A1"] = "COMMERCIAL  INVOICE"
    ws["G4"] = ino
    heads = {"A": "12) Marks and Number of PKGS", "D": "13) Description of Goods", "F": " 14) Q'ty /Unit", "H": "15) Unit Price",
             "I": "16) Amount", "K": "COMPANY", "L": "관리코드", "M": "MODEL", "N": "품명", "O": "Q'TY", "P": "ORIGIN",
             "Q": "SUPPLIER", "R": "CATEGORY", "S": "HS CODE", "T": "Tên hàng", "U": "SHIPPING DATE"}
    for col, v in heads.items():
        ws[col + "31"] = v
    r = 34
    for desc, hs, vn in items:
        ws["D%d" % r], ws["F%d" % r], ws["H%d" % r], ws["I%d" % r] = desc, 2, 3.5, 7
        if hs is not None:
            ws["S%d" % r] = hs
        if vn is not None:
            ws["T%d" % r] = vn
        r += 1
    ws["A%d" % (r + 1)] = "TOTAL"
    wb.create_sheet("PACKING LIST")["A1"] = "PACKING LIST"
    for i, _cached in formula_cells:
        ws["T%d" % (34 + i)] = '=IFERROR(INDEX(Z1:Z2,MATCH("x",Z1:Z2,0)),"0")'
    b = io.BytesIO()
    wb.save(b)
    data = b.getvalue()
    if not formula_cells:
        return data
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for n in zin.namelist():
            buf = zin.read(n)
            if n == "xl/worksheets/sheet1.xml":
                x = buf.decode("utf-8")
                for i, cached in formula_cells:
                    ref = "T%d" % (34 + i)
                    x = re.sub(r'<c r="%s"([^>]*)><f>(.*?)</f>(?:<v\s*/>|<v></v>)?</c>' % ref,
                               lambda m: '<c r="%s"%s t="str"><f>%s</f><v>%s</v></c>' % (
                                   ref, re.sub(r'\s+t="[^"]*"', "", m.group(1)), m.group(2), cached), x)
                buf = x.encode("utf-8")
            zo.writestr(n, buf)
    return out.getvalue()


def cell_xml(data, ref):
    x = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml").decode("utf-8")
    m = re.search(r'<c r="%s"[^>]*?(?:/>|>.*?</c>)' % ref, x, re.S)
    return m.group(0) if m else ""


print("=" * 78)
print("§1 원문")
print("=" * 78)
cases_true = ["0", " 0 ", "0.0", "00", "-0", 0, 0.0, "#N/A", "#n/a", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!", "#SPILL!"]
cases_false = [None, "", "10", 5, "Cảm biến quang", "TIMING BELT#&Dây đai răng", "N/A", "0A", "#&Van", False, True]
check("1 판정 한 곳 hs_fill.formula_junk — 「0」·숫자 0·#N/A 같은 오류 표시는 참 · 빈칸·숫자·글자·「#&」 들어간 품명은 거짓",
      all(H.formula_junk(v) for v in cases_true) and not any(H.formula_junk(v) for v in cases_false),
      ([v for v in cases_true if not H.formula_junk(v)], [v for v in cases_false if H.formula_junk(v)]))
SRC = open(os.path.join(ROOT, "app", "main.py"), encoding="utf-8").read()
FN = {n.name: ast.get_source_segment(SRC, n) for n in ast.parse(SRC).body
      if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))
      and n.name in ("_hs_parse_invoice_xlsx", "_hs_learn_invoice_rows", "_hs_hist_dict_sync", "_hs_hist_latest", "_hs_hist_summary",
                     "_hs_junk_fn")}
check("2 사전 성장(③ · 사전 올림) — 읽을 때 품명 칸 원래 값(숫자 0 포함)으로 junk 표시 · 그 줄은 건너뛰고 수만(junk)",
      "_hf_p.formula_junk(_vraw)" in FN["_hs_parse_invoice_xlsx"] and '"junk": _junk' in FN["_hs_parse_invoice_xlsx"]
      and 'if row.get("junk"):' in FN["_hs_learn_invoice_rows"] and '"junk": junk' in FN["_hs_learn_invoice_rows"])
check("3 📚 사전 판단 — 함수 값 이력 줄 빼고 · 깨끗한 줄이 없으면 안 넣고 셈 · 📌·같은 날 두 HS — SQL 에 같은 판정(formula_junk)",
      "clean = [r for r in rows if not _hf.formula_junk(r[\"vn_name\"])]" in FN["_hs_hist_dict_sync"]
      and FN["_hs_hist_latest"].count("formula_junk(") >= 4 and "_hs_junk_fn(c)" in FN["_hs_hist_latest"]
      and FN["_hs_hist_summary"].count("formula_junk(") == 2
      and 'create_function("formula_junk", 1, _hf.formula_junk_sql, deterministic=True)' in FN["_hs_junk_fn"])
TF = open(os.path.join(ROOT, "app", "templates", "export_hs_fill.html"), encoding="utf-8").read()
TD = open(os.path.join(ROOT, "app", "templates", "export_hs_dict.html"), encoding="utf-8").read()
check("4 화면 — ③ 설명·결과 · 📚 결과 · HS 사전 올림 결과·요약에 「함수 값 줄은 안 넣음」",
      "품명이 「0」·「#N/A」(엑셀 함수가 못 찾은 값)인 줄은 사전에 넣지 않습니다" in TF and "if(x.junk){" in TF and "if(d.junk)" in TF
      and "x.junk? ' · ⚠품명 「0」·「#N/A」(엑셀 함수 값) '" in TD and "jk? ' · ⚠품명 「0」·「#N/A」(엑셀 함수 값) '" in TD
      and "var a=0,u2=0,cf=0,er=0,oc=0,jk=0" in TD)

print()
print("=" * 78)
print("§2 ① 채우기")
print("=" * 78)
ENTRIES = [
    {"part_key": "PART-A1", "model": "PART-A1", "hs_code": "11111111", "vn_name": "Tên A", "status": "CONFIRMED"},
    {"part_key": "PART-B2", "model": "PART-B2", "hs_code": "22222222", "vn_name": "", "status": "CONFIRMED"},
    {"part_key": "PART-D4", "model": "PART-D4", "hs_code": "33333333", "vn_name": "Tên D", "status": "CONFIRMED"},
    {"part_key": "PART-E5", "model": "PART-E5", "hs_code": "55555555", "vn_name": "Tên E sổ", "status": "CONFIRMED"},
    {"part_key": "PART-F6", "model": "PART-F6", "hs_code": "66666666", "vn_name": "0", "status": "CONFIRMED"},     # 사전에 남은 함수 값
    {"part_key": "PART-G7", "model": "PART-G7", "hs_code": "77777777", "vn_name": "Tên G", "status": "CONFIRMED"},
]
INV = make_form([("PART-A1", None, "0"), ("PART-B2", None, "#N/A"), ("PART-C3", None, 0), ("PART-D4", "33333333", "0"),
                 ("PART-E5", None, "Tên E thật"), ("PART-F6", None, None), ("PART-G7", None, None)], "KNK-261029-1",
                formula_cells=[(6, "0")])
sc0 = H.scan_fast(INV)
check("5 시험 파일 — 품명 칸 저장값 그대로 읽힘(「0」 · #N/A · 숫자 0 → 「0」 · 수식 칸 저장값 「0」)",
      [it["vn_raw"] for it in sc0["sheets"][0]["items"]] == ["0", "#N/A", "0", "0", "Tên E thật", "", "0"],
      [it["vn_raw"] for it in sc0["sheets"][0]["items"]])
res = H.fill(INV, ENTRIES)
acts = {a["desc"]: a for a in res["acts"]}
sc1 = H.scan_fast(res["out"])
vn1 = {it["desc"]: it["vn_raw"] for it in sc1["sheets"][0]["items"]}
hs1 = {it["desc"]: it["hs_raw"] for it in sc1["sheets"][0]["items"]}
check("6 적힌 「0」 + 사전 품명 있음 → 사전 품명으로 덮음(자동) · 까닭에 「엑셀 함수 값 → 사전 품명으로 채움」",
      vn1["PART-A1"] == "Tên A" and hs1["PART-A1"] == "11111111" and H.category(acts["PART-A1"]) == H.AUTO
      and "엑셀 함수 값) → 사전 품명으로 채움" in acts["PART-A1"]["why"], (vn1["PART-A1"], acts["PART-A1"]["why"]))
check("7 적힌 #N/A + 사전 품명 없음 → 칸을 비우고 노랑(법인 입력) · HS 는 사전 값",
      vn1["PART-B2"] == "" and hs1["PART-B2"] == "22222222" and acts["PART-B2"]["color_vn"] == H.YELLOW
      and H.category(acts["PART-B2"]) == H.NEED and "엑셀 함수 값이라 비움" in acts["PART-B2"]["why"]
      and re.fullmatch(r'<c r="T35" s="\d+"/>', cell_xml(res["out"], "T35")), (vn1["PART-B2"], cell_xml(res["out"], "T35")))
check("8 적힌 숫자 0 + 사전에 없음 → HS·품명 노랑 · 품명 칸 비움",
      vn1["PART-C3"] == "" and acts["PART-C3"]["color_hs"] == H.YELLOW and acts["PART-C3"]["color_vn"] == H.YELLOW
      and acts["PART-C3"]["status"] == H.NEED, (vn1["PART-C3"], acts["PART-C3"]))
check("9 HS 적힘(사전과 같음) + 품명 「0」 → 품명만 사전 값으로(「HS 는 이미 같음 · 품명만 채움」)",
      vn1["PART-D4"] == "Tên D" and acts["PART-D4"]["status"] == H.AUTO and "품명만 채움" in acts["PART-D4"]["why"], acts["PART-D4"])
check("10 깨끗한 적힌 품명은 그대로(덮지 않음) — 함수 값만 빈칸으로 본다",
      vn1["PART-E5"] == "Tên E thật" and acts["PART-E5"]["write_vn"] is None, acts["PART-E5"])
check("11 사전에 남은 함수 값(「0」)은 품명으로 안 씀 → 노랑(법인 입력) · 「0」을 채우지 않음",
      vn1["PART-F6"] == "" and acts["PART-F6"]["color_vn"] == H.YELLOW and acts["PART-F6"]["dict_vn"] == "", (vn1["PART-F6"], acts["PART-F6"]))
f40 = cell_xml(res["out"], "T40")
check("12 수식이 든 품명 칸(저장값 「0」) — 값을 안 쓰고(수식 그대로) 「직접 확인」(수식은 건드리지 않는다)",
      "<f>" in f40 and acts["PART-G7"]["status"] == H.CONFLICT and "수식이 든 칸" in acts["PART-G7"]["why"]
      and "T40" in res["skipped"], (f40[:120], acts["PART-G7"]["status"], res["skipped"]))
check("13 셈 — 법인 입력 = B2 · C3 · F6 (· G7 은 확인) · 회신 필요 있음",
      res["counts"]["need"] >= 3 and not res["counts"]["no_reply"], res["counts"])

print()
print("=" * 78)
print("§3 ③ 회신 · HS 사전 올림")
print("=" * 78)
TMP = tempfile.mkdtemp(prefix="knk_z1191v_")
_REAL = D.DB_PATH
D.DB_PATH = os.path.join(TMP, "t.db")
assert D.DB_PATH != _REAL
with contextlib.redirect_stdout(io.StringIO()):
    D.init_db()
with D.db_session() as conn:
    have = {r[1] for r in conn.execute("PRAGMA table_info(users)")}
    for col, typ in [("employee_no", "TEXT"), ("entity", "TEXT"), ("name_vi", "TEXT"), ("rank", "TEXT"), ("team_id", "INTEGER"),
                     ("password_version", "INTEGER DEFAULT 1")]:
        if col not in have:
            conn.execute("ALTER TABLE users ADD COLUMN %s %s" % (col, typ))
    conn.execute("INSERT INTO users (name, role, employee_no, login_id, password, is_active, password_version, "
                 "can_view_sales, can_use_sales, can_view_logistics, can_use_logistics) VALUES ('가상V901','member','V901','V901','x',1,1,1,1,0,0)")
    U = dict(conn.execute("SELECT * FROM users WHERE employee_no='V901'").fetchone())
import app.main as M                                        # noqa: E402
from fastapi.testclient import TestClient                  # noqa: E402

M._HS_FILL_DIR = os.path.join(TMP, "hs_fill")
M._HS_HIST_IMG_DIR = os.path.join(TMP, "hs_hist_img")
M.get_user = lambda request, _u=U: _u
cl = TestClient(M.app, follow_redirects=False)
with D.db_session() as c:
    M._hs_dict_ensure(c)
    M._hs_hist_ensure(c)
    for t in ("export_hs_dict", "export_hist_line", "export_hist_inv"):
        c.execute("DELETE FROM %s" % t)
    c.execute("INSERT INTO export_hs_dict(part_key,model,vn_name,hs_code,status,seen_count,note,updated_at) "
              "VALUES('PART-SAME','PART-SAME','Tên cũ','12345678','CONFIRMED',4,'',datetime('now','localtime'))")


def drow(k):
    with D.db_session() as c:
        r = c.execute("SELECT * FROM export_hs_dict WHERE part_key=?", (k,)).fetchone()
        return dict(r) if r else None


def q1(sql, *a):
    with D.db_session() as c:
        return c.execute(sql, a).fetchone()[0]


REPLY = make_form([("PART-H8", "66666661", "Tên H"), ("PART-I9", "77777771", "0"), ("PART-J10", "88888881", "#N/A"),
                   ("PART-K11", "99999991", 0), ("PART-SAME", "12345678", "0")], "KNK-261029-2")
rr = quiet(cl.post, "/export/hs-fill/reply", files=[("files", ("Invoice&packinglist_VN(261029-2)-re.xlsx", REPLY))]).json()
x = (rr.get("results") or [{}])[0]
check("14 ③ 회신 — 깨끗한 줄만 사전에(새 1) · 함수 값 줄 4개(「0」·#N/A·숫자 0·같은 HS 재출현도)는 안 넣고 junk 4",
      rr.get("ok") and x.get("added") == 1 and x.get("junk") == 4 and drow("PART-H8") and not drow("PART-I9")
      and not drow("PART-J10") and not drow("PART-K11"), x)
same = drow("PART-SAME") or {}
check("15 ③ — 사전에 있던 규격의 함수 값 줄은 횟수·품명을 안 건드림(그대로 4번 · 「Tên cũ」)",
      same.get("seen_count") == 4 and same.get("vn_name") == "Tên cũ", same)
check("16 ③ — 지난 수출 이력에는 5줄 그대로(실제 인보이스 기록 · 품명 칸 값도 그대로)",
      x.get("hist", {}).get("lines") == 5 and q1("SELECT COUNT(*) FROM export_hist_line WHERE invoice_no='KNK-261029-2'") == 5
      and q1("SELECT vn_name FROM export_hist_line WHERE invoice_no='KNK-261029-2' AND part_key='PART-I9'") == "0", x.get("hist"))
UP = make_form([("PART-N14", "11223344", "Tên N"), ("PART-O15", "55667788", "#N/A")], "KNK-261029-3")
ru = quiet(cl.post, "/export/hs-dict/upload", files=[("files", ("Invoice&packinglist_VN(261029-3).xlsx", UP))]).json()
y = (ru.get("results") or [{}])[0]
check("17 HS 사전 올림 — 같은 규칙(새 1 · junk 1 · #N/A 규격은 사전에 없음)",
      ru.get("ok") and y.get("added") == 1 and y.get("junk") == 1 and drow("PART-N14") and not drow("PART-O15"), y)

print()
print("=" * 78)
print("§4 📚 등록")
print("=" * 78)
HIST = make_form([("PART-P16", "10101010", "Tên P"), ("PART-Q17", "20202020", "0"), ("PART-R18", "30303030", "#N/A")], "KNK-261029-4")
rh = quiet(cl.post, "/export/hs-fill/hist", data={"kind": "수책통관"}, files={"file": ("Invoice&packinglist_VN(261029-4).xlsx", HIST)})
jh = rh.json()
check("18 📚 등록 — 이력 3줄 그대로 · 사전은 깨끗한 규격만(P16) · 함수 값 줄만 있는 규격 2개는 안 넣고 junk 2",
      rh.status_code == 200 and jh.get("lines") == 3 and (jh.get("dict") or {}).get("junk") == 2 and drow("PART-P16")
      and not drow("PART-Q17") and not drow("PART-R18"), (rh.status_code, jh.get("dict"), jh.get("error")))

print()
print("=" * 78)
print("§5 📌 최근 수출 값으로 확정 · 같은 날 두 HS")
print("=" * 78)
with D.db_session() as c:
    for k, hs, st in (("PART-L12", "12121212", "CONFIRMED"), ("PART-M13", "13131313", "PENDING"), ("PART-S19", "19191919", "PENDING")):
        c.execute("INSERT INTO export_hs_dict(part_key,model,vn_name,hs_code,status,seen_count,note,updated_at) "
                  "VALUES(?,?,?,?,?,1,'',datetime('now','localtime'))", (k, k, "", hs, st))
    rows = [("KNK-260901-9", "2026-09-01", "PART-L12", "12121212", "Tên L"), ("KNK-261001-9", "2026-10-01", "PART-L12", "34343434", "0"),
            ("KNK-261001-9", "2026-10-01", "PART-M13", "13131313", "Tên M"), ("KNK-261001-9", "2026-10-01", "PART-M13", "45454545", "#N/A"),
            ("KNK-261001-9", "2026-10-01", "PART-S19", "56565656", "0")]
    for i, (ino, d, k, hs, vn) in enumerate(rows):
        c.execute("INSERT INTO export_hist_line(invoice_no,inv_date,kind,sheet,row_no,part_key,descr,qty,unit_price,amount,hs_code,vn_name) "
                  "VALUES(?,?,'수책통관','INVOICE',?,?,?,1,1,1,?,?)", (ino, d, 34 + i, k, k, hs, vn))
    latest = {x["part_key"]: x for x in M._hs_hist_latest(c)}
    summ = M._hs_hist_summary(c)
check("19 📌 — 최근 줄이 함수 값(「0」)이면 그 앞의 깨끗한 줄이 기준 → PART-L12 은 사전과 같아 대상 아님(34343434 · 「0」으로 안 되돌림)",
      "PART-L12" not in latest, latest.get("PART-L12"))
check("20 📌 — 같은 날 줄 중 함수 값(#N/A) 줄은 빼고 셈 → PART-M13 은 HS 하나(13131313) · 확인 대기 → 확정 대상(품명 「Tên M」)",
      latest.get("PART-M13", {}).get("hs") == "13131313" and latest["PART-M13"].get("vn") == "Tên M"
      and "PART-M13" not in {s["part_key"] for s in summ.get("split", [])}, (latest.get("PART-M13"), summ.get("split")))
check("21 📌 — 함수 값 줄밖에 없는 규격(PART-S19)은 대상 아님(이력 HS 56565656 로 확정하지 않음)",
      "PART-S19" not in latest, latest.get("PART-S19"))

print()
print("결과: %d개 중 실패 %d" % (OK + FAIL, FAIL))
sys.exit(1 if FAIL else 0)
