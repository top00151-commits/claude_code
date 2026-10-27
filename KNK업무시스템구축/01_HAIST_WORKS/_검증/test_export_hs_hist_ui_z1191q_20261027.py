# -*- coding: utf-8 -*-
"""z1191q 시험 — 📚 등록이 끝난 것이 보이게(대표 2026-10-27 「수책통관 14개 올렸는데.. 지금 등록이 완료 된거야??
   완료된 표가 안나는데... 파일올리는데 파일 이름도 그대로 있고...」)
  원문(화면 틀·스크립트): 등록된 인보이스 목록은 펼친 채로 · 줄마다 「✅ 등록 완료」 · 끝나면 완료 띠(줄·대조·사전) ·
       처리 뒤 고른 파일과 상자 안 이름 지우기(📚·①·③ 셋 다) · 화면이 그려지는지(임시 DB · 등록 1개 뒤 목록 펼침)
  (실제 누름·지움은 브라우저 실측으로 — scratchpad z1191lw/_shot_page_q.py)
"""
import contextlib
import io
import os
import re
import sys
import tempfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
os.environ["KNK_SSO_SERVICE_KEY"] = "z1191q-test-only-not-a-real-key"   # 시험용 가짜 값

import openpyxl                                            # noqa: E402
import app.database as D                                   # noqa: E402

OK = FAIL = 0


def check(label, cond, got=None):
    global OK, FAIL
    if cond:
        OK += 1
        print("  [OK] " + label)
    else:
        FAIL += 1
        print("  [FAIL] " + label + ("   ← %r" % (got,) if got is not None else ""))


tpl = open(os.path.join(ROOT, "app", "templates", "export_hs_fill.html"), encoding="utf-8").read()
js = tpl[tpl.index("<script>"):]
print("=" * 78)
print("§1 원문")
print("=" * 78)
check("1 등록된 인보이스 목록은 펼친 채로(<details … open>) · 사전과 다른 규격 목록은 접힌 채로",
      re.search(r'<details class="hf-det" open><summary>등록된 인보이스 목록', tpl) is not None
      and re.search(r'<details class="hf-det"><summary style="color:#b91c1c;">HS 사전과 HS 가 다른 규격', tpl) is not None)
check("2 줄마다 「✅ 등록 완료」 + 새로 넣음 / 같은 번호 있던 것 바꿔 넣음",
      "✅ 등록 완료" in js and "새로 넣음" in js and "같은 번호 있던 것 바꿔 넣음" in js)
check("3 끝나면 완료 띠 — 「N개 등록 완료」·품목 줄·원본↔저장·HS 사전 새 규격 · 다 못 넣으면 「N개 중 M개」",
      "개 등록 완료" in js and "원본↔저장 " in js and "HS 사전 새 규격 확정 " in js and "개 중 " in js and "hf-hist-done" in js)
check("4 처리 뒤 고른 파일·상자 안 이름 지우기 — 📚·①·③ 셋 다(clearPicked 정의 1 + 부름 3)",
      js.count("function clearPicked(") == 1 and js.count("clearPicked(inp);") == 3 and "inp.value = ''" in js
      and ".dz-files" in js)

print()
print("=" * 78)
print("§2 화면")
print("=" * 78)
TMP = tempfile.mkdtemp(prefix="knk_z1191q_")
_REAL = D.DB_PATH
D.DB_PATH = os.path.join(TMP, "t.db")
assert D.DB_PATH != _REAL
with contextlib.redirect_stdout(io.StringIO()):
    D.init_db()
with D.db_session() as conn:
    have = {r[1] for r in conn.execute("PRAGMA table_info(users)")}
    for col, typ in [("employee_no", "TEXT"), ("entity", "TEXT"), ("name_vi", "TEXT"), ("rank", "TEXT"), ("team_id", "INTEGER"),
                     ("password_version", "INTEGER DEFAULT 1")]:
        if col not in have:
            conn.execute("ALTER TABLE users ADD COLUMN %s %s" % (col, typ))
    conn.execute("INSERT INTO users (name, role, employee_no, login_id, password, is_active, password_version, "
                 "can_view_sales, can_use_sales, can_view_logistics, can_use_logistics) VALUES ('가상M901','member','M901','M901','x',1,1,1,1,0,0)")
    U = dict(conn.execute("SELECT * FROM users WHERE employee_no='M901'").fetchone())
import app.main as M                                        # noqa: E402
from fastapi.testclient import TestClient                  # noqa: E402

M._HS_FILL_DIR = os.path.join(TMP, "hs_fill")
M.get_user = lambda request, _u=U: _u
cl = TestClient(M.app, follow_redirects=False)


def quiet(fn, *a, **k):
    with contextlib.redirect_stdout(io.StringIO()):
        return fn(*a, **k)


pg0 = quiet(cl.get, "/export/hs-fill").text
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "INVOICE"
ws["G4"], ws["H4"], ws["I4"] = "KNK-260101-1", "DATE:", "2026-01-01"
for col, v in {"D": "13) Description of Goods", "F": "14) Q'ty /Unit", "H": "15) Unit Price", "I": "16) Amount",
               "S": "HS CODE", "T": "Tên hàng"}.items():
    ws[col + "31"] = v
ws["D34"], ws["F34"], ws["H34"], ws["I34"], ws["S34"], ws["T34"] = "DCP32DN 15A", 2, 10.0, 20.0, "85362011", "Cầu dao"
b = io.BytesIO()
wb.save(b)
r = quiet(cl.post, "/export/hs-fill/hist", data={"kind": "수책통관"}, files={"file": ("a.xlsx", b.getvalue())})
pg1 = quiet(cl.get, "/export/hs-fill").text
check("5 화면 200 · 등록 전에는 목록 없음 · 등록 1개 뒤 목록이 펼친 채로 그 번호가 보임",
      "<summary>등록된 인보이스 목록" not in pg0 and r.status_code == 200
      and re.search(r'<details class="hf-det" open><summary>등록된 인보이스 목록 1개</summary>', pg1) is not None
      and "KNK-260101-1" in pg1, (r.status_code, re.findall(r"<details[^>]*><summary>등록된[^<]*", pg1)))

print()
print("결과: %d개 중 실패 %d" % (OK + FAIL, FAIL))
sys.exit(1 if FAIL else 0)
