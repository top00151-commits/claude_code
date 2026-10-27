# -*- coding: utf-8 -*-
"""z1191o 시험 — 인보이스 표준 양식 이름: 수책 양식(설비 부품) → 수책통관 · 소모품 양식 → 일반통관
   대표 지시 2026-10-27 「여기서 표준 약식 이름을 "수책통관","일반통관" 으로 변경하자..
                          기존에는 "수책양식(설비부품)","소모품양식" 으로 되어 있어..」
  §1 모듈(app/hs_fill): form_name — 새 이름 · 다른 값 그대로 · 여러 시트 · 빈 값 / 알아보는 값(scan form)은 그대로
  §2 경로(임시 DB · TestClient): 화면 단추·안내 글 새 이름(옛 이름 없음) · 받는 파일 이름 ·
     기록 표(옛 값 「수책」·「소모품」으로 저장된 줄도 새 이름) · ① 결과 forms 새 이름 + 저장 값 그대로 · 예전 양식 안내 글
"""
import contextlib
import io
import os
import re
import sys
import tempfile
from urllib.parse import quote as _q

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
os.environ["KNK_SSO_SERVICE_KEY"] = "z1191o-test-only-not-a-real-key"   # 시험용 가짜 값

import openpyxl                                            # noqa: E402
import app.database as D                                   # noqa: E402
from app import hs_fill as H                               # noqa: E402

OK = FAIL = 0


def check(label, cond, got=None):
    global OK, FAIL
    if cond:
        OK += 1
        print("  [OK] " + label)
    else:
        FAIL += 1
        print("  [FAIL] " + label + ("   ← %r" % (got,) if got is not None else ""))


SEED = os.path.join(ROOT, "app", "seed")


def filled(fn, ino, items):
    """표준 양식(seed)에 번호·선적일·품목 줄을 적은 파일 — 직원이 실제로 하는 그대로."""
    wb = openpyxl.load_workbook(os.path.join(SEED, fn))
    ws = wb["INVOICE"]
    ws["G4"], ws["D29"] = ino, "2026-10-30"
    for i, (d, q, pr) in enumerate(items):
        ws["D%d" % (34 + i)], ws["F%d" % (34 + i)], ws["H%d" % (34 + i)] = d, q, pr
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def old_form():
    """예전 양식 — 머리글만 맞는 간단한 인보이스(수책통관·일반통관 어느 쪽도 아님)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INVOICE"
    ws.append(["No", "Description", "Q'ty", "HS CODE", "Tên hàng"])
    ws.append([1, "DCP32DN 15A", 2, None, None])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


print("=" * 78)
print("§1 모듈")
print("=" * 78)
check("1 form_name: 수책 → 수책통관 · 소모품 → 일반통관",
      H.form_name("수책") == "수책통관" and H.form_name("소모품") == "일반통관", (H.form_name("수책"), H.form_name("소모품")))
check("2 form_name: 예전 양식 그대로 · 여러 시트 「수책·소모품」 → 「수책통관·일반통관」 · 빈 값·None → 빈 글자",
      H.form_name("예전 양식") == "예전 양식" and H.form_name("수책·소모품") == "수책통관·일반통관"
      and H.form_name("") == "" and H.form_name(None) == "",
      (H.form_name("예전 양식"), H.form_name("수책·소모품"), H.form_name(None)))
SUCHAEK = filled("invoice_form_suchaek.xlsx", "KNK-261101-1", [("DCP32DN 15A", 11, 15.16), ("NEW-PART-O1", 2, 3.0)])
SOMO = filled("invoice_form_somo.xlsx", "KNK-261101-2", [("DCP32DN 15A", 3, 15.16)])
OLD = old_form()
forms = {n: [s["form"] for s in H.scan(d)["sheets"]] for n, d in (("suchaek", SUCHAEK), ("somo", SOMO), ("old", OLD))}
check("3 알아보는 값은 그대로(수책 · 소모품 · 예전 양식) — 기록 표에 이미 저장된 값과 같아야 옛 줄도 새 이름으로 보인다",
      forms == {"suchaek": ["수책"], "somo": ["소모품"], "old": ["예전 양식"]}, forms)

print()
print("=" * 78)
print("§2 경로")
print("=" * 78)
TMP = tempfile.mkdtemp(prefix="knk_z1191o_")
_REAL = D.DB_PATH
D.DB_PATH = os.path.join(TMP, "t.db")
assert D.DB_PATH != _REAL
with contextlib.redirect_stdout(io.StringIO()):
    D.init_db()
with D.db_session() as conn:
    have = {r[1] for r in conn.execute("PRAGMA table_info(users)")}
    for col, typ in [("employee_no", "TEXT"), ("entity", "TEXT"), ("name_vi", "TEXT"), ("rank", "TEXT"), ("team_id", "INTEGER"),
                     ("password_version", "INTEGER DEFAULT 1")]:
        if col not in have:
            conn.execute("ALTER TABLE users ADD COLUMN %s %s" % (col, typ))
    conn.execute("INSERT INTO users (name, role, employee_no, login_id, password, is_active, password_version, "
                 "can_view_sales, can_use_sales, can_view_logistics, can_use_logistics) VALUES ('가상M901','member','M901','M901','x',1,1,1,1,0,0)")
    U = dict(conn.execute("SELECT * FROM users WHERE employee_no='M901'").fetchone())
import app.main as M                                        # noqa: E402
from fastapi.testclient import TestClient                  # noqa: E402

M._HS_FILL_DIR = os.path.join(TMP, "hs_fill")
M._HS_HIST_IMG_DIR = os.path.join(TMP, "hs_hist_img")       # z1191r: 회신·최종 파일의 품목 사진도 임시 폴더에(작업 폴더 data/ 에 남기지 않게)
M.get_user = lambda request, _u=U: _u
cl = TestClient(M.app, follow_redirects=False)


def quiet(fn, *a, **k):
    with contextlib.redirect_stdout(io.StringIO()):
        return fn(*a, **k)


with D.db_session() as c:
    M._hs_dict_ensure(c)
    c.execute("DELETE FROM export_hs_dict")
    c.execute("INSERT INTO export_hs_dict(part_key,model,hs_code,vn_name,status,seen_count) "
              "VALUES('DCP32DN15A','DCP32DN 15A','85362011','Cầu dao điện','CONFIRMED',2)")
    M._hs_fill_log_ensure(c)
    # 이름을 바꾸기 전에 저장된 기록 줄(운영에 이미 있는 값 그대로)
    for i, f in enumerate(("수책", "소모품", "수책·소모품", "예전 양식")):
        c.execute("INSERT INTO export_hs_fill_log(invoice_no,file_name,form,n_total,n_auto,n_check,n_need,n_keep,filled_by) "
                  "VALUES(?,?,?,1,1,0,0,0,?)", ("KNK-26090%d-1" % i, "before_%d.xlsx" % i, f, U["id"]))

pg = quiet(cl.get, "/export/hs-fill")
t = pg.text
check("4 화면 단추 「수책통관」·「일반통관」(주소 그대로) · 옛 이름 「수책 양식(설비 부품)」·「소모품 양식」 없음",
      pg.status_code == 200 and 'href="/export/hs-fill/form/suchaek">수책통관</a>' in t
      and 'href="/export/hs-fill/form/somo">일반통관</a>' in t and "수책 양식" not in t and "소모품 양식" not in t,
      pg.status_code)
check("5 ① 안내 글 「최종 양식(수책통관·일반통관)」", "최종 양식(<b>수책통관</b>·<b>일반통관</b>)" in t and "<b>소모품</b>" not in t)
before = dict(re.findall(r'before_(\d)\.xlsx<div class="hf-sub">([^<]*)</div>', t))
check("6 기록 표 — 바꾸기 전에 저장된 줄도 새 이름(수책 → 수책통관 · 소모품 → 일반통관 · 둘 다 · 예전 양식 그대로)",
      before == {"0": "수책통관", "1": "일반통관", "2": "수책통관·일반통관", "3": "예전 양식"}, before)

got = {k: quiet(cl.get, "/export/hs-fill/form/" + k) for k in ("suchaek", "somo")}
cd = {k: r.headers.get("content-disposition", "") for k, r in got.items()}
check("7 받는 파일 이름 「Invoice&packinglist_VN(수책통관)_표준양식.xlsx」·「…(일반통관)_표준양식.xlsx」 · 내용 = seed 그대로",
      all(r.status_code == 200 for r in got.values())
      and _q("Invoice&packinglist_VN(수책통관)_표준양식.xlsx") in cd["suchaek"]
      and _q("Invoice&packinglist_VN(일반통관)_표준양식.xlsx") in cd["somo"]
      and got["suchaek"].content == open(os.path.join(SEED, "invoice_form_suchaek.xlsx"), "rb").read()
      and got["somo"].content == open(os.path.join(SEED, "invoice_form_somo.xlsx"), "rb").read(), cd)

res = {}
for n, data, fname in (("suchaek", SUCHAEK, "Invoice&packinglist_VN(261101-1).xlsx"),
                       ("somo", SOMO, "Invoice&packinglist_VN(261101-2).xlsx"),
                       ("old", OLD, "invoice_old_261101.xlsx")):
    r = quiet(cl.post, "/export/hs-fill", files={"file": (fname, data)})
    res[n] = (r.status_code, r.json(), fname)
with D.db_session() as c:
    saved = {r["file_name"]: r["form"] for r in c.execute("SELECT file_name, form FROM export_hs_fill_log")}
check("8 ① 결과 — 수책통관 양식 · 일반통관 양식으로 알아봄(카드 머리 글) · 기록에는 알아보는 값 그대로(수책 · 소모품)",
      res["suchaek"][0] == 200 and res["suchaek"][1].get("forms") == ["수책통관"]
      and res["somo"][0] == 200 and res["somo"][1].get("forms") == ["일반통관"]
      and saved.get(res["suchaek"][2]) == "수책" and saved.get(res["somo"][2]) == "소모품",
      ({k: v[1].get("forms") for k, v in res.items()}, saved))
check("9 예전 양식 — forms 「예전 양식」 · 안내 글 「앞으로는 최종 양식(수책통관·일반통관)으로」",
      res["old"][0] == 200 and res["old"][1].get("forms") == ["예전 양식"]
      and "최종 양식(수책통관·일반통관)" in (res["old"][1].get("hint") or ""), (res["old"][1].get("forms"), res["old"][1].get("hint")))
t2 = quiet(cl.get, "/export/hs-fill").text
after = {fn: m for fn, m in re.findall(r'(Invoice&amp;packinglist_VN\(261101-\d\)\.xlsx)<div class="hf-sub">([^<]*)</div>', t2)}
check("10 ① 뒤 기록 표 — 새로 올린 두 줄도 수책통관 · 일반통관",
      after == {"Invoice&amp;packinglist_VN(261101-1).xlsx": "수책통관", "Invoice&amp;packinglist_VN(261101-2).xlsx": "일반통관"}, after)

print()
print("결과: %d개 중 실패 %d" % (OK + FAIL, FAIL))
sys.exit(1 if FAIL else 0)
