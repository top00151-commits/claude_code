# -*- coding: utf-8 -*-
"""z1191l 시험 — 베트남 인보이스 HS 자동 채움(엑셀 왕복)
   대표 지시 2026-10-26 「엑셀 양식은 지금 올려준 양식으로 최종 적용해 … 직원들이 가장 편하고 실수없이 빠르게」
  §1 모듈(app/hs_fill): 최종 양식 2종(수책 S·T 칸 / 소모품 R·S 칸 · 머리글 31줄) 칸 찾기 · 사전 대조 7갈래 ·
     채운 값 · 색(노랑 FFFF00 · 주황 FFC000) · 사진 바이트 그대로 · 바꾼 칸 말고 안 바뀜 · 인쇄 영역 그대로 · 수식 칸은 안 덮음
  §2 경로(임시 DB · TestClient): 권한 · 화면 · ① 자동 채움(기록·보관·내려받기) · ③ 회신 반영(사전 성장·확인 대기·회신 대기 해제)
     · 회신 전 파일 거절 · 사전 화면 업로드도 같은 규칙 · 세관 파싱 실패 길에서 NameError 안 남(log 이름)
  §3 원문: 사전 성장 코드는 한 곳(_hs_learn_invoice_rows) · 수출 hub·사전 화면 단추
"""
import contextlib
import io
import os
import re
import sys
import tempfile
import zipfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
os.environ["KNK_SSO_SERVICE_KEY"] = "z1191l-test-only-not-a-real-key"   # 시험용 가짜 값

import openpyxl                                            # noqa: E402
from openpyxl.styles import Border, Side                   # noqa: E402
import app.database as D                                   # noqa: E402
from app import hs_fill as H                               # noqa: E402

OK = FAIL = 0


def check(label, cond, got=None):
    global OK, FAIL
    if cond:
        OK += 1
        print("  [OK] " + label)
    else:
        FAIL += 1
        print("  [FAIL] " + label + ("   ← %r" % (got,) if got is not None else ""))


# ── 가짜 최종 양식 만들기(대표 10-26 양식과 같은 칸) ──
def make_form(kind, items, png=None):
    """kind='수책'(HS=S·Tên hàng=T) / '소모품'(HS=R·Tên hàng=S). items=[(규격, HS, 품명, HS칸 수식?)]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INVOICE"
    ws["A1"] = "COMMERCIAL  INVOICE"
    ws["G4"] = "KNK-261099-1"
    if kind == "수책":
        heads = {"A": "12) Marks and Number of PKGS", "D": "13) Description of Goods", "F": " 14) Q'ty /Unit", "H": "15) Unit Price",
                 "I": "16) Amount", "K": "COMPANY", "L": "관리코드", "M": "MODEL", "N": "품명", "O": "Q'TY", "P": "ORIGIN",
                 "Q": "SUPPLIER", "R": "CATEGORY", "S": "HS CODE", "T": "Tên hàng", "U": "SHIPPING DATE"}
        chs, cvn = "S", "T"
    else:
        heads = {"A": "12) Marks and Number of PKGS", "D": "13) Description of Goods", "F": " 14) Q'ty /Unit", "H": "15) Unit Price",
                 "I": "16) Amount", "K": "COMPANY", "L": "IMAGE", "M": "Origin", "N": "Manufacturer", "O": "invoice\n%", "P": "C/O",
                 "Q": "SIZE", "R": "HS CODE", "S": "Tên hàng"}
        chs, cvn = "R", "S"
    for col, v in heads.items():
        ws[col + "31"] = v
    thin = Side(style="thin", color="FF000000")
    r = 34
    for desc, hs, vn, formula in items:
        ws["D%d" % r] = desc
        ws["F%d" % r] = 2
        ws["H%d" % r] = 3.5
        ws["I%d" % r] = "=F%d*H%d" % (r, r)
        for col in (chs, cvn):
            ws["%s%d" % (col, r)].border = Border(left=thin, right=thin, top=thin, bottom=thin)
        if formula:
            ws["%s%d" % (chs, r)] = '=IF(1=1,"","")'
        elif hs is not None:
            ws["%s%d" % (chs, r)] = hs
        if vn is not None:
            ws["%s%d" % (cvn, r)] = vn
        r += 1
    ws["A%d" % (r + 1)] = "TOTAL"
    ws["I%d" % (r + 1)] = "=SUM(I33:I%d)" % r
    ws.print_area = "A1:I%d" % (r + 5)
    pl = wb.create_sheet("PACKING LIST")
    pl["A1"] = "PACKING LIST"
    pl["D31"] = "Description"
    pl["S31"] = "HS CODE"
    if png:
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(io.BytesIO(png))
        ws.add_image(img, "L34")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), chs, cvn


def tiny_png():
    from PIL import Image
    im = Image.new("RGB", (6, 5), (200, 30, 30))
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


DICT = [
    {"part_key": "DCP32DN15A", "model": "DCP32DN 15A", "hs_code": "85362011", "vn_name": "Cầu dao điện, mã hàng: DCP32DN 15A", "status": "CONFIRMED", "seen_count": 3},
    {"part_key": "SP100F120H6.5TTBM", "model": "SP100F120H6.5TTBM", "hs_code": "85389019", "vn_name": "Chân pin tiếp xúc, mã hàng: SP100F120H6.5TTBM", "status": "CONFIRMED", "seen_count": 2},
    {"part_key": "PINBLOCKSM-F971U-SUBTOP-R0.5C", "model": "PIN BLOCK SM-F971U-SUB TOP-R0.5C", "hs_code": "84669390", "vn_name": "x", "status": "CONFIRMED", "seen_count": 1},
    {"part_key": "FLOATINGSM-F971U-SUBTOP-R0.5C", "model": "FLOATING SM-F971U-SUB TOP-R0.5C", "hs_code": "84798939", "vn_name": "y", "status": "CONFIRMED", "seen_count": 1},
    {"part_key": "EN-AGS111107", "model": "EN-AGS111107", "hs_code": "39231090", "vn_name": "Hộp nhựa", "status": "PENDING", "seen_count": 2},
    {"part_key": "XFHT4060", "model": "XFHT4060", "hs_code": "73181590", "vn_name": "", "status": "CONFIRMED", "seen_count": 1},
    {"part_key": "G7SA-2A2B", "model": "G7SA-2A2B", "hs_code": "85364199", "vn_name": "Rơ le", "status": "CONFIRMED", "seen_count": 1},
]
ITEMS = [  # 규격 · HS · 품명 · HS칸 수식?
    ("DCP32DN 15A", None, None, False),                      # 34 사전 그대로 → 자동(HS+품명)
    ("PARTS OF CONNECTORS SP100F120H6.5TTBM", None, None, False),   # 35 모델만 일치 → 확인(주황) + 값
    ("SM-F971U-SUB TOP-R0.5C", None, None, False),           # 36 후보 HS 가 갈림 → 노랑(값 없음)
    ("EN-AGS111107", None, None, False),                     # 37 사전 확인 대기 → 노랑
    ("NEW-PART-99999", None, None, False),                   # 38 새 품목 → 노랑
    ("G7SA-2A2B", "85364990", "Rơ le an toàn", False),       # 39 본사가 적은 HS 가 사전과 다름 → 주황 · 값 그대로
    ("XFHT4060", None, None, False),                         # 40 사전 확정 · 품명 없음 → HS 자동 · 품명 노랑
    ("DCP32DN 15A", "85362011", None, False),                # 41 HS 같음 · 품명 비었음 → 품명만 자동
    ("DCP32DN 15A", None, None, True),                       # 42 HS 칸에 수식 → 값 안 씀
]
WANT = {34: "AUTO", 35: "CHECK", 36: "NEED", 37: "NEED", 38: "NEED", 39: "CONFLICT", 40: "AUTO", 41: "AUTO", 42: "CONFLICT"}


def fills(data, cols):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["INVOICE"]
    out = {}
    for r in range(34, 43):
        out[r] = tuple((ws["%s%d" % (c, r)].value, (ws["%s%d" % (c, r)].fill.fgColor.rgb if ws["%s%d" % (c, r)].fill and ws["%s%d" % (c, r)].fill.fill_type == "solid" else None)) for c in cols)
    pa = ws.print_area
    wb.close()
    return out, pa


print("=" * 78)
print("§1 모듈 — 최종 양식 2종")
print("=" * 78)
PNG = tiny_png()
for kind in ("수책", "소모품"):
    src, chs, cvn = make_form(kind, ITEMS, png=PNG)
    res = H.fill(src, DICT)
    sc = res["scan"]
    sh = sc["sheets"][0]
    check("1 %s: 칸 찾기 — 머리글 31줄 · HS=%s · Tên hàng=%s · 품목 %d줄(합계 줄 빠짐) · 양식 '%s' · 인보이스 번호" % (kind, chs, cvn, len(ITEMS), kind),
          sh["header_row"] == 31 and H._col_letter(sh["col_hs"]) == chs and H._col_letter(sh["col_vn"]) == cvn
          and len(sh["items"]) == len(ITEMS) and sh["form"] == kind and sc["invoice_no"] == "KNK-261099-1",
          (sh["header_row"], sh["col_hs"], sh["col_vn"], len(sh["items"]), sh["form"], sc["invoice_no"]))
    st = {a["row"]: a["status"] for a in res["acts"]}
    check("2 %s: 줄마다 판정 7갈래(자동·모델만=확인·후보 갈림·확인 대기·새 품목·사전과 다름·수식 칸)" % kind, st == WANT, st)
    got, pa = fills(res["out"], (chs, cvn))
    Y, O = "FFFFFF00", "FFFFC000"
    exp = {
        34: (("85362011", None), ("Cầu dao điện, mã hàng: DCP32DN 15A", None)),
        35: (("85389019", O), ("Chân pin tiếp xúc, mã hàng: SP100F120H6.5TTBM", O)),
        36: ((None, Y), (None, Y)),
        37: ((None, Y), (None, Y)),
        38: ((None, Y), (None, Y)),
        39: (("85364990", O), ("Rơ le an toàn", None)),
        40: (("73181590", None), (None, Y)),
        41: (("85362011", None), ("Cầu dao điện, mã hàng: DCP32DN 15A", None)),
    }
    bad = {r: (got[r], e) for r, e in exp.items() if got[r] != e}
    check("3 %s: 채운 값·색 — 사전 확정은 흰 칸 · 모델만은 주황 · 법인 입력은 노랑 · 본사가 적은 값은 안 덮음" % kind, not bad, bad)
    check("4 %s: 수식이 든 HS 칸은 값을 안 씀(수식 그대로)" % kind, got[42][0][0] == '=IF(1=1,"","")' and "%s42" % chs in res["skipped"], (got[42], res["skipped"]))
    zi, zo = zipfile.ZipFile(io.BytesIO(src)), zipfile.ZipFile(io.BytesIO(res["out"]))
    media = [n for n in zi.namelist() if n.startswith("xl/media/")]
    check("5 %s: 사진 바이트 그대로(%d장) · 압축 정상 · 바뀐 항목은 시트·서식·통합 문서뿐" % (kind, len(media)),
          media and all(zo.read(n) == zi.read(n) for n in media) and zo.testzip() is None
          and {n for n in zi.namelist() if zi.read(n) != zo.read(n)} <= {"xl/worksheets/sheet1.xml", "xl/styles.xml", "xl/workbook.xml"},
          [n for n in zi.namelist() if zi.read(n) != zo.read(n)])
    wi = openpyxl.load_workbook(io.BytesIO(src))
    wo = openpyxl.load_workbook(io.BytesIO(res["out"]))
    tgt = {(r, c) for r in range(34, 43) for c in (chs, cvn)}
    diff = []
    for name in wi.sheetnames:
        a, b = wi[name], wo[name]
        for row in a.iter_rows(min_row=1, max_row=max(a.max_row, b.max_row), max_col=max(a.max_column, b.max_column)):
            for cell in row:
                if name == "INVOICE" and (cell.row, cell.column_letter) in tgt:
                    continue
                if cell.value != b[cell.coordinate].value:
                    diff.append((name, cell.coordinate))
    check("6 %s: 바꾼 칸 말고 달라진 칸 0 · 인쇄 영역 그대로(A~I)" % kind, not diff and "A$1:$I$" in (pa or ""), (diff[:5], pa))
    wi.close()
    wo.close()

d_only = make_form("수책", [("DCP32DN 15A", "85362011", "x", False)])[0]
r0 = H.fill(d_only, DICT)
check("7 이미 다 적힌 파일 → 바꿀 것 없음(원본 그대로 돌려줌)", r0["out"] == d_only and r0["summary"]["KEEP"] == 1, r0["summary"])
bad_zip = H.fill(b"not a zip", DICT)
check("8 엑셀이 아니면 알기 쉬운 오류", bad_zip["ok"] is False and ".xlsx" in bad_zip["error"], bad_zip)
check("9 회신 파일 이름 앞부분: -re · _RE_DOAN · -LAN · (2) 떼기",
      H.file_stem("Invoice&packinglist_VN(261022)_RE_DOAN.xlsx") == H.file_stem("Invoice&packinglist_VN(261022).xlsx")
      == H.file_stem("Invoice&packinglist_VN(261022)-LAN (2).xlsx") == "invoice&packinglist_vn(261022)")

print()
print("=" * 78)
print("§2 경로 — 임시 DB · TestClient")
print("=" * 78)
TMP = tempfile.mkdtemp(prefix="knk_z1191l_")
_REAL = D.DB_PATH
D.DB_PATH = os.path.join(TMP, "t.db")
assert D.DB_PATH != _REAL
with contextlib.redirect_stdout(io.StringIO()):
    D.init_db()
with D.db_session() as conn:
    have = {r[1] for r in conn.execute("PRAGMA table_info(users)")}
    for col, typ in [("employee_no", "TEXT"), ("entity", "TEXT"), ("name_vi", "TEXT"), ("dept_code", "TEXT"), ("rank", "TEXT"),
                     ("team_id", "INTEGER"), ("password_version", "INTEGER DEFAULT 1")]:
        if col not in have:
            conn.execute("ALTER TABLE users ADD COLUMN %s %s" % (col, typ))
    if "entity" not in {r[1] for r in conn.execute("PRAGMA table_info(teams)")}:
        conn.execute("ALTER TABLE teams ADD COLUMN entity TEXT")
    conn.execute("INSERT INTO teams (name, entity, code, display_order) VALUES ('가상기술영업팀','KOR','ZLHQ',1)")
    TM = conn.execute("SELECT MAX(id) FROM teams").fetchone()[0]
    U = {}
    for sb, role, fl in (("L901", "member", (1, 1, 0, 0)), ("L902", "member", (0, 0, 0, 0)), ("L903", "admin", (0, 0, 0, 0))):
        conn.execute("INSERT INTO users (name, role, employee_no, login_id, password, is_active, password_version, team_id, rank, "
                     "can_view_sales, can_use_sales, can_view_logistics, can_use_logistics) VALUES (?,?,?,?,'x',1,1,?,'프로',?,?,?,?)",
                     ("가상" + sb, role, sb, sb, TM) + fl)
        U[sb] = dict(conn.execute("SELECT * FROM users WHERE employee_no=?", (sb,)).fetchone())

import app.main as M                                        # noqa: E402
from fastapi.testclient import TestClient                  # noqa: E402

M._HS_FILL_DIR = os.path.join(TMP, "hs_fill")              # 시험 파일은 임시 폴더에
cl = TestClient(M.app, follow_redirects=False)
with D.db_session() as c:
    M._hs_dict_ensure(c)
    c.execute("DELETE FROM export_hs_dict")                 # 시드 대신 시험 사전만
    for e in DICT:
        c.execute("INSERT INTO export_hs_dict(part_key,model,hs_code,vn_name,status,seen_count,source_file) VALUES(?,?,?,?,?,?,'시험')",
                  (e["part_key"], e["model"], e["hs_code"], e["vn_name"], e["status"], e["seen_count"]))


def as_user(sb):
    M.get_user = lambda request, _u=U[sb]: _u


def quiet(fn, *a, **k):
    with contextlib.redirect_stdout(io.StringIO()):
        return fn(*a, **k)


as_user("L902")
r1, r2 = quiet(cl.get, "/export/hs-fill"), quiet(cl.post, "/export/hs-fill", files={"file": ("a.xlsx", b"x")})
as_user("L903")
r3 = quiet(cl.get, "/export/hs-fill")
check("10 권한: 매출 등록 체크 없는 팀원·관리자 → 화면 303 /home · 올리기 403 (수출 권한 _export_guard 그대로)",
      r1.status_code == 303 and r1.headers.get("location") == "/home" and r2.status_code == 403 and r3.status_code == 303,
      (r1.status_code, r2.status_code, r3.status_code))
as_user("L901")
pg = quiet(cl.get, "/export/hs-fill")
check("11 화면 열림(200) · 단계 안내 · ①·③ 올리기 칸", pg.status_code == 200 and "인보이스 HS 자동 채움" in pg.text
      and 'id="hfFile"' in pg.text and 'id="hfReplyFiles"' in pg.text, pg.status_code)
draft, chs, cvn = make_form("소모품", ITEMS[:8], png=PNG)
rr = quiet(cl.post, "/export/hs-fill", files={"file": ("Invoice&packinglist_VN(261099).xlsx", draft)})
j = rr.json()
# z1191n(2026-10-27): 셈은 칸 색으로 — 40줄(사전 확정 · 품명은 사전에 없음 → HS 자동 · 품명 노랑)은 '법인 입력'으로 센다
check("12 ① 자동 채움: 요약(자동 2·확인 2·법인 입력 4 — z1191n: 품명만 노란 40줄도 법인 입력) · 인보이스 번호 · 이음 방 글",
      rr.status_code == 200 and j.get("ok")
      and j["summary"] == {"total": 8, "auto": 2, "check": 2, "need": 4, "keep": 0} and j["invoice_no"] == "KNK-261099-1"
      and "노란 칸 4줄" in j["message"] and "주황 칸 2줄" in j["message"], j.get("summary") or j)
with D.db_session() as c:
    lg = dict(c.execute("SELECT * FROM export_hs_fill_log ORDER BY id DESC LIMIT 1").fetchone())
check("13 자동 채움 기록 1줄(회신 대기 · 채운 사람 · 숫자)", lg["invoice_no"] == "KNK-261099-1" and lg["reply_at"] is None
      and lg["filled_by"] == U["L901"]["id"] and (lg["n_auto"], lg["n_check"], lg["n_need"]) == (2, 2, 4), lg)
dl = quiet(cl.get, j["download_url"])
dl_ok = dl.status_code == 200 and "attachment" in dl.headers.get("content-disposition", "")
filled = dl.content if dl_ok else b""
got, _ = fills(filled, (chs, cvn)) if dl_ok else ({}, None)
check("14 내려받기: 올린 이름 그대로 · 채운 값이 들어 있음", dl_ok and "Invoice%26packinglist_VN%28261099%29.xlsx" in dl.headers.get("content-disposition", "")
      and got.get(34, ((None,),))[0][0] == "85362011", (dl.status_code, dl.headers.get("content-disposition")))
pg2 = quiet(cl.get, "/export/hs-fill")
check("15 화면 기록: 회신 대기 1 · 다시 받기 단추", "회신 대기 1" in pg2.text and "/download" in pg2.text)

# 법인 회신 흉내 — 노란 칸 채움 + 모델만 일치 줄을 다른 값으로 고침(법인이 고친 경우)
wb = openpyxl.load_workbook(io.BytesIO(filled))
ws = wb["INVOICE"]
ws["%s36" % chs], ws["%s36" % cvn] = "84669390", "Khối chân pin"
ws["%s37" % chs], ws["%s37" % cvn] = "39269099", "Hộp nhựa"
ws["%s38" % chs], ws["%s38" % cvn] = "84799030", "Linh kiện mới"
ws["%s40" % cvn] = "Bu lông"
ws["%s35" % chs] = "85389099"                     # 사전 확정(85389019)과 다르게 — 확인 대기로 가야 함
buf = io.BytesIO()
wb.save(buf)
reply = buf.getvalue()
rp = quiet(cl.post, "/export/hs-fill/reply", files=[("files", ("Invoice&packinglist_VN(261099)_RE_DOAN.xlsx", reply))])
x = (rp.json().get("results") or [{}])[0]
check("16 ③ 회신 반영: 8줄 → 새 2 · 같은 값 3 · 다른 값 → 확인 대기 3(사전과 다름 2 + 같은 모델과 다름 1) · 빈 HS 0",
      rp.status_code == 200 and not x.get("error") and (x.get("rows"), x.get("added"), x.get("updated"), x.get("conflict"), x.get("empty")) == (8, 2, 3, 3, 0), x)
with D.db_session() as c:
    nw = c.execute("SELECT status, hs_code FROM export_hs_dict WHERE part_key='NEW-PART-99999'").fetchone()
    sp = c.execute("SELECT status, hs_code FROM export_hs_dict WHERE part_key='SP100F120H6.5TTBM'").fetchone()
    pc = c.execute("SELECT status, hs_code, note FROM export_hs_dict WHERE part_key='PARTSOFCONNECTORSSP100F120H6.5TTBM'").fetchone()
    lg2 = dict(c.execute("SELECT * FROM export_hs_fill_log WHERE id=?", (lg["id"],)).fetchone())
check("17 사전: 새 품목은 확정 추가 · 원래 확정 값(SP100F…=85389019)은 그대로",
      nw and nw["status"] == "CONFIRMED" and nw["hs_code"] == "84799030" and sp["status"] == "CONFIRMED" and sp["hs_code"] == "85389019",
      (dict(nw) if nw else None, dict(sp)))
check("17b 분류 글이 붙은 새 규격이 같은 모델과 HS 가 다르면 '확인 대기'로 들어감(말없이 갈라지지 않음)",
      pc and pc["status"] == "PENDING" and pc["hs_code"] == "85389099" and "같은 모델" in (pc["note"] or ""), dict(pc) if pc else None)
check("18 회신 대기 해제: 같은 인보이스 번호 기록 → 반영(회신 파일·사람·숫자)", lg2["reply_at"] and lg2["reply_file"].endswith("_RE_DOAN.xlsx")
      and lg2["reply_by"] == U["L901"]["id"] and x.get("log", {}).get("id") == lg["id"], lg2)
rp2 = quiet(cl.post, "/export/hs-fill/reply", files=[("files", ("draft.xlsx", draft))])
e2 = (rp2.json().get("results") or [{}])[0]
check("19 회신 전 파일(HS 빈 줄 6/8 — 본사가 2줄 미리 적은 초안)을 ③에 올리면 거절 — 알기 쉬운 말",
      "회신 전 파일" in (e2.get("error") or "") and "6/8" in (e2.get("error") or ""), e2)
rp3 = quiet(cl.post, "/export/hs-fill/reply", files=[("files", ("Invoice&packinglist_VN(261099).xlsx", filled))])
e3 = (rp3.json().get("results") or [{}])[0]
up3 = quiet(cl.post, "/export/hs-dict/upload", files=[("files", ("Invoice&packinglist_VN(261099).xlsx", filled))])
e4 = (up3.json().get("results") or [{}])[0]
check("19b ① 자동 채움 파일 그대로를 ③·사전 화면에 올리면 거절(주황 값이 확인 없이 확정되지 않게)",
      "그대로입니다" in (e3.get("error") or "") and "그대로입니다" in (e4.get("error") or ""), (e3, e4))
rx = quiet(cl.post, "/export/hs-fill", files={"file": ("a.xls", b"xx")})
check("20 .xls 는 거절(.xlsx 로 저장 안내)", rx.status_code == 400 and ".xlsx" in rx.json().get("error", ""), rx.text[:120])

# 사전 화면 업로드(기존 단추)도 같은 규칙 · 그 길에서 세관 파싱이 실패해도 멈추지 않음(log 이름)
quiet(cl.post, "/export/hs-fill", files={"file": ("Invoice&packinglist_VN(261099).xlsx", draft)})   # 두 번째 기록(회신 대기)
_orig = M._hs_parse_customs_xls


def _boom(raw, fname):
    raise ValueError("시험: 세관 파싱 실패")


M._hs_parse_customs_xls = _boom
try:
    up = quiet(cl.post, "/export/hs-dict/upload", files=[("files", ("Invoice&packinglist_VN(261099)-re.xlsx", reply))])
finally:
    M._hs_parse_customs_xls = _orig
u0 = (up.json().get("results") or [{}])[0] if up.status_code == 200 else {}
check("21 사전 화면 업로드: 세관 파싱 실패 길에서도 200(전에는 log 이름이 없어 NameError) · 같은 규칙(두 번째라 새 0)",
      up.status_code == 200 and u0.get("added") == 0 and u0.get("rows") == x.get("rows"), (up.status_code, u0))
check("22 사전 화면 업로드도 회신 대기를 품(두 번째 기록 → 반영)", (u0.get("log") or {}).get("id") not in (None, lg["id"]), u0.get("log"))

print()
print("=" * 78)
print("§3 원문")
print("=" * 78)
src = io.open(os.path.join(ROOT, "app", "main.py"), encoding="utf-8").read()
up_src = re.search(r'@app\.post\("/export/hs-dict/upload"\).*?(?=\n@app\.)', src, re.S).group(0)
check("23 사전 성장 코드는 한 곳 — 업로드·회신 반영 둘 다 _hs_learn_invoice_rows 를 부름",
      src.count("def _hs_learn_invoice_rows(") == 1 and "_hs_learn_invoice_rows(" in up_src and "INSERT INTO export_hs_dict" not in up_src
      and src.count("_hs_learn_invoice_rows(c, rows, fname, today, u)") == 2)
check("24 모듈에 log 이름 정의(아래 log.warning 들이 NameError 안 남)", re.search(r"^log = ", src, re.M) is not None and hasattr(M, "log"))
hub = io.open(os.path.join(ROOT, "app", "templates", "export_home.html"), encoding="utf-8").read()
dic = io.open(os.path.join(ROOT, "app", "templates", "export_hs_dict.html"), encoding="utf-8").read()
check("25 수출 hub 맨 앞 단추 · HS 사전 화면 안내 줄 → /export/hs-fill",
      hub.find('href="/export/hs-fill"') != -1 and hub.find('href="/export/hs-fill"') < hub.find('href="/export/prep"')
      and 'href="/export/hs-fill"' in dic)

print()
print("결과: %d개 중 실패 %d" % (OK + FAIL, FAIL))
sys.exit(1 if FAIL else 0)
