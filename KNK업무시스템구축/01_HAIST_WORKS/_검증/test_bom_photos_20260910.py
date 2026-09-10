#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""BOM 부품 사진 이어붙이기 — 시험 (대표 지시 2026-09-10)

대표 지시: "인벤터 엑셀에는 부품 사진이 있는데 우리 양식에서는 빠진다. 마지막 열에
사진도 넣어 달라. 영업팀이나 다른 팀들이 참고할 수 있게." → 대표 선택 (나) 통일판까지.

무엇을 지키는가
  ① 인벤터 사진이 간이판(Q) → 통일판(AB) 까지 이어진다
  ② 기존 칸(비고 P/AA)을 밀어내지 않는다 · 수식이 그대로다
  ③ ⭐개정으로 줄을 넣어도 사진이 제 부품에 붙어 있다
  ④ ⭐역검사: 사진 밀기를 빼면 시험이 실제로 실패한다(검사기가 헛돌지 않음을 증명)

실행:  python _검증/test_bom_photos_20260910.py
"""
import hashlib
import os
import shutil
import sys
import tempfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
os.chdir(ROOT)

from openpyxl import load_workbook, Workbook            # noqa: E402

import app.bom_tools.inventor_to_partlist as B          # noqa: E402
import app.bom_tools.merge_to_master as M               # noqa: E402
import app.bom_tools.revise_master as R                 # noqa: E402
import app.bom_tools.photos as PH                       # noqa: E402

_TAIL = ("참고자료", "설계팀", "BOM 업무 자동화_2026.08.04",
         "1. BUDS 단차 검사기", "INVENTOR DOWN")
_CANDS = [os.path.join(ROOT, "..", *_TAIL),
          os.path.join(ROOT, "..", "..", "..", "KNK업무시스템구축", *_TAIL)]
INV = next((os.path.abspath(c) for c in _CANDS if os.path.isdir(c)), None)

FAIL, CNT = [], 0


def check(name, ok, detail=""):
    global CNT
    CNT += 1
    print(("  OK  " if ok else "  FAIL") + " %02d %s" % (CNT, name)
          + ("" if ok else "  -- " + str(detail)))
    if not ok:
        FAIL.append(name)


def sig(b):
    return hashlib.md5(b).hexdigest()[:10]


def spec_photo_map(path, spec_col):
    """{형번: 사진서명} · 형번 빈 줄에 붙은 사진 수 · 못 읽은 수."""
    ws = load_workbook(path).active
    pics, bad = PH.read_photos(ws)
    out, blank = {}, 0
    for r, data in pics.items():
        spec = str(ws.cell(row=r, column=spec_col).value or "").strip()
        if spec:
            out.setdefault(spec, sig(data))
        else:
            blank += 1
    return out, blank, bad, ws


print("=" * 72)
print("  BOM 부품 사진 이어붙이기 - 시험")
print("=" * 72)

if not INV:
    print("  [중단] 인벤터 실물 폴더를 찾지 못했습니다.")
    sys.exit(1)

T = tempfile.mkdtemp(prefix="knk_bomphoto_")

# 1) 인벤터 원본
rows, excluded, per_file = B.read_inventor_files([INV])
n_pic = sum(1 for x in rows if x.get("사진"))
check("인벤터 구매품 줄을 읽는다", len(rows) == 60, "%d줄" % len(rows))
check("구매품 줄에 사진이 실린다", n_pic == 60, "%d/%d" % (n_pic, len(rows)))
check("파일별 통계에 사진 수가 남는다",
      all(len(t) == 5 for t in per_file) and sum(t[3] for t in per_file) == 60, str(per_file))
check("제외 집계에 사진 이야기를 섞지 않는다",
      not any("사진" in k for k in excluded), str(excluded))

# 2) 간이판
draft = os.path.join(T, "간이판.xlsx")
res = B.write_partlist(rows, "005M2606", "BUDS 단차 검사기", "한재운", draft)
dws = load_workbook(draft).active
check("간이판에 사진이 실린다", res["사진"] == 60 and res["사진없음"] == 0, str(res.get("사진")))
check("간이판 사진은 Q열(17)에 있다",
      {i.anchor._from.col + 1 for i in dws._images} == {17},
      str({i.anchor._from.col + 1 for i in dws._images}))
check("간이판 사진 머리글이 있다", dws.cell(row=7, column=17).value == PH.HEADER_TEXT,
      repr(dws.cell(row=7, column=17).value))
check("비고(P)를 밀어내지 않는다",
      "REMARK" in str(dws.cell(row=7, column=16).value or "").upper(),
      repr(dws.cell(row=7, column=16).value))
check("자료 칸이 그대로다(첫 줄 품명·형번)",
      str(dws.cell(row=8, column=5).value or "") != ""
      and str(dws.cell(row=8, column=6).value or "") != "",
      "%s / %s" % (dws.cell(row=8, column=5).value, dws.cell(row=8, column=6).value))
check("사진 때문에 표가 길어지지 않는다(줄 높이 그대로)",
      abs((dws.row_dimensions[8].height or 0) - 24.95) < 0.01,
      "%s (양식 기본 24.95)" % dws.row_dimensions[8].height)
check("⭐사진이 칸에 매여 있다 - 칸을 키우면 사진도 커진다",
      all(type(i.anchor).__name__ == "TwoCellAnchor" for i in dws._images)
      and all(i.anchor.to.row == i.anchor._from.row + 1
              and i.anchor.to.col == i.anchor._from.col + 1 for i in dws._images),
      str({type(i.anchor).__name__ for i in dws._images}))
check("사진 칸이 네모다(사진이 찌그러지지 않는다)",
      abs((dws.column_dimensions["Q"].width or 0) * 7.0 * 0.75
          - (dws.row_dimensions[8].height or 0)) < 6.0,
      "너비 %s · 높이 %s" % (dws.column_dimensions["Q"].width,
                            dws.row_dimensions[8].height))

# 3) 통일판
d_rows, rep = M.read_draft_files([draft])
check("간이판에서 사진을 다시 읽는다",
      sum(1 for x in d_rows if x.get("사진")) == 60 and not rep.get("사진못읽음"),
      str(rep.get("사진못읽음")))
master = os.path.join(T, "통일판.xlsx")
mres = M.write_master(d_rows, "005M2606", "BUDS 단차 검사기", "한재운", master, sets=2)
mws = load_workbook(master).active
check("통일판까지 사진이 이어진다", mres["사진"] == 60 and mres["사진없음"] == 0,
      str(mres.get("사진")))
check("통일판 사진은 AB열(28)에 있다",
      {i.anchor._from.col + 1 for i in mws._images} == {28},
      str({i.anchor._from.col + 1 for i in mws._images}))
check("통일판 사진도 칸에 매여 있다",
      all(type(i.anchor).__name__ == "TwoCellAnchor" for i in mws._images),
      str({type(i.anchor).__name__ for i in mws._images}))
check("통일판 비고(AA)를 밀어내지 않는다",
      "REMARK" in str(mws.cell(row=7, column=27).value or "").upper(),
      repr(mws.cell(row=7, column=27).value))
check("통일판 수식이 그대로다",
      mws["K9"].value == "=J9*$J$5" and mws["W9"].value == "=U9+V9",
      "%s / %s" % (mws["K9"].value, mws["W9"].value))
check("합계줄 SUBTOTAL 이 그대로다",
      str(mws.cell(row=mres["합계줄"], column=10).value or "").startswith("=SUBTOTAL"),
      str(mws.cell(row=mres["합계줄"], column=10).value))

# 4) 개정 - 사진이 제 부품에 남는가
before, blank0, bad0, _ = spec_photo_map(master, 6)
inv2 = os.path.join(T, "INV2")
shutil.copytree(INV, inv2)
src = os.path.join(inv2, "AC00.xlsx")
wb = load_workbook(src)
ws = wb.active
last = ws.max_row
for k in (1, 2):
    r = last + k
    for c in range(1, (ws.max_column or 0) + 1):
        h = str(ws.cell(row=1, column=c).value or "").strip().lower()
        if h.startswith("구분"):
            ws.cell(row=r, column=c, value="구매품")
        elif h == "주제":
            ws.cell(row=r, column=c, value="TEST BLOCK")
        elif h == "code":
            ws.cell(row=r, column=c, value="AC00")
        elif "description" in h:
            ws.cell(row=r, column=c, value="시험부품%d" % k)
        elif "부품번호" in h or "productcode" in h:
            ws.cell(row=r, column=c, value="NEW-SPEC-%d" % k)
        elif h == "수량" or h == "qty" or h.startswith("q"):
            ws.cell(row=r, column=c, value=1)
wb.save(src)

out2 = os.path.join(T, "개정본.xlsx")
rrep = R.revise(master, [inv2], out2)
after, blank1, bad1, _ = spec_photo_map(out2, 6)
mis = [k for k in before if after.get(k) != before[k]]
check("개정 뒤에도 사진이 제 부품에 붙어 있다", not mis,
      "어긋남 %d개 %s" % (len(mis), mis[:3]))
check("개정으로 사진이 사라지지 않는다", len(after) >= len(before),
      "%d -> %d" % (len(before), len(after)))
check("형번 빈 줄에 사진이 남지 않는다", blank1 == 0, "%d개" % blank1)
check("새로 추가된 줄에도 사진이 붙는다", rrep.get("사진추가", 0) > 0, str(rrep.get("사진추가")))

# 5) 역검사 - 사진 밀기를 빼면 정말 어긋나는가
_real = PH.shift_photos
try:
    PH.shift_photos = lambda ws, at_row, n: 0
    out3 = os.path.join(T, "역검사.xlsx")
    R.revise(master, [inv2], out3)
    broken, _b2, _b3, _ = spec_photo_map(out3, 6)
    mis2 = [k for k in before if broken.get(k) != before[k]]
    check("역검사: 사진 밀기를 빼면 짝이 어긋난다(시험이 헛돌지 않음)",
          len(mis2) > 0, "어긋남 %d개 -- 0이면 이 시험은 아무것도 안 잡는 것" % len(mis2))
finally:
    PH.shift_photos = _real

# 6) 사진 없는 파일 · 크기 지정
nopic = os.path.join(T, "사진없음")
os.makedirs(nopic, exist_ok=True)
wb2 = load_workbook(os.path.join(INV, "AD00.xlsx"))
ws2 = wb2.active
ws2._images = []
wb2.save(os.path.join(nopic, "AD00.xlsx"))
r2, _e2, pf2 = B.read_inventor_files([nopic])
d2 = os.path.join(T, "간이판2.xlsx")
res2 = B.write_partlist(r2, "005M2606", "BUDS", "한재운", d2)
check("사진이 없는 인벤터 파일도 그대로 동작한다",
      res2["품목"] == len(r2) and res2["사진"] == 0, str(res2.get("사진")))

d3 = os.path.join(T, "간이판_큰사진.xlsx")
B.write_partlist(rows, "005M2606", "BUDS", "한재운", d3, photo_pt=120)
w3 = load_workbook(d3).active
check("크게 쓰고 싶으면 지정할 수 있다(줄 높이가 따라 커짐)",
      (w3.row_dimensions[8].height or 0) > (dws.row_dimensions[8].height or 0),
      "%s vs %s" % (w3.row_dimensions[8].height, dws.row_dimensions[8].height))


# 7) 못 읽는 그림을 조용히 넘기지 않는가
class _Broken:
    ref = None
    anchor = None


ws4 = Workbook().active
ws4._images = [_Broken()]
got, bad = PH.read_photos(ws4)
check("못 읽는 그림은 수로 드러난다(조용히 넘기지 않음)", got == {} and bad == 1,
      "%s %s" % (got, bad))

print("-" * 72)
print("  시험 %d건 · 실패 %d건" % (CNT, len(FAIL))
      + ("" if not FAIL else " -> " + ", ".join(FAIL)))
sys.exit(1 if FAIL else 0)
