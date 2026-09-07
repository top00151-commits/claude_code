#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""WP-04 2026-09-07 결함 수리 회귀 시험 — 외부 검토로 드러난 8건

배경: 방향전환 문서(2026-09-07)에 대한 외부 검토가 결함을 지적했고, 전부 재현됐다.
      이 시험은 **그 결함이 되살아나면 실패**하도록 고정한다.

다루는 것 (라이브러리 계층)
  11-8-①  두 유닛 동시 추가 — 보고 줄번호가 실제와 일치 · 신규줄 K·O 수식이 채워짐
  11-8-②  옛 형번 2 → 새 형번 1 — 자동 확정하지 않고 후보로만 보고
  11-6    직전 개정에서 도구가 찍은 「추가」가 다음 설계에서 빠지면 삭제 표시됨
  11-7    개정하지 않은 유닛의 손 수정 수식은 모양 그대로 보존됨
  11-5    기존단가에 적힌 명시적 0 을 실적가로 덮지 않음
  11-3    수량·대수·재고 해석 실패를 1대분 수량으로 조용히 대체하지 않음
  음수     재고가 필요량보다 많아도 음수 발주수량을 내보내지 않음

다루지 않는 것 (웹 보고 계층 — test_bom_tools_routes.py 가 확인)
  11-5 타매입처 출처 행별 표시 · 11-8-③ 「신규만」 안내

실행: python _검증/test_bom_wp04_defects_20260907.py → "실패 0" 이어야 통과
"""
import os
import shutil
import sys
import tempfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from app.bom_tools import inventor_to_partlist as A   # noqa: E402
from app.bom_tools import merge_to_master as B        # noqa: E402
from app.bom_tools import revise_master as R          # noqa: E402
from app.bom_tools import make_po as P                # noqa: E402
from app.bom_tools import fill_prices as F            # noqa: E402
from openpyxl import load_workbook, Workbook          # noqa: E402

KIT = os.path.normpath(os.path.join(HERE, "..", "..", "참고자료", "설계팀",
                                    "BOM 업무 자동화_2026.08.04", "1. BUDS 단차 검사기", "INVENTOR DOWN"))
TMP = tempfile.mkdtemp(prefix="knk_wp04def_")
UNITS = ("AA00", "AB00", "AC00", "AD00")
G, Q, TP, CD, DS, SP, MK = 4, 5, 6, 7, 8, 9, 10       # 인벤터 열 (실물 실측)

FAIL, CNT = [], 0


def check(name, ok, detail=""):
    global CNT
    CNT += 1
    print(("  ✅" if ok else "  ❌") + f" {CNT:02d} {name}" + ("" if ok else f" — {detail}"))
    if not ok:
        FAIL.append(name)


def inv_copy(name, d):
    os.makedirs(d, exist_ok=True)
    dst = os.path.join(d, name + ".xlsx")
    shutil.copy(os.path.join(KIT, name + ".xlsx"), dst)
    return dst


def inv_add(path, topic, code, desc, spec, maker, qty=1):
    wb = load_workbook(path); ws = wb.active
    r = ws.max_row + 1
    for c, v in ((G, "구매품"), (Q, qty), (TP, topic), (CD, code),
                 (DS, desc), (SP, spec), (MK, maker)):
        ws.cell(r, c, v)
    wb.save(path)


def build_master(inv_paths, tag, sets=1):
    rows, _e, _p = A.read_inventor_files(inv_paths)
    draft = os.path.join(TMP, tag + "_간이판.xlsx")
    A.write_partlist(rows, "A005M2606", "BUDS 단차 검사기", "시험", draft)
    mrows, _r = B.read_draft_files([draft])
    out = os.path.join(TMP, tag + ".xlsx")
    B.write_master(mrows, "A005M2606", "BUDS 단차 검사기", "시험", out, sets=sets)
    return out


print("=" * 70)
print("  WP-04 결함 수리 회귀 (2026-09-07 외부 검토 지적 8건)")
print("=" * 70)

base = [inv_copy(u, os.path.join(TMP, "base")) for u in UNITS]
OLD = build_master(base, "옛마스터")

# ── 11-8-① 두 유닛 동시 추가 ──
d = os.path.join(TMP, "d1")
aa = inv_copy("AA00", d); ac = inv_copy("AC00", d)
inv_add(aa, "BOTTOM FRAME", "AA00", "NEW-AA-PART", "NEW-AA-001", "테스트메이커")
inv_add(ac, "SHUTTLE UNIT", "AC00", "NEW-AC-PART", "NEW-AC-001", "테스트메이커")
out = os.path.join(TMP, "d1_개정.xlsx")
rep = R.revise(OLD, [aa, ac], out)
ws = load_workbook(out).active
mismatch = [(s, r) for (r, _c, s) in rep["추가"] if str(ws.cell(r, 6).value or "") != s]
check("11-8-① 두 유닛 동시 추가: 보고 줄번호 = 실제 줄번호", not mismatch, str(mismatch))
noform = []
for spec in ("NEW-AA-001", "NEW-AC-001"):
    for rr in [x for x in range(9, ws.max_row + 1) if str(ws.cell(x, 6).value or "") == spec]:
        k, o = ws.cell(rr, 11).value, ws.cell(rr, 15).value
        if not (isinstance(k, str) and k.startswith("=")) or not (isinstance(o, str) and o.startswith("=")):
            noform.append((spec, rr, k, o))
check("11-8-① 신규줄에 K·O 수식이 채워짐", not noform, str(noform))

# ── 11-8-② 옛 2 → 새 1 은 자동 확정 금지 ──
d = os.path.join(TMP, "d2")
aa2 = inv_copy("AA00", d)
inv_add(aa2, "BOTTOM FRAME", "AA00", "DUP-PART", "OLD-1", "듀얼메이커")
inv_add(aa2, "BOTTOM FRAME", "AA00", "DUP-PART", "OLD-2", "듀얼메이커")
old2 = build_master([aa2] + [inv_copy(u, d) for u in ("AB00", "AC00", "AD00")], "옛마스터2")
d = os.path.join(TMP, "d2b")
aa2b = inv_copy("AA00", d)
inv_add(aa2b, "BOTTOM FRAME", "AA00", "DUP-PART", "NEW-ONLY", "듀얼메이커")
rep2 = R.revise(old2, [aa2b], os.path.join(TMP, "d2_개정.xlsx"))
auto = [x for x in rep2["형번개정"] if x[2] == "NEW-ONLY"]
check("11-8-② 옛 2 → 새 1: 자동 확정하지 않음", not auto, f"자동확정 {auto}")
check("11-8-② 대신 형번 후보로 보고", bool(rep2["형번후보"]), str(rep2["형번후보"]))

# ── 11-6 도구가 찍은 「추가」는 다음 개정에서 삭제 표시 ──
d = os.path.join(TMP, "d3")
aa3 = inv_copy("AA00", d)
inv_add(aa3, "BOTTOM FRAME", "AA00", "TEMP-PART", "TEMP-001", "임시메이커")
o1 = os.path.join(TMP, "d3_1차.xlsx")
R.revise(OLD, [aa3], o1)
o2 = os.path.join(TMP, "d3_2차.xlsx")
rep3 = R.revise(o1, [inv_copy("AA00", os.path.join(TMP, "d3b"))], o2)
ws3 = load_workbook(o2).active
st = [(r, ws3.cell(r, 2).value) for r in range(9, ws3.max_row + 1)
      if str(ws3.cell(r, 6).value or "") == "TEMP-001"]
check("11-6 직전 「추가」가 빠지면 삭제 표시됨",
      bool(st) and all(str(v or "") == "삭제" for _r, v in st), str(st))

# ── 11-7 개정 안 한 유닛의 손 수정 수식 보존 ──
old4 = os.path.join(TMP, "옛마스터4.xlsx"); shutil.copy(OLD, old4)
wb4 = load_workbook(old4); ws4 = wb4.active
tgt = None
for r in range(9, ws4.max_row + 1):
    if str(ws4.cell(r, 4).value or "").strip().upper() == "AB00":
        tgt = r; ws4.cell(r, 15, f"=$K{r}-$M{r}"); break
wb4.save(old4)
rep4 = R.revise(old4, [inv_copy("AA00", os.path.join(TMP, "d4"))], os.path.join(TMP, "d4_개정.xlsx"))
got = load_workbook(os.path.join(TMP, "d4_개정.xlsx")).active.cell(tgt, 15).value
check("11-7 개정 안 한 유닛의 손 수정 수식 보존", got == f"=$K{tgt}-$M{tgt}",
      f"'=$K{tgt}-$M{tgt}' → '{got}'")

# ── 11-5 명시적 0 보존 ──
led = os.path.join(TMP, "수불부.xlsx")
wb = Workbook(); ws = wb.active; ws.title = "발주관리대장"
ws.append(["규격/형번", "협력사", "발주일자", "단가"])
ws.append(["ZERO-SPEC", "가나상사", "2026-01-05", 12345])
wb.save(led)
m5 = os.path.join(TMP, "m5.xlsx"); shutil.copy(OLD, m5)
wb5 = load_workbook(m5); ws5 = wb5.active
ws5.cell(9, 5, "제로단가품"); ws5.cell(9, 6, "ZERO-SPEC")
ws5.cell(9, 8, "가나상사"); ws5.cell(9, 16, 0)
wb5.save(m5)
o5 = os.path.join(TMP, "m5_out.xlsx")
F.fill_prices(m5, led, o5)
check("11-5 기존단가의 명시적 0 을 덮지 않음",
      load_workbook(o5).active.cell(9, 16).value == 0,
      str(load_workbook(o5).active.cell(9, 16).value))

# ── 11-3 · 음수 ──
m6 = os.path.join(TMP, "m6.xlsx"); shutil.copy(OLD, m6)
wb6 = load_workbook(m6); ws6 = wb6.active
ws6["J5"] = 12
ws6.cell(9, 5, "문자재고품"); ws6.cell(9, 6, "TXT-STOCK"); ws6.cell(9, 8, "가나상사")
ws6.cell(9, 10, 10); ws6.cell(9, 13, "미확인")
ws6.cell(10, 5, "과잉재고품"); ws6.cell(10, 6, "OVER-STOCK"); ws6.cell(10, 8, "가나상사")
ws6.cell(10, 10, 10); ws6.cell(10, 13, 200); ws6.cell(10, 14, 0)
wb6.save(m6)
_c, _n, _s, rows6, _v = P.read_master(m6)
txt = next((x for x in rows6 if x["형번"] == "TXT-STOCK"), None)
over = next((x for x in rows6 if x["형번"] == "OVER-STOCK"), None)
check("11-3 재고 해석 실패를 1대분 수량으로 대체하지 않음",
      txt is not None and txt["수량"] is None and bool(txt["수량경고"]),
      f"수량={txt and txt['수량']!r} 경고={txt and txt['수량경고']!r}")
check("음수 발주수량을 내보내지 않음 (구매 필요 0 + 초과 보고)",
      over is not None and over["수량"] == 0 and "초과" in (over["수량경고"] or ""),
      f"수량={over and over['수량']!r} 경고={over and over['수량경고']!r}")

print("-" * 70)
print(f"  시험 {CNT}건 · 실패 {len(FAIL)}건" + ("" if not FAIL else " → " + ", ".join(FAIL)))
shutil.rmtree(TMP, ignore_errors=True)
sys.exit(1 if FAIL else 0)
