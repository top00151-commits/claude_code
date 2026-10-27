# -*- coding: utf-8 -*-
"""z1191r 시험 — 지난 수출 사진 등록·보기 · HS 사전을 최근 수출 값으로 확정
   대표 2026-10-27 「등록할때 이미지도 가능하면 같이 사진,단가 부분도 볼수있게 등록을해줘.. 그리고 지금 기준으로 AI추정 및
                    확인대기 이부분은 등록된 정보를 가장 최근걸로 입력을해줘.. 그럼 모두 등록확정이 될것 같은데..」
                  → (숫자·방법 보고) → 「12개도 최근 수출 값으로」(확정인데 HS 가 마지막 수출과 다른 것도)
  §1 모듈(app/hs_fill): _rels(상대·절대 경로 · 외부 링크 뺌) · row_images(한 칸·두 칸 그림 · 한 줄은 첫 장만 · 품목 밖 서명 뺌 ·
     예전 양식 B열) · compare 사진
  §2 경로(임시 DB · TestClient): 사진 줄이기(크기·흰 바탕·지문 이름·한 파일·깨진 그림 뺌) · 📚 등록 사진 ·
     이어 받기(같은 줄 · 줄이 밀림 · HS 사전 올림) · 사진 경로(200·404·403) · 인보이스 품목 · ① 회신 필요 없음 사진 ·
     규격으로 찾기(최근 순 · LIKE 글자 그대로) · ① 비교 사진 · ③ 회신 사진(여러 파일 · 엑셀 아님 · ① 받은 파일 그대로 거절) ·
     최근 수출 값으로 확정(대상 · 409 · 403 · 사유 · 품명 · 두 번째 0 · 다음 자동 채움) · 화면
  §3 원문: 화면 스크립트
"""
import contextlib
import datetime
import hashlib
import io
import os
import sys
import tempfile
import zipfile
from html.parser import HTMLParser

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
os.environ["KNK_SSO_SERVICE_KEY"] = "z1191r-test-only-not-a-real-key"   # 시험용 가짜 값

import openpyxl                                                          # noqa: E402
from openpyxl.drawing.image import Image as XImage                       # noqa: E402
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor  # noqa: E402
from PIL import Image as PI, ImageDraw                                   # noqa: E402
import app.database as D                                                 # noqa: E402
from app import hs_fill as H                                             # noqa: E402

OK = FAIL = 0


def check(label, cond, got=None):
    global OK, FAIL
    if cond:
        OK += 1
        print("  [OK] " + label)
    else:
        FAIL += 1
        print("  [FAIL] " + label + ("   ← %r" % (got,) if got is not None else ""))


def _save(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def png(color, size=(60, 40)):
    b = io.BytesIO()
    PI.new("RGB", size, color).save(b, "PNG")
    return b.getvalue()


def jpg(color, size=(60, 40)):
    b = io.BytesIO()
    PI.new("RGB", size, color).save(b, "JPEG", quality=92)
    return b.getvalue()


def big_rgba():
    """1200×800 — 바탕은 투명(빨강 0%) · 가운데 파랑 네모(불투명)."""
    im = PI.new("RGBA", (1200, 800), (255, 0, 0, 0))
    ImageDraw.Draw(im).rectangle((400, 300, 800, 500), fill=(0, 0, 255, 255))
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def sha(b):
    return hashlib.sha1(b).hexdigest() + ".jpg"


def two(c0, r0, c1, r1):
    """두 칸 그림 자리(0부터 세는 칸·줄)."""
    a = TwoCellAnchor()
    a._from = AnchorMarker(col=c0, row=r0)
    a.to = AnchorMarker(col=c1, row=r1)
    return a


def add_imgs(ws, imgs):
    for data, anchor in imgs:
        im = XImage(io.BytesIO(data))
        if isinstance(anchor, str):
            ws.add_image(im, anchor)
        else:
            im.anchor = anchor
            ws.add_image(im)


def _top(ws, ino, date):
    ws["A1"] = "COMMERCIAL  INVOICE"
    ws["F2"] = "8) No. & Date of Invoice"
    if ino:
        ws["G4"] = ino
    if date is not None:
        ws["H4"] = "DATE:"
        ws["I4"] = date


def make_final(kind, items, ino, date=None, imgs=()):
    """최종(표준) 양식(z1191p 시험과 같은 모양) + 그림 — kind '수책'(HS=S·Tên hàng=T) / '소모품'(HS=R·Tên hàng=S · IMAGE=L)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INVOICE"
    _top(ws, ino, date)
    if kind == "수책":
        heads = {"A": "12) Marks and Number of PKGS", "D": "13) Description of Goods", "F": " 14) Q'ty /Unit", "H": "15) Unit Price",
                 "I": "16) Amount", "K": "COMPANY", "P": "ORIGIN", "Q": "SUPPLIER", "R": "CATEGORY", "S": "HS CODE", "T": "Tên hàng"}
        chs, cvn = "S", "T"
    else:
        heads = {"A": "12) Marks and Number of PKGS", "D": "13) Description of Goods", "F": " 14) Q'ty /Unit", "H": "15) Unit Price",
                 "I": "16) Amount", "K": "COMPANY", "L": "IMAGE", "M": "Origin", "P": "C/O", "R": "HS CODE", "S": "Tên hàng"}
        chs, cvn = "R", "S"
    for col, v in heads.items():
        ws[col + "31"] = v
    r = 34
    for desc, q, p, hs, vn in items:
        ws["D%d" % r], ws["F%d" % r], ws["H%d" % r], ws["I%d" % r] = desc, q, p, "=F%d*H%d" % (r, r)
        if hs is not None:
            ws["%s%d" % (chs, r)] = hs
        if vn is not None:
            ws["%s%d" % (cvn, r)] = vn
        r += 1
    ws["A%d" % (r + 1)] = "TOTAL"
    add_imgs(ws, imgs)
    wb.create_sheet("PACKING LIST")["A1"] = "PACKING LIST"
    return _save(wb)


def make_old_ilban(items, ino, date, imgs=()):
    """예전 일반통관 양식(수출자료 폴더 실물 모양) — HS 머리글 32~33줄 · 그림은 B열."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INVOICE"
    _top(ws, ino, date)
    for col, v in {"A": "12) Marks and Number of PKGS", "D": "13) Description of Goods", "F": "14) Q'ty /Unit",
                   "H": "15) Unit Price", "I": "16) Amount"}.items():
        ws[col + "31"] = v
    ws["L32"], ws["Q32"], ws["S32"], ws["Q33"], ws["R33"] = "Origin", "HS CODE CHECK", "Duty", "HS code", "Tên hàng"
    r = 34
    for desc, q, p, hs, vn in items:
        ws["D%d" % r], ws["F%d" % r], ws["H%d" % r], ws["I%d" % r] = desc, q, p, round(q * p, 4)
        ws["Q%d" % r], ws["R%d" % r] = hs, vn
        r += 1
    ws["A%d" % (r + 1)] = "TOTAL"
    add_imgs(ws, imgs)
    return _save(wb)


IA, IB, IC, ID, IE = png((200, 30, 30)), jpg((30, 160, 40)), png((240, 200, 0)), png((10, 10, 10), (200, 60)), png((90, 0, 160))
ITEMS1 = [("IMG-PART-A1", 2, 10.0, "85365099", "Công tắc A"), ("IMG-PART-B2", 1, 4.0, "39269099", "Hộp B"),
          ("IMG-PART-C3", 5, 1.5, "84799030", "Chốt C"), ("IMG-PART-D4", 1, 2.0, "90314990", "Cảm biến D")]

print("=" * 78)
print("§1 모듈")
print("=" * 78)
X1 = make_final("소모품", ITEMS1, "KNK-260501-1", datetime.datetime(2026, 5, 1),
                imgs=[(IA, "L34"), (IB, two(11, 34, 12, 35)), (IC, "L34"), (ID, "A60")])
ri = H.row_images(X1, H.scan(X1))
check("1 row_images — 34줄 한 칸 그림(같은 줄 둘째 그림은 뺌 · 첫 장만) · 35~36줄에 걸친 두 칸 그림 → 35줄 · 60줄 서명(품목 밖) 뺌",
      sorted(ri) == [("INVOICE", 34), ("INVOICE", 35)] and ri[("INVOICE", 34)] == ("png", IA)
      and ri[("INVOICE", 35)][1] == IB and ri[("INVOICE", 35)][0] in ("jpeg", "jpg"),
      {k: (v[0], len(v[1])) for k, v in ri.items()})
zb = io.BytesIO()
with zipfile.ZipFile(zb, "w") as z:
    z.writestr("xl/drawings/_rels/drawing1.xml.rels",
               '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
               '<Relationship Id="rId1" Type="http://x/image" Target="../media/image1.png"/>'
               '<Relationship Id="rId2" Type="http://x/image" Target="http://example.invalid/x.png" TargetMode="External"/>'
               '<Relationship Id="rId3" Type="http://x/image" Target="/xl/media/image3.png"/></Relationships>')
rl = H._rels(zipfile.ZipFile(io.BytesIO(zb.getvalue())), "xl/drawings/drawing1.xml")
check("2 _rels — 상대 경로(../media) 풀기 · 절대 경로(/xl/…) · 외부 링크(TargetMode=External)는 뺌",
      rl == {"rId1": "xl/media/image1.png", "rId3": "xl/media/image3.png"}, rl)
X0 = make_old_ilban([("OLD-PART-Q1", 3, 2.0, "85389019", "Chân pin")], "KNK-260415-1", datetime.datetime(2026, 4, 15), imgs=[(IC, "B34")])
r0 = H.row_images(X0, H.scan(X0))
check("3 예전 일반통관 양식 — 그림이 B열이어도 줄로 잇는다(칸은 안 봄)",
      list(r0) == [("INVOICE", 34)] and r0[("INVOICE", 34)][1] == IC, list(r0))
acts = [{"sheet": "INVOICE", "row": 34, "key": "AAA-11111", "desc": "AAA-11111", "write_hs": "85362011", "hs_written": ""}]
past = {"AAA-11111": [{"id": 1, "invoice_no": "KNK-260101-1", "inv_date": "2026-01-01", "hs": "85362012", "price": 10.0, "img": "a.jpg"},
                      {"id": 2, "invoice_no": "KNK-260301-1", "inv_date": "2026-03-01", "hs": "85362012", "price": 10.0, "img": ""}]}
cm = H.compare(acts, {("INVOICE", 34): 20.0}, past, "")
check("4 compare — 그 규격의 마지막 사진(사진 없는 마지막 줄은 건너뛰고 그 앞 사진) · 단가·HS 항목에도 img",
      cm["img"] == {("INVOICE", 34): "a.jpg"} and cm["price"][0]["img"] == "a.jpg" and cm["hs"][0]["img"] == "a.jpg",
      (cm["img"], cm["price"], cm["hs"]))

print()
print("=" * 78)
print("§2 경로")
print("=" * 78)
TMP = tempfile.mkdtemp(prefix="knk_z1191r_")
_REAL = D.DB_PATH
D.DB_PATH = os.path.join(TMP, "t.db")
assert D.DB_PATH != _REAL
with contextlib.redirect_stdout(io.StringIO()):
    D.init_db()
with D.db_session() as conn:
    have = {r[1] for r in conn.execute("PRAGMA table_info(users)")}
    for col, typ in [("employee_no", "TEXT"), ("entity", "TEXT"), ("name_vi", "TEXT"), ("rank", "TEXT"), ("team_id", "INTEGER"),
                     ("password_version", "INTEGER DEFAULT 1")]:
        if col not in have:
            conn.execute("ALTER TABLE users ADD COLUMN %s %s" % (col, typ))
    U = {}
    for sb, fl in (("M901", (1, 1, 0, 0)), ("M902", (0, 0, 0, 0))):
        conn.execute("INSERT INTO users (name, role, employee_no, login_id, password, is_active, password_version, "
                     "can_view_sales, can_use_sales, can_view_logistics, can_use_logistics) VALUES (?,?,?,?,'x',1,1,?,?,?,?)",
                     ("가상" + sb, "member", sb, sb) + fl)
        U[sb] = dict(conn.execute("SELECT * FROM users WHERE employee_no=?", (sb,)).fetchone())
import app.main as M                                        # noqa: E402
from fastapi.testclient import TestClient                  # noqa: E402

M._HS_FILL_DIR = os.path.join(TMP, "hs_fill")
M._HS_HIST_IMG_DIR = os.path.join(TMP, "hs_hist_img")       # ⚠ 작업 폴더 data/ 에 시험 사진을 남기지 않게
cl = TestClient(M.app, follow_redirects=False)
IMGDIR = M._HS_HIST_IMG_DIR


def as_user(sb):
    M.get_user = lambda request, _u=U[sb]: _u


def quiet(fn, *a, **k):
    with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
        return fn(*a, **k)


def reg(kind, data, name):
    r = quiet(cl.post, "/export/hs-fill/hist", data={"kind": kind}, files={"file": (name, data)})
    return r.status_code, r.json()


def lines(ino):
    with D.db_session() as c:
        return [(r["row_no"], r["part_key"], r["img"]) for r in c.execute(
            "SELECT row_no, part_key, img FROM export_hist_line WHERE invoice_no=? ORDER BY id", (ino,))]


def drow(k):
    with D.db_session() as c:
        r = c.execute("SELECT * FROM export_hs_dict WHERE part_key=?", (k,)).fetchone()
        return dict(r) if r else None


def has_img(name):
    return os.path.exists(os.path.join(IMGDIR, name))


as_user("M901")
with D.db_session() as c:
    M._hs_dict_ensure(c)
    c.execute("DELETE FROM export_hs_dict")

t0 = M._hs_hist_thumbs({})
check("5 사진이 없으면 빈 결과 · 사진 폴더도 만들지 않음", t0 == {} and not os.path.exists(IMGDIR))
BIG = big_rgba()
tt = quiet(M._hs_hist_thumbs, {("S", 1): ("png", BIG), ("S", 2): ("png", BIG), ("S", 3): ("png", b"not an image")})
im = PI.open(os.path.join(IMGDIR, sha(BIG))) if has_img(sha(BIG)) else None
check("6 사진 줄이기 — 긴 변 480(1200×800 → 480×320) · JPEG · 투명 자리는 흰 바탕 · 이름 = 원본 지문 · 같은 그림 두 줄 = 한 파일 · 깨진 그림은 뺌",
      tt == {("S", 1): sha(BIG), ("S", 2): sha(BIG)} and im is not None and im.format == "JPEG" and im.size == (480, 320)
      and min(im.getpixel((5, 5))) >= 240 and im.getpixel((240, 160))[2] >= 200 and im.getpixel((240, 160))[0] <= 60
      and sorted(os.listdir(IMGDIR)) == [sha(BIG)],
      (tt, im and (im.format, im.size, im.getpixel((5, 5)), im.getpixel((240, 160))), os.listdir(IMGDIR) if os.path.exists(IMGDIR) else None))

c1, j1 = reg("일반통관", X1, "Invoice&packinglist_VN(260501).xlsx")
L1 = lines("KNK-260501-1")
check("7 📚 등록 — 사진 2줄(34 A · 35 B) · 나머지 줄은 사진 없음 · 원본↔저장 대조 같음 · 둘째 그림·서명은 안 넣음",
      c1 == 200 and j1.get("imgs") == 2 and (j1.get("check") or {}).get("ok")
      and [(r, i) for r, k, i in L1] == [(34, sha(IA)), (35, sha(IB)), (36, None), (37, None)]
      and has_img(sha(IA)) and has_img(sha(IB)) and not has_img(sha(IC)) and not has_img(sha(ID)),
      (c1, j1.get("imgs"), j1.get("error"), L1))

c2, j2 = reg("일반통관", make_final("소모품", ITEMS1, "KNK-260501-1", datetime.datetime(2026, 5, 1)), "Invoice&packinglist_VN(260501)_b.xlsx")
L2 = lines("KNK-260501-1")
check("8 이어 받기 — 같은 번호를 사진 없는 파일로 다시 올려도(바꿔 넣음) 같은 줄·규격의 사진은 그대로",
      c2 == 200 and j2.get("replaced") and j2.get("imgs") == 2 and L2 == L1, (j2.get("imgs"), L2))
X1c = make_final("소모품", [("IMG-NEW-TOP0", 1, 1.0, "85365099", "Mới")] + ITEMS1, "KNK-260501-1", datetime.datetime(2026, 5, 1))
c3, j3 = reg("일반통관", X1c, "Invoice&packinglist_VN(260501)_c.xlsx")
L3 = lines("KNK-260501-1")
check("9 이어 받기 — 맨 위에 줄이 끼어 번호가 밀려도 같은 규격이면 그 사진(35 A · 36 B) · 새 줄은 사진 없음",
      c3 == 200 and j3.get("imgs") == 2 and L3[:3] == [(34, "IMG-NEW-TOP0", None), (35, "IMG-PART-A1", sha(IA)), (36, "IMG-PART-B2", sha(IB))],
      L3)
with D.db_session() as c:
    jd = M._hs_hist_from_raw(c, X1c, "dict.xlsx", "dict", U["M901"])
L4 = lines("KNK-260501-1")
check("10 HS 사전 올림(사진을 안 읽는 길)으로 같은 번호가 다시 들어와도 사진은 그대로",
      jd.get("imgs") == 2 and L4 == L3, (jd, L4))

ra = quiet(cl.get, "/export/hs-fill/hist-img/" + sha(IA))
bad = [quiet(cl.get, "/export/hs-fill/hist-img/" + x).status_code
       for x in ("0" * 40 + ".jpg", sha(IA)[:-4].upper() + ".jpg", sha(IA)[:-4] + ".png", "abc.jpg", "..jpg", "0" * 41 + ".jpg")]
as_user("M902")
r403 = quiet(cl.get, "/export/hs-fill/hist-img/" + sha(IA)).status_code
as_user("M901")
check("11 사진 경로 — 있는 사진 200 · image/jpeg · 개인 캐시 · 없는 이름·대문자·다른 확장자·짧은·긴 이름·점점 404 · 수출 권한 없음 403",
      ra.status_code == 200 and ra.headers.get("content-type") == "image/jpeg" and ra.headers.get("cache-control", "").startswith("private")
      and ra.content[:2] == b"\xff\xd8" and bad == [404] * 6 and r403 == 403,
      (ra.status_code, ra.headers.get("content-type"), ra.headers.get("cache-control"), bad, r403))

ri1 = quiet(cl.get, "/export/hs-fill/hist/inv", params={"no": "KNK-260501-1"})
jl = ri1.json()
la = next((x for x in jl.get("lines", []) if x["descr"] == "IMG-PART-A1"), {})
ri404 = quiet(cl.get, "/export/hs-fill/hist/inv", params={"no": "KNK-269999-9"}).status_code
as_user("M902")
ri403 = quiet(cl.get, "/export/hs-fill/hist/inv", params={"no": "KNK-260501-1"}).status_code
as_user("M901")
check("12 인보이스 품목 — 5줄 · 줄마다 사진·규격·수량·단가·금액(수식이면 수량×단가)·HS·품명 · 없는 번호 404 · 권한 없음 403",
      ri1.status_code == 200 and jl.get("ok") and len(jl.get("lines", [])) == 5 and jl["invoice"]["invoice_no"] == "KNK-260501-1"
      and (la.get("qty"), la.get("unit_price"), la.get("amount"), la.get("hs_code"), la.get("vn_name"), la.get("img"), la.get("row_no"))
      == (2.0, 10.0, 20.0, "85365099", "Công tắc A", sha(IA), 35) and ri404 == 404 and ri403 == 403,
      (ri1.status_code, la, ri404, ri403))

XN = make_final("소모품", [("IMG-PART-A1", 4, 10.0, None, None), ("IMG-PART-B2", 2, 4.0, None, None)], "KNK-260601-1",
                datetime.datetime(2026, 6, 1), imgs=[(IA, "L34"), (IB, "L35")])
rn = quiet(cl.post, "/export/hs-fill", files={"file": ("Invoice&packinglist_VN(260601).xlsx", XN)}).json()
check("13 ① 회신 필요 없음(사전 확정 값으로 다 채움) → 받은 파일로 이력 · 사진 2줄도",
      rn.get("no_reply") and (rn.get("hist_saved") or {}).get("imgs") == 2
      and lines("KNK-260601-1") == [(34, "IMG-PART-A1", sha(IA)), (35, "IMG-PART-B2", sha(IB))],
      (rn.get("no_reply"), rn.get("hist_saved"), rn.get("error"), lines("KNK-260601-1")))

reg("일반통관", make_final("소모품", [("AB_CD-1", 1, 3.0, "85365099", "x"), ("ABXCD-1", 1, 3.0, "85365099", "y")],
                         "KNK-260502-1", datetime.datetime(2026, 5, 2)), "Invoice&packinglist_VN(260502).xlsx")


def find(q):
    r = quiet(cl.get, "/export/hs-fill/hist/find", params={"q": q})
    return r.status_code, r.json()


f1 = find(" img-part-a1 ")          # 규격 키 = 공백 빼고 대문자(하이픈은 그대로)
f2 = find("AB_CD")
f3 = find("%%")
f4 = find(" a ")
as_user("M902")
f5 = find("IMG-PART-A1")
as_user("M901")
check("14 규격으로 찾기 — 공백·대소문자 무시 · 최근 순(260601 → 260501) · 단가·사진 함께 · more 없음",
      f1[0] == 200 and [(x["invoice_no"], x["unit_price"], x["img"]) for x in f1[1]["lines"]]
      == [("KNK-260601-1", 10.0, sha(IA)), ("KNK-260501-1", 10.0, sha(IA))] and f1[1]["more"] is False,
      f1)
check("15 찾기 — '_'·'%' 는 글자 그대로(AB_CD → AB_CD-1 만 · ABXCD-1 안 걸림 · '%%' → 0줄) · 한 글자 400 · 권한 없음 403",
      [x["part_key"] for x in f2[1]["lines"]] == ["AB_CD-1"] and f3[0] == 200 and f3[1]["lines"] == []
      and f4[0] == 400 and f5[0] == 403, (f2, f3[1].get("lines"), f4[0], f5[0]))

with D.db_session() as c:
    c.execute("UPDATE export_hs_dict SET hs_code='85365010' WHERE part_key='IMG-PART-A1'")
    c.execute("DELETE FROM export_hs_dict WHERE part_key='IMG-PART-B2'")
XC = make_final("소모품", [("IMG-PART-A1", 1, 15.0, None, None), ("IMG-PART-B2", 1, 4.0, None, None)], "KNK-260701-1",
                datetime.datetime(2026, 7, 1))
rc = quiet(cl.post, "/export/hs-fill", files={"file": ("Invoice&packinglist_VN(260701).xlsx", XC)}).json()
hi = rc.get("hist") or {}
attB = next((a for a in rc.get("attention", []) if a["desc"] == "IMG-PART-B2"), {})
check("16 ① 비교 — 단가 15 ↔ 마지막 10(+50%) 줄 · HS(사전 85365010 ↔ 마지막 수출 85365099) 줄에 지난 사진 · 노란 줄(사전에 없음) 근거에 사진 · 이음 글에 단가 없음",
      [(p["desc"], p.get("img")) for p in hi.get("price", [])] == [("IMG-PART-A1", sha(IA))]
      and [(x["desc"], x.get("img")) for x in hi.get("hs", [])] == [("IMG-PART-A1", sha(IA))]
      and attB.get("hist_img") == sha(IB) and attB.get("st") == H.NEED and "15" not in rc.get("message", ""),
      (hi.get("price"), hi.get("hs"), attB))

XR = make_final("소모품", [("IMG-PART-E5", 3, 7.0, "85369099", "Đầu nối E")], "KNK-260801-1", datetime.datetime(2026, 8, 1),
                imgs=[(IE, "L34")])
own = quiet(cl.get, "/export/hs-fill/%d/download" % rn["log_id"]).content
rr = quiet(cl.post, "/export/hs-fill/reply", files=[("files", ("reply_260801.xlsx", XR)), ("files", ("note.txt", b"x")),
                                                   ("files", ("own.xlsx", own))]).json()
res = rr.get("results", [])
check("17 ③ 회신 여러 개 — 차례 그대로 · 회신 이력에 사진 1줄 · 엑셀 아님 · ① 받은 파일 그대로는 거절(느린 일을 DB 밖으로 옮긴 뒤에도)",
      rr.get("ok") and [x["file"] for x in res] == ["reply_260801.xlsx", "note.txt", "own.xlsx"]
      and (res[0].get("hist") or {}).get("imgs") == 1 and lines("KNK-260801-1") == [(34, "IMG-PART-E5", sha(IE))]
      and res[0].get("added") == 1 and res[1].get("error") == "엑셀(.xlsx)만 됩니다" and "그대로입니다" in (res[2].get("error") or ""),
      res)

# ─── 최근 수출 값으로 확정 — 이력·사전을 비우고 정해 둔 줄로 ───
IMGN = sha(IE)
with D.db_session() as c:
    for t in ("export_hs_dict", "export_hist_line", "export_hist_inv"):
        c.execute("DELETE FROM %s" % t)
    for k, hs, st, vn, note in (
            ("K-AI-SAME", "85365099", "AI", "Công tắc", ""),
            ("K-PEND-DIFF", "39269099", "PENDING", "Nhựa", "지난 수출 등록 · 이력끼리 HS 갈림: 39269099×1 · 84799030×1 (마지막 KNK-260301-9)"),
            ("K-CONF-DIFF", "8536.50.99", "CONFIRMED", "Rơ le", ""),
            ("K-CONF-SAME", "90314990", "CONFIRMED", "Cảm biến", ""),
            ("K-CONF-DOTSAME", "9031.49.90", "CONFIRMED", "Cảm biến 2", ""),
            ("K-AI-NOHIST", "85389019", "AI", "Chân", ""),
            ("K-PEND-NOHIST", "85369099", "PENDING", "Đầu nối", ""),
            ("K-AI-EMPTYVN", "85365099", "AI", "", ""),
            ("K-HIST-EMPTYHS", "85365099", "AI", "x", ""),
            ("K-BY-DATE", "11111111", "PENDING", "y", "")):
        c.execute("INSERT INTO export_hs_dict(part_key,model,vn_name,hs_code,status,seen_count,note,updated_at) "
                  "VALUES(?,?,?,?,?,1,?,'2026-07-23 10:00:00')", (k, k, vn, hs, st, note))
    for ino, d, k, hs, vn, img in (
            ("KNK-260105-9", "2026-01-05", "K-AI-SAME", "85365099", "Công tắc A", None),
            ("KNK-260201-9", "2026-02-01", "K-PEND-DIFF", "39269099", "Nhựa", None),
            ("KNK-260301-9", "2026-03-01", "K-PEND-DIFF", "84799030", "Chốt#&VN", IMGN),
            ("KNK-260401-9", "2026-04-01", "K-CONF-DIFF", "85364990", "Rơ le điện", None),
            ("KNK-260401-9", "2026-04-01", "K-CONF-SAME", "90314990", "Cảm biến", None),
            ("KNK-260401-9", "2026-04-01", "K-CONF-DOTSAME", "90314990", "Cảm biến 2", None),
            ("KNK-260401-9", "2026-04-01", "K-AI-EMPTYVN", "85365099", "Công tắc B", None),
            ("KNK-260401-9", "2026-04-01", "K-HIST-EMPTYHS", "", "", None),
            ("KNK-260501-9", "2026-05-01", "K-BY-DATE", "22222222", "y2", None),
            ("KNK-260301-9", "2026-03-01", "K-BY-DATE", "33333333", "y3", None)):   # 나중에 넣었지만(번호 큼) 날짜는 앞 → 마지막 아님
        c.execute("INSERT INTO export_hist_line(invoice_no,inv_date,kind,sheet,row_no,part_key,descr,qty,unit_price,amount,hs_code,vn_name,img) "
                  "VALUES(?,?,'일반통관','INVOICE',34,?,?,1,1,1,?,?,?)", (ino, d, k, k, hs, vn, img))
    for ino, d in (("KNK-260105-9", "2026-01-05"), ("KNK-260201-9", "2026-02-01"), ("KNK-260301-9", "2026-03-01"),
                   ("KNK-260401-9", "2026-04-01"), ("KNK-260501-9", "2026-05-01")):
        c.execute("INSERT INTO export_hist_inv(invoice_no,inv_date,kind,file_name,source,n_lines,total_amount,reg_by) "
                  "VALUES(?,?,'일반통관','x.xlsx','register',1,1,?)", (ino, d, U["M901"]["id"]))
    today = c.execute("SELECT date('now','localtime')").fetchone()[0]
    sm = M._hs_hist_summary(c)
lt = {x["part_key"]: x for x in sm["latest"]}
check("18 대상 — 이력 있는 사전 줄 중 AI 추정·확인 대기 + 확정인데 HS 다름 = 5(HS 바뀜 3 · 확정만 2) · 확정·같은 HS(점 섞여도) 빠짐 · "
      "이력 없는 AI 2(빈 HS 이력 포함)·확인 대기 1 은 남음 · 마지막 = 날짜 순(번호 순 아님)",
      sorted(lt) == ["K-AI-EMPTYVN", "K-AI-SAME", "K-BY-DATE", "K-CONF-DIFF", "K-PEND-DIFF"]
      and (sm["n_latest"], sm["n_latest_hs"], sm["n_latest_st"], sm["rest_ai"], sm["rest_pending"]) == (5, 3, 2, 2, 1)
      and (lt["K-BY-DATE"]["hs"], lt["K-BY-DATE"]["invoice_no"], lt["K-BY-DATE"]["n"]) == ("22222222", "KNK-260501-9", 2)
      and lt["K-PEND-DIFF"]["img"] == IMGN and lt["K-CONF-DIFF"]["hs_change"] and not lt["K-AI-SAME"]["hs_change"]
      and (lt["K-AI-EMPTYVN"]["vn_change"], lt["K-AI-SAME"]["vn_change"], lt["K-PEND-DIFF"]["vn"]) == (True, False, "Chốt#&VN"),
      (sorted(lt), sm["n_latest"], sm["n_latest_hs"], sm["n_latest_st"], sm["rest_ai"], sm["rest_pending"], lt.get("K-BY-DATE")))

pg = quiet(cl.get, "/export/hs-fill").text


class Nest(HTMLParser):
    """그 id 들이 hfHistSum 안에 있는가(div 만 셈)."""
    def __init__(self):
        super().__init__()
        self.stack, self.seen = [], {}

    def handle_starttag(self, tag, attrs):
        i = dict(attrs).get("id")
        if i in ("hfLatest", "hfLatestOut", "hfHistView", "hfFindQ"):
            self.seen[i] = "hfHistSum" in self.stack
        if tag == "div":
            self.stack.append(i or "")

    def handle_endtag(self, tag):
        if tag == "div" and self.stack:
            self.stack.pop()


nest = Nest()
nest.feed(pg)
check("19 화면 — 📌 맞출 규격 5개(HS 바뀜 3 · 확정만 2) · 단추 data-n=5 · 사진 · 이력 없는 AI 2·확인 대기 1 고리 · 사진 있는 줄 수 · "
      "목록 번호 누름 · 확정 결과·찾기·품목 보기 칸은 hfHistSum 밖(새로 그려도 안 지워짐)",
      "📌 HS 사전을 최근 수출 값으로 맞출 규격 5개" in pg and "(HS 바뀜 3 · 같은 HS 로 확정만 2)" in pg and 'data-n="5"' in pg
      and "최근 수출 값으로 확정 (5개)" in pg and '<img class="hf-thumb" src="/export/hs-fill/hist-img/%s"' % IMGN in pg
      and 'href="/export/hs-dict?status=AI">AI 추정 2</a>' in pg and 'href="/export/hs-dict?status=PENDING">확인 대기 1</a>' in pg
      and "사진 있는 줄 <b>1</b>" in pg and 'data-no="KNK-260501-9" onclick="hfInvView(this)"' in pg
      and nest.seen == {"hfLatest": True, "hfLatestOut": False, "hfFindQ": False, "hfHistView": False},
      (nest.seen, [s for s in ("📌 HS 사전을 최근 수출 값으로 맞출 규격 5개", 'data-n="5"', "사진 있는 줄 <b>1</b>") if s not in pg]))


def snap():
    with D.db_session() as c:
        return [tuple(r) for r in c.execute("SELECT part_key, hs_code, status, vn_name, name_en, note, updated_by, updated_at "
                                            "FROM export_hs_dict ORDER BY part_key")]


S0 = snap()
r409 = quiet(cl.post, "/export/hs-fill/hist/apply-latest", json={"expect": 4})
as_user("M902")
r403 = quiet(cl.post, "/export/hs-fill/hist/apply-latest", json={"expect": 5}).status_code
as_user("M901")
check("20 화면이 본 개수와 다르면 409(그 사이 바뀜 · 「4 → 5개」) · 권한 없음 403 · 둘 다 사전은 그대로",
      r409.status_code == 409 and "4 → 5개" in r409.json().get("error", "") and r403 == 403 and snap() == S0,
      (r409.status_code, r409.json(), r403))
ra5 = quiet(cl.post, "/export/hs-fill/hist/apply-latest", json={"expect": 5})
ja = ra5.json()
d = {k: drow(k) for k in ("K-AI-SAME", "K-PEND-DIFF", "K-CONF-DIFF", "K-AI-EMPTYVN", "K-BY-DATE", "K-CONF-SAME", "K-CONF-DOTSAME",
                          "K-AI-NOHIST", "K-PEND-NOHIST", "K-HIST-EMPTYHS")}
check("21 확정 — 5개(HS 바뀜 3 · 확정만 2 · 품명 4) · 모두 확정 · HS = 최근 수출(날짜 순 22222222) · 누가(updated_by)",
      ra5.status_code == 200 and ja == {"ok": True, "applied": 5, "hs_changed": 3, "status_only": 2, "vn_changed": 4}
      and all(d[k]["status"] == "CONFIRMED" and d[k]["updated_by"] == U["M901"]["id"]
              for k in ("K-AI-SAME", "K-PEND-DIFF", "K-CONF-DIFF", "K-AI-EMPTYVN", "K-BY-DATE"))
      and [d[k]["hs_code"] for k in ("K-AI-SAME", "K-PEND-DIFF", "K-CONF-DIFF", "K-AI-EMPTYVN", "K-BY-DATE")]
      == ["85365099", "84799030", "85364990", "85365099", "22222222"],
      (ja, {k: (v["hs_code"], v["status"]) for k, v in d.items()}))
check("22 품명 — HS 가 바뀌면 그 HS 로 나간 가장 최근 품명(#& 앞은 영문명) · 사전 품명이 비었으면 채움 · HS 같고 품명 있으면 그대로",
      (d["K-PEND-DIFF"]["vn_name"], d["K-PEND-DIFF"]["name_en"], d["K-CONF-DIFF"]["vn_name"], d["K-AI-EMPTYVN"]["vn_name"],
       d["K-AI-SAME"]["vn_name"], d["K-BY-DATE"]["vn_name"]) == ("Chốt#&VN", "Chốt", "Rơ le điện", "Công tắc B", "Công tắc", "y2"),
      {k: (v["vn_name"], v["name_en"]) for k, v in d.items()})
check("23 사유 — 「지난 수출 최근 값으로 확정 날짜 이름: 전 HS(상태) → 새 HS (인보이스 날짜)」 · 전 사유가 있으면 뒤에 · 점 섞인 전 HS 는 숫자로",
      d["K-PEND-DIFF"]["note"].startswith("지난 수출 최근 값으로 확정 %s 가상M901: 39269099(확인 대기) → 84799030 (KNK-260301-9 2026-03-01)"
                                          " · 전 사유: 지난 수출 등록 · 이력끼리 HS 갈림" % today)
      and d["K-CONF-DIFF"]["note"] == "지난 수출 최근 값으로 확정 %s 가상M901: 85365099(확정) → 85364990 (KNK-260401-9 2026-04-01)" % today
      and d["K-AI-SAME"]["note"] == "지난 수출 최근 값으로 확정 %s 가상M901: 85365099(AI 추정) → 85365099 (KNK-260105-9 2026-01-05)" % today,
      (d["K-PEND-DIFF"]["note"], d["K-CONF-DIFF"]["note"], d["K-AI-SAME"]["note"]))
S1 = {r[0]: r for r in snap()}
S0d = {r[0]: r for r in S0}
check("24 대상이 아닌 줄은 그대로 — 확정·같은 HS · 점 섞인 같은 HS(9031.49.90 그대로) · 이력 없는 AI·확인 대기 · 빈 HS 이력",
      all(S1[k] == S0d[k] for k in ("K-CONF-SAME", "K-CONF-DOTSAME", "K-AI-NOHIST", "K-PEND-NOHIST", "K-HIST-EMPTYHS"))
      and S1["K-CONF-DOTSAME"][1] == "9031.49.90", {k: S1[k][1:3] for k in ("K-CONF-DOTSAME", "K-AI-NOHIST", "K-HIST-EMPTYHS")})
again = quiet(cl.post, "/export/hs-fill/hist/apply-latest", json={"expect": 0}).json()
again2 = quiet(cl.post, "/export/hs-fill/hist/apply-latest", json={}).json()
pg2 = quiet(cl.get, "/export/hs-fill").text
check("25 두 번째 — 0개(이미 다 맞음) · 화면은 「모두 최근 수출 값으로 확정돼 있습니다」 + 이력 없는 AI 2·확인 대기 1 · 📌 칸 없음",
      again == {"ok": True, "applied": 0, "hs_changed": 0, "status_only": 0, "vn_changed": 0} and again2.get("applied") == 0
      and "✅ HS 사전에서 지난 수출 이력이 있는 규격은 모두 최근 수출 값으로 확정돼 있습니다" in pg2
      and 'status=AI">AI 추정 2</a>' in pg2 and 'id="hfLatest"' not in pg2 and "📌 HS 사전을" not in pg2,
      (again, again2))
r2 = quiet(cl.post, "/export/hs-fill", files={"file": ("Invoice&packinglist_VN(260901).xlsx",
                                                     make_final("수책", [("K-PEND-DIFF", 1, 1.0, None, None)], "KNK-260901-1"))}).json()
check("26 다음 자동 채움부터 그 값 — 확인 대기였던 K-PEND-DIFF 가 자동(사전 확정) · 회신 필요 없음",
      (r2.get("summary") or {}).get("auto") == 1 and r2.get("no_reply") is True, (r2.get("summary"), r2.get("attention")))

print()
print("=" * 78)
print("§3 원문")
print("=" * 78)
tpl = open(os.path.join(ROOT, "app", "templates", "export_hs_fill.html"), encoding="utf-8").read()
js = tpl[tpl.index("<script>"):]
check("27 원문 — 사진(thumb · lazy 안 씀 — 숨은 칸 + lazy = 안 받음) · ① 단가·HS 표 사진 칸 · 노랑·주황 근거 사진 · 📚 결과·완료 띠 사진 수",
      "function thumb(" in js and 'loading="lazy"' not in tpl and "loading=\\'lazy\\'" not in tpl and ".loading" not in js
      and "thumb(p.img, true)" in js and "thumb(x.img, true)" in js
      and "thumb(a.hist_img, true)" in js and "j.imgs" in js and "tot.imgs" in js)
check("28 원문 — 최근 수출 값으로 확정(hfApplyLatest · 본 개수 expect 보냄 · 409 면 다시 그림) · 품목 보기(hfInvView) · 찾기(hfFind · 한글 조합 중 Enter 안 받음)",
      all(s in js for s in ("window.hfApplyLatest", "/export/hs-fill/hist/apply-latest", "expect: +d.n", "r.status === 409",
                            "window.hfInvView", "/export/hs-fill/hist/inv?no=", "window.hfFind", "/export/hs-fill/hist/find?q=",
                            "window.hfFindKey"))
      and "!event.isComposing" in tpl)

print()
print("결과: %d개 중 실패 %d" % (OK + FAIL, FAIL))
sys.exit(1 if FAIL else 0)
