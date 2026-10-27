# -*- coding: utf-8 -*-
"""z1191n 시험 — 인보이스 HS 자동 채움 다듬기 (대표 「진행해」 2026-10-27)
   대표 질문: 「hs코드 자동 모두 기입됐을때에는 어떻게 처리되지? 법인 회신이 필요없는 사항인데」
             「사전 확정을 하고 싶을때 어떻게해??」
  §1 모듈: 셈은 칸 색으로(품명만 노란 줄도 '법인 입력') · 노랑·주황이 없으면 no_reply · 고를 값(choices) · 근거 ·
           본사 값(overrides)으로 그 규격 칸만 덮기
  §2 경로(임시 DB · TestClient): ① 전부 자동 → 회신 필요 없음(글·기록·화면 단추) · 원본 보관 ·
           본사 확정 4갈래(적힌 값 · 사전 값 · 모델만 일치 · 확인 대기 값) → 사전·파일·숫자 · 먼저 받아 간 판도 ③ 거절 ·
           잘못된 요청 · 권한 · 원본 없는 기록 · 회신 짝 찾기(이름 같은 후보 여럿 → 정하지 않음 · 같은 번호 여러 줄 · 이미 반영)
  §3 원문: ① 여러 파일 차례로 · 표 새로 그리기 · 단추 이름 · 본사 확정 경로
"""
import contextlib
import io
import json
import os
import re
import sys
import tempfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
os.environ["KNK_SSO_SERVICE_KEY"] = "z1191n-test-only-not-a-real-key"   # 시험용 가짜 값

import openpyxl                                            # noqa: E402
from openpyxl.styles import Border, Side                   # noqa: E402
import app.database as D                                   # noqa: E402
from app import hs_fill as H                               # noqa: E402

OK = FAIL = 0
Y, O = "FFFFFF00", "FFFFC000"


def check(label, cond, got=None):
    global OK, FAIL
    if cond:
        OK += 1
        print("  [OK] " + label)
    else:
        FAIL += 1
        print("  [FAIL] " + label + ("   ← %r" % (got,) if got is not None else ""))


def make_form(kind, items, ino="KNK-261099-1"):
    """kind='수책'(HS=S·Tên hàng=T) / '소모품'(HS=R·Tên hàng=S). items=[(규격, HS, 품명, HS칸 수식?)] · ino=None 이면 번호 칸 비움"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INVOICE"
    ws["A1"] = "COMMERCIAL  INVOICE"
    if ino:
        ws["G4"] = ino
    if kind == "수책":
        heads = {"A": "12) Marks and Number of PKGS", "D": "13) Description of Goods", "F": " 14) Q'ty /Unit", "H": "15) Unit Price",
                 "I": "16) Amount", "K": "COMPANY", "L": "관리코드", "M": "MODEL", "N": "품명", "O": "Q'TY", "P": "ORIGIN",
                 "Q": "SUPPLIER", "R": "CATEGORY", "S": "HS CODE", "T": "Tên hàng", "U": "SHIPPING DATE"}
        chs, cvn = "S", "T"
    else:
        heads = {"A": "12) Marks and Number of PKGS", "D": "13) Description of Goods", "F": " 14) Q'ty /Unit", "H": "15) Unit Price",
                 "I": "16) Amount", "K": "COMPANY", "L": "IMAGE", "M": "Origin", "N": "Manufacturer", "O": "invoice\n%", "P": "C/O",
                 "Q": "SIZE", "R": "HS CODE", "S": "Tên hàng"}
        chs, cvn = "R", "S"
    for col, v in heads.items():
        ws[col + "31"] = v
    thin = Side(style="thin", color="FF000000")
    r = 34
    for desc, hs, vn, formula in items:
        ws["D%d" % r] = desc
        ws["F%d" % r] = 2
        ws["H%d" % r] = 3.5
        ws["I%d" % r] = "=F%d*H%d" % (r, r)
        for col in (chs, cvn):
            ws["%s%d" % (col, r)].border = Border(left=thin, right=thin, top=thin, bottom=thin)
        if formula:
            ws["%s%d" % (chs, r)] = '=IF(1=1,"","")'
        elif hs is not None:
            ws["%s%d" % (chs, r)] = hs
        if vn is not None:
            ws["%s%d" % (cvn, r)] = vn
        r += 1
    ws["A%d" % (r + 1)] = "TOTAL"
    ws["I%d" % (r + 1)] = "=SUM(I33:I%d)" % r
    ws.print_area = "A1:I%d" % (r + 5)
    pl = wb.create_sheet("PACKING LIST")
    pl["A1"] = "PACKING LIST"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), chs, cvn


def cell(data, ref):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    c = wb["INVOICE"][ref]
    v = c.value
    f = c.fill.fgColor.rgb if (c.fill is not None and c.fill.fill_type == "solid") else None
    wb.close()
    return v, f


DICT = [
    {"part_key": "DCP32DN15A", "model": "DCP32DN 15A", "hs_code": "85362011", "vn_name": "Cầu dao điện, mã hàng: DCP32DN 15A", "status": "CONFIRMED", "seen_count": 3, "last_seen": "2026-07-23"},
    {"part_key": "SP100F120H6.5TTBM", "model": "SP100F120H6.5TTBM", "hs_code": "85389019", "vn_name": "Chân pin tiếp xúc, mã hàng: SP100F120H6.5TTBM", "status": "CONFIRMED", "seen_count": 2, "last_seen": "2026-09-21"},
    {"part_key": "PINBLOCKSM-F971U-SUBTOP-R0.5C", "model": "PIN BLOCK SM-F971U-SUB TOP-R0.5C", "hs_code": "84669390", "vn_name": "x", "status": "CONFIRMED", "seen_count": 1, "last_seen": "2026-08-01"},
    {"part_key": "FLOATINGSM-F971U-SUBTOP-R0.5C", "model": "FLOATING SM-F971U-SUB TOP-R0.5C", "hs_code": "84798939", "vn_name": "y", "status": "CONFIRMED", "seen_count": 1, "last_seen": "2026-08-01"},
    {"part_key": "EN-AGS111107", "model": "EN-AGS111107", "hs_code": "39231090", "vn_name": "Hộp nhựa", "status": "PENDING", "seen_count": 2, "last_seen": "2026-07-23"},
    {"part_key": "XFHT4060", "model": "XFHT4060", "hs_code": "73181590", "vn_name": "", "status": "CONFIRMED", "seen_count": 1, "last_seen": "2026-06-01"},
    {"part_key": "G7SA-2A2B", "model": "G7SA-2A2B", "hs_code": "85364199", "vn_name": "Rơ le", "status": "CONFIRMED", "seen_count": 1, "last_seen": "2026-05-02"},
]
ITEMS = [  # 규격 · HS · 품명 · HS칸 수식?  (z1191l 시험과 같은 9줄)
    ("DCP32DN 15A", None, None, False),                      # 34 자동
    ("PARTS OF CONNECTORS SP100F120H6.5TTBM", None, None, False),   # 35 모델만 일치 → 확인(주황)
    ("SM-F971U-SUB TOP-R0.5C", None, None, False),           # 36 후보 HS 갈림 → 노랑
    ("EN-AGS111107", None, None, False),                     # 37 사전 확인 대기 → 노랑(고를 값 있음)
    ("NEW-PART-99999", None, None, False),                   # 38 새 품목 → 노랑
    ("G7SA-2A2B", "85364990", "Rơ le an toàn", False),       # 39 적힌 HS ↔ 사전 다름 → 주황
    ("XFHT4060", None, None, False),                         # 40 HS 자동 · 품명 노랑(사전에 품명 없음)
    ("DCP32DN 15A", "85362011", None, False),                # 41 품명만 자동
    ("DCP32DN 15A", None, None, True),                       # 42 HS 칸 수식
]

print("=" * 78)
print("§1 모듈 — 셈 · 고를 값 · 근거 · 본사 값")
print("=" * 78)
src8, _c1, _c2 = make_form("소모품", ITEMS[:8])
r8 = H.fill(src8, DICT)
cats = {a["row"]: H.category(a) for a in r8["acts"]}
check("1 셈은 칸 색으로: 자동 2 · 확인 2 · 법인 입력 4(40줄 = HS 자동 · 품명 노랑) · no_reply 아님",
      r8["counts"] == {"total": 8, "auto": 2, "check": 2, "need": 4, "keep": 0, "no_reply": False} and cats[40] == H.NEED,
      (r8["counts"], cats))
src_all, _a, _b = make_form("수책", [("DCP32DN 15A", None, None, False), ("G7SA-2A2B", None, None, False)])
src_keep, _a, _b = make_form("수책", [("DCP32DN 15A", "85362011", "Cầu dao", False)])
ra, rk = H.fill(src_all, DICT), H.fill(src_keep, DICT)
check("2 노랑·주황이 하나도 없으면 no_reply — 전부 자동(2) · 이미 다 적힘(1)",
      ra["counts"]["no_reply"] and ra["counts"]["auto"] == 2 and rk["counts"]["no_reply"] and rk["counts"]["keep"] == 1,
      (ra["counts"], rk["counts"]))
src9, chs9, cvn9 = make_form("수책", ITEMS)
r9 = H.fill(src9, DICT)
ch = {a["row"]: a["choices"] for a in r9["acts"]}
check("3 고를 값: 39 적힌·사전 둘 · 35 모델만 일치 값 · 37 확인 대기 값 · 36·38·42(후보 갈림·새 품목·수식)와 34·40·41 은 없음",
      ch[39] == [{"hs": "85364990", "kind": "file"}, {"hs": "85364199", "kind": "dict"}]
      and ch[35] == [{"hs": "85389019", "kind": "dict"}] and ch[37] == [{"hs": "39231090", "kind": "dict"}]
      and all(ch[r] == [] for r in (34, 36, 38, 40, 41, 42)), ch)
ac = {a["row"]: a for a in r9["acts"]}
check("4 근거: 35 = 다른 규격 SP100F120H6.5TTBM · 2번 · 마지막 2026-09-21 / 37 = 확인 대기 / 34 = 같은 규격이라 규격 표시 없음",
      ac[35]["dict_model"] == "SP100F120H6.5TTBM" and ac[35]["dict_seen"] == 2 and ac[35]["dict_last"] == "2026-09-21"
      and ac[37]["dict_status"] == "PENDING" and ac[34]["dict_model"] == "" and ac[34]["dict_seen"] == 3,
      {r: (ac[r]["dict_model"], ac[r]["dict_seen"], ac[r]["dict_last"], ac[r]["dict_status"]) for r in (34, 35, 37)})
check("4b 품명 근거: 39줄 적힌 품명 「Rơ le an toàn」 ↔ 사전 품명 「Rơ le」",
      ac[39]["vn_written"] == "Rơ le an toàn" and ac[39]["dict_vn"] == "Rơ le", (ac[39]["vn_written"], ac[39]["dict_vn"]))
ro = H.fill(src8, DICT, {"G7SA-2A2B": "85364199"})
a39 = next(a for a in ro["acts"] if a["row"] == 39)
check("5 본사 값(overrides): 적힌 85364990 → 85364199 로 쓰고 색 없음 · 품명(적힌 값)은 그대로 · 확인 2 → 1",
      a39["status"] == H.AUTO and a39["write_hs"] == "85364199" and a39["color_hs"] is None and a39["write_vn"] is None
      and cell(ro["out"], "R39") == ("85364199", None) and ro["counts"]["check"] == 1, (a39, cell(ro["out"], "R39"), ro["counts"]))
src_v, _a, _b = make_form("수책", [("G7SA-2A2B", "85364990", None, False)])
rv = H.fill(src_v, DICT, {"G7SA-2A2B": "85364199"})
check("6 본사 값 + 품명 빈 칸 → 사전 품명도 채움(사전 값과 같을 때)",
      cell(rv["out"], "S34") == ("85364199", None) and cell(rv["out"], "T34") == ("Rơ le", None) and rv["counts"]["no_reply"],
      (cell(rv["out"], "S34"), cell(rv["out"], "T34"), rv["counts"]))

print()
print("=" * 78)
print("§2 경로 — 임시 DB · TestClient")
print("=" * 78)
TMP = tempfile.mkdtemp(prefix="knk_z1191n_")
_REAL = D.DB_PATH
D.DB_PATH = os.path.join(TMP, "t.db")
assert D.DB_PATH != _REAL
with contextlib.redirect_stdout(io.StringIO()):
    D.init_db()
with D.db_session() as conn:
    have = {r[1] for r in conn.execute("PRAGMA table_info(users)")}
    for col, typ in [("employee_no", "TEXT"), ("entity", "TEXT"), ("name_vi", "TEXT"), ("dept_code", "TEXT"), ("rank", "TEXT"),
                     ("team_id", "INTEGER"), ("password_version", "INTEGER DEFAULT 1")]:
        if col not in have:
            conn.execute("ALTER TABLE users ADD COLUMN %s %s" % (col, typ))
    if "entity" not in {r[1] for r in conn.execute("PRAGMA table_info(teams)")}:
        conn.execute("ALTER TABLE teams ADD COLUMN entity TEXT")
    conn.execute("INSERT INTO teams (name, entity, code, display_order) VALUES ('가상기술영업팀','KOR','ZNHQ',1)")
    TM = conn.execute("SELECT MAX(id) FROM teams").fetchone()[0]
    U = {}
    for sb, role, fl in (("N901", "member", (1, 1, 0, 0)), ("N902", "member", (0, 0, 0, 0))):
        conn.execute("INSERT INTO users (name, role, employee_no, login_id, password, is_active, password_version, team_id, rank, "
                     "can_view_sales, can_use_sales, can_view_logistics, can_use_logistics) VALUES (?,?,?,?,'x',1,1,?,'프로',?,?,?,?)",
                     ("가상" + sb, role, sb, sb, TM) + fl)
        U[sb] = dict(conn.execute("SELECT * FROM users WHERE employee_no=?", (sb,)).fetchone())

import app.main as M                                        # noqa: E402
from fastapi.testclient import TestClient                  # noqa: E402

M._HS_FILL_DIR = os.path.join(TMP, "hs_fill")              # 시험 파일은 임시 폴더에
cl = TestClient(M.app, follow_redirects=False)
with D.db_session() as c:
    M._hs_dict_ensure(c)
    c.execute("DELETE FROM export_hs_dict")                 # 시드 대신 시험 사전만
    for e in DICT:
        c.execute("INSERT INTO export_hs_dict(part_key,model,hs_code,vn_name,status,seen_count,last_seen,source_file) VALUES(?,?,?,?,?,?,?,'시험')",
                  (e["part_key"], e["model"], e["hs_code"], e["vn_name"], e["status"], e["seen_count"], e["last_seen"]))


def as_user(sb):
    M.get_user = lambda request, _u=U[sb]: _u


def quiet(fn, *a, **k):
    with contextlib.redirect_stdout(io.StringIO()):
        return fn(*a, **k)


def log_row(lid):
    with D.db_session() as c:
        return dict(c.execute("SELECT * FROM export_hs_fill_log WHERE id=?", (lid,)).fetchone())


def dict_row(key):
    with D.db_session() as c:
        r = c.execute("SELECT * FROM export_hs_dict WHERE part_key=?", (key,)).fetchone()
        return dict(r) if r else None


def confirm(lid, sheet, row, hs):
    return quiet(cl.post, "/export/hs-fill/%d/confirm" % lid, json={"sheet": sheet, "row": row, "hs": hs})


as_user("N901")
fa, _a, _b = make_form("수책", [("DCP32DN 15A", None, None, False), ("G7SA-2A2B", None, None, False)], ino="KNK-261100-1")
ja = quiet(cl.post, "/export/hs-fill", files={"file": ("Invoice&packinglist_VN(261100).xlsx", fa)}).json()
check("7 ① 전부 자동 → no_reply · 전달용 글(「회신은 필요 없습니다」) · 볼 줄 없음",
      ja.get("ok") and ja["no_reply"] is True and ja["summary"]["need"] == 0 and ja["summary"]["check"] == 0 and ja["attention"] == []
      and ja["message"].startswith("[인보이스 전달] KNK-261100-1") and "회신은 필요 없습니다" in ja["message"], ja)
la = log_row(ja["log_id"])
check("8 기록: 원본도 보관(_in.xlsx · 파일 있음) · 법인 입력 0 · 확인 0",
      (la["orig_stored"] or "").endswith("_in.xlsx") and os.path.exists(os.path.join(M._HS_FILL_DIR, la["orig_stored"]))
      and la["n_need"] == 0 and la["n_check"] == 0, la)
pg = quiet(cl.get, "/export/hs-fill").text
check("9 화면 기록: 그 줄 「회신 필요 없음 · 받은 파일이 최종 파일」 · 단추 「최종 파일 받기」 · 회신 대기 표시 없음",
      '<span class="hf-done">회신 필요 없음</span><div class="hf-sub">받은 파일이 최종 파일</div>' in pg
      and "최종 파일 받기</a>" in pg and 'class="hf-wait"' not in pg)

f8, chs8, cvn8 = make_form("소모품", ITEMS[:8], ino="KNK-261100-2")
r1 = quiet(cl.post, "/export/hs-fill", files={"file": ("Invoice&packinglist_VN(261100-2).xlsx", f8)})
j1 = r1.json()
L2 = j1.get("log_id")
at = {a["row"]: a for a in j1.get("attention", [])}
check("10 ① 노랑·주황 있음: 자동 2·확인 2·법인 입력 4 · 글 「노란 칸 4줄」「주황 칸 2줄」 · 40줄(품명 노랑)도 볼 줄에 · 고를 값 35·37·39",
      j1.get("ok") and j1["summary"] == {"total": 8, "auto": 2, "check": 2, "need": 4, "keep": 0} and j1["no_reply"] is False
      and "노란 칸 4줄" in j1["message"] and "주황 칸 2줄" in j1["message"] and set(at) == {35, 36, 37, 38, 39, 40}
      and at[40]["st"] == "NEED" and "품명은 사전에 없음" in at[40]["why"]
      and {r for r, a in at.items() if a["choices"]} == {35, 37, 39}, (j1.get("summary"), sorted(at)))
check("10b 결과 JSON 에 품명 근거(39줄 적힌·사전)", at.get(39, {}).get("vn_written") == "Rơ le an toàn" and at.get(39, {}).get("dict_vn") == "Rơ le",
      at.get(39))
pg2 = quiet(cl.get, "/export/hs-fill").text
check("11 화면 기록: 회신 대기 1 · 단추 「보낸 파일 받기」", 'class="hf-wait">회신 대기 1' in pg2 and "보낸 파일 받기</a>" in pg2)
old_out = quiet(cl.get, j1["download_url"]).content

c1 = confirm(L2, "INVOICE", 39, "85364990")
k1 = c1.json()
g7 = dict_row("G7SA-2A2B")
check("12 본사 확정(적힌 값 85364990): 사전 85364199 → 85364990 확정 · note(누가·전 값·인보이스·줄) · 확인 2 → 1",
      c1.status_code == 200 and k1.get("ok") and g7["hs_code"] == "85364990" and g7["status"] == "CONFIRMED"
      and "본사 확정" in (g7["note"] or "") and "85364199 → 85364990" in (g7["note"] or "") and "KNK-261100-2 39줄" in (g7["note"] or "")
      and "가상N901" in (g7["note"] or "") and g7["updated_by"] == U["N901"]["id"]
      and k1["summary"]["check"] == 1 and k1["confirmed"] == {"row": 39, "desc": "G7SA-2A2B", "hs": "85364990", "prev": "85364199", "dict_changed": True},
      (c1.status_code, k1.get("summary"), k1.get("confirmed"), g7))
new_out = quiet(cl.get, j1["download_url"]).content
check("13 다시 만든 파일: 39줄 HS 85364990 · 색 없음 / 35줄은 아직 주황",
      cell(new_out, "R39") == ("85364990", None) and cell(new_out, "R35")[1] == O, (cell(new_out, "R39"), cell(new_out, "R35")))
e_old = (quiet(cl.post, "/export/hs-fill/reply", files=[("files", ("x.xlsx", old_out))]).json().get("results") or [{}])[0]
e_new = (quiet(cl.post, "/export/hs-fill/reply", files=[("files", ("y.xlsx", new_out))]).json().get("results") or [{}])[0]
check("14 ③: 확정 전에 받아 간 판·다시 만든 판 모두 「① 자동 채움 파일 그대로」로 거절(지문 기록)",
      "그대로입니다" in (e_old.get("error") or "") and "그대로입니다" in (e_new.get("error") or ""), (e_old, e_new))

k2 = confirm(L2, "INVOICE", 35, "85389019").json()
pc = dict_row("PARTSOFCONNECTORSSP100F120H6.5TTBM")
o2 = quiet(cl.get, j1["download_url"]).content
check("15 본사 확정(모델만 일치 85389019): 이 규격을 확정으로 새로 넣음(후보 품명) · 파일 35줄 HS·품명 색 없음 · 확인 0",
      k2.get("ok") and pc and pc["status"] == "CONFIRMED" and pc["hs_code"] == "85389019" and (pc["vn_name"] or "").startswith("Chân pin")
      and cell(o2, "R35") == ("85389019", None) and cell(o2, "S35")[1] is None and k2["summary"]["check"] == 0,
      (k2.get("summary"), pc, cell(o2, "R35"), cell(o2, "S35")))
k3 = confirm(L2, "INVOICE", 37, "39231090").json()
en = dict_row("EN-AGS111107")
o3 = quiet(cl.get, j1["download_url"]).content
check("16 본사 확정(확인 대기 값 39231090): 사전 확인 대기 → 확정 · 파일 37줄 HS·품명 채움 · 법인 입력 4 → 3",
      k3.get("ok") and en["status"] == "CONFIRMED" and en["hs_code"] == "39231090"
      and cell(o3, "R37") == ("39231090", None) and cell(o3, "S37") == ("Hộp nhựa", None) and k3["summary"]["need"] == 3,
      (k3.get("summary"), en, cell(o3, "R37"), cell(o3, "S37")))

fd, _a, _b = make_form("수책", [("DCP32DN 15A", "85369099", "Cầu dao", False), ("G7SA-2A2B", None, None, False)], ino="KNK-261100-3")
jd = quiet(cl.post, "/export/hs-fill", files={"file": ("Invoice&packinglist_VN(261100-3).xlsx", fd)}).json()
L3 = jd["log_id"]
kd = confirm(L3, "INVOICE", 34, "85362011").json()
dc = dict_row("DCP32DN15A")
od = quiet(cl.get, jd["download_url"]).content
l3 = log_row(L3)
pg3 = quiet(cl.get, "/export/hs-fill").text
check("17 본사 확정(사전 값 85362011 — 적힌 85369099 가 틀림): 사전 그대로 · 파일 칸을 사전 값으로 · 고른 값 기록 · 회신 필요 없음",
      kd.get("ok") and kd["confirmed"]["dict_changed"] is False and dc["hs_code"] == "85362011" and dc["status"] == "CONFIRMED"
      and cell(od, "S34") == ("85362011", None) and json.loads(l3["overrides"] or "{}") == {"DCP32DN15A": "85362011"}
      and kd["no_reply"] is True and l3["n_check"] == 0 and l3["n_need"] == 0
      and pg3.count('<span class="hf-done">회신 필요 없음</span><div class="hf-sub">받은 파일이 최종 파일</div>') == 2,
      (kd.get("summary"), kd.get("confirmed"), dc, cell(od, "S34"), l3.get("overrides")))

bad1 = confirm(L2, "INVOICE", 35, "85389019")                 # 이미 확정 → 고를 값 없음
bad2 = quiet(cl.post, "/export/hs-fill/%d/confirm" % L2, content=b"not json", headers={"Content-Type": "application/json"})
bad3 = quiet(cl.post, "/export/hs-fill/%d/confirm" % L2, json={"sheet": "INVOICE", "row": 0, "hs": "1234"})
bad4 = confirm(999999, "INVOICE", 34, "85362011")
with D.db_session() as c:
    c.execute("UPDATE export_hs_fill_log SET orig_stored=NULL WHERE id=?", (ja["log_id"],))
bad5 = confirm(ja["log_id"], "INVOICE", 34, "85362011")
as_user("N902")
bad6 = confirm(L2, "INVOICE", 36, "84669390")
as_user("N901")
check("18 막힘: 고를 수 없는 값 409 · JSON 아님 400 · HS·줄 틀림 400 · 없는 기록 404 · 원본 없는 기록 409(사전 화면 안내) · 권한 없음 403",
      (bad1.status_code, bad2.status_code, bad3.status_code, bad4.status_code, bad5.status_code, bad6.status_code) == (409, 400, 400, 404, 409, 403)
      and "HS 사전 화면" in bad5.json().get("error", ""),
      (bad1.status_code, bad2.status_code, bad3.status_code, bad4.status_code, bad5.status_code, bad6.status_code))

# 회신 짝 찾기 — 표준 양식을 이름 그대로 쓴 두 인보이스(번호 칸 비움 · z1191o 받는 이름 = 일반통관)
same = "Invoice&packinglist_VN(일반통관)_표준양식.xlsx"
for _ in range(2):
    fx, _a, _b = make_form("소모품", [("NEW-AAA-11111", None, None, False)], ino=None)
    quiet(cl.post, "/export/hs-fill", files={"file": (same, fx)})
with D.db_session() as c:
    two = [dict(r) for r in c.execute("SELECT id, reply_at FROM export_hs_fill_log WHERE file_name=? ORDER BY id", (same,))]
rx, _a, _b = make_form("소모품", [("NEW-AAA-11111", "84799030", "Linh kiện", False)], ino=None)
xa = (quiet(cl.post, "/export/hs-fill/reply", files=[("files", ("Invoice&packinglist_VN(일반통관)_표준양식 (1).xlsx", rx))]).json().get("results") or [{}])[0]
with D.db_session() as c:
    two2 = [r[0] for r in c.execute("SELECT reply_at FROM export_hs_fill_log WHERE file_name=? ORDER BY id", (same,))]
check("19 이름이 같은 후보 2개 → 어느 것인지 정하지 않음(ambiguous 2) · 두 기록 모두 회신 대기 그대로 · 사전 반영은 됨",
      len(two) == 2 and (xa.get("log") or {}).get("ambiguous") == 2 and two2 == [None, None] and xa.get("added") == 1
      and dict_row("NEW-AAA-11111")["hs_code"] == "84799030", (xa, two2))
for _ in range(2):
    fy, _a, _b = make_form("소모품", [("NEW-BBB-22222", None, None, False)], ino="KNK-261100-9")
    quiet(cl.post, "/export/hs-fill", files={"file": ("Invoice&packinglist_VN(261100-9).xlsx", fy)})
with D.db_session() as c:
    nine = [r[0] for r in c.execute("SELECT id FROM export_hs_fill_log WHERE invoice_no='KNK-261100-9' ORDER BY id")]
ry, _a, _b = make_form("소모품", [("NEW-BBB-22222", "84799030", "Linh kiện B", False)], ino="KNK-261100-9")
xb = (quiet(cl.post, "/export/hs-fill/reply", files=[("files", ("Invoice&packinglist_VN(261100-9)-re.xlsx", ry))]).json().get("results") or [{}])[0]
with D.db_session() as c:
    nine2 = [r[0] for r in c.execute("SELECT reply_at FROM export_hs_fill_log WHERE invoice_no='KNK-261100-9' ORDER BY id")]
check("20 같은 번호로 ①에 두 번 올린 인보이스 → 회신 한 번에 두 줄 모두 반영(n 2 · 새 줄 번호)",
      len(nine) == 2 and (xb.get("log") or {}).get("id") == nine[-1] and (xb.get("log") or {}).get("n") == 2 and all(nine2), (xb.get("log"), nine2))
xc = (quiet(cl.post, "/export/hs-fill/reply", files=[("files", ("Invoice&packinglist_VN(261100-9)-re.xlsx", ry))]).json().get("results") or [{}])[0]
check("21 이미 반영된 회신을 다시 올림 → again(그 기록 번호) · 사전 반영은 됨(같은 값)",
      (xc.get("log") or {}).get("again") == nine[-1] and not xc.get("error") and xc.get("updated") == 1, xc)

print()
print("=" * 78)
print("§3 원문")
print("=" * 78)
tpl = io.open(os.path.join(ROOT, "app", "templates", "export_hs_fill.html"), encoding="utf-8").read()
src = io.open(os.path.join(ROOT, "app", "main.py"), encoding="utf-8").read()
check("22 ① 파일 칸 여러 개(multiple) · 하나씩 차례로 도는 반복 · 파일마다 결과 카드",
      'id="hfFile" accept=".xlsx" multiple' in tpl and "for(var i = 0; i < n; i++)" in tpl and "slot.innerHTML = renderCard(j)" in tpl)
check("23 ①·③·본사 확정 뒤 기록 표 새로 그리기(hfLogCard) · 단추 이름 「보낸 파일 받기」/「최종 파일 받기」(「다시 받기」 없음)",
      'id="hfLogCard"' in tpl and tpl.count("hfReloadLog()") >= 3 and "보낸 파일 받기" in tpl and "최종 파일 받기" in tpl
      and ">다시 받기<" not in tpl)
check("23b 화면: 품명 근거 줄(vnDiff) · 사전이 이미 그 값이면 「파일 칸을 이 값으로 고침」 안내",
      "function vnDiff(" in tpl and "+ vnDiff(a) +" in tpl and "(사전은 이미 이 값) 파일 칸을 이 값으로 고침" in tpl)
check("24 본사 확정 경로 · 다시 만들기 잠금 · 결과 JSON 한 곳(①·확정 둘 다 _hs_fill_payload)",
      '@app.post("/export/hs-fill/{lid:int}/confirm")' in src and "_HS_FILL_REGEN_LOCK" in src
      and src.count("_hs_fill_payload(") == 3 and re.search(r"def _hs_fill_confirm_work\(", src) is not None)

print()
print("결과: %d개 중 실패 %d" % (OK + FAIL, FAIL))
sys.exit(1 if FAIL else 0)
