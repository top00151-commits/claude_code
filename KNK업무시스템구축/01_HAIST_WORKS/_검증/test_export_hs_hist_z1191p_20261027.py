# -*- coding: utf-8 -*-
"""z1191p 시험 — 지난 수출 이력(📚 등록) · ① 결과 비교(단가 ±20% · HS 다름) · HS 사전 새 규격 · 회신·최종 파일 이력 쌓기
   대표 2026-10-27 「지금 경로를 하나 줄꺼야... 이건 올해 수출한 내용을 수책통관, 일반통관 정리한 파일이야...
                    최종 인보이스 양식이 아니라 조금씩 달라... 예전 파일을 보고 HS코드 및 금액도 비교할수있게 미리 등록을 해 줬으면해...」
                  → 제안(비교 기준 ±20% · 사전에 없는 규격은 사전에 · ③ 회신도 이력에) → 「진행해」
  §1 모듈(app/hs_fill): 수량·단가·금액·날짜 읽기(예전 일반통관 · 예전 수책통관 · 최종 양식) · hist_lines · date_of · kind_of ·
     compare(단가 ±20% · HS 다름 · 같은 인보이스 빼기 · 근거 글) · model_conflict
  §2 경로(임시 DB · TestClient): 📚 등록 → 이력·원본↔저장 대조·사전(새 규격 확정 · 원래 있던 줄 그대로 · 다른 HS 는 셈만) ·
     같은 번호 다시 → 바꿔 넣음 · 이력이 늘면 등록이 넣은 사전 줄만 맞춤(이력끼리 갈림 → 확인 대기) · 같은 모델 충돌 → 확인 대기 ·
     사람이 고친 줄은 손대지 않음 · 대조가 다르면 되돌림(409) · 막힘(400·403) · 화면 📚 칸 ·
     ① 결과 비교 + 근거 + 이음 글에 단가 없음 · ① 회신 필요 없음 → 이력 · ③ 회신 → 이력 · HS 사전 올림 → 이력 · 본사 확정 → 회신 필요 없음 → 이력
  §3 원문: 화면 스크립트 이름
"""
import contextlib
import datetime
import io
import os
import re
import sys
import tempfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
os.environ["KNK_SSO_SERVICE_KEY"] = "z1191p-test-only-not-a-real-key"   # 시험용 가짜 값

import openpyxl                                            # noqa: E402
import app.database as D                                   # noqa: E402
from app import hs_fill as H                               # noqa: E402

OK = FAIL = 0


def check(label, cond, got=None):
    global OK, FAIL
    if cond:
        OK += 1
        print("  [OK] " + label)
    else:
        FAIL += 1
        print("  [FAIL] " + label + ("   ← %r" % (got,) if got is not None else ""))


def _save(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _top(ws, ino, date):
    ws["A1"] = "COMMERCIAL  INVOICE"
    ws["F2"] = "8) No. & Date of Invoice"
    if ino:
        ws["G4"] = ino
    if date is not None:
        ws["H4"] = "DATE:"
        ws["I4"] = date


def make_final(kind, items, ino, date=None):
    """최종(표준) 양식 — kind '수책'(HS=S·Tên hàng=T · CATEGORY·SUPPLIER) / '소모품'(HS=R·Tên hàng=S · IMAGE·C/O).
    items=[(규격, 수량, 단가, HS, 품명)] · 금액 칸은 수식(계산 값 없음 — openpyxl 로 만든 파일)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INVOICE"
    _top(ws, ino, date)
    if kind == "수책":
        heads = {"A": "12) Marks and Number of PKGS", "D": "13) Description of Goods", "F": " 14) Q'ty /Unit", "H": "15) Unit Price",
                 "I": "16) Amount", "K": "COMPANY", "P": "ORIGIN", "Q": "SUPPLIER", "R": "CATEGORY", "S": "HS CODE", "T": "Tên hàng"}
        chs, cvn = "S", "T"
    else:
        heads = {"A": "12) Marks and Number of PKGS", "D": "13) Description of Goods", "F": " 14) Q'ty /Unit", "H": "15) Unit Price",
                 "I": "16) Amount", "K": "COMPANY", "L": "IMAGE", "M": "Origin", "P": "C/O", "R": "HS CODE", "S": "Tên hàng"}
        chs, cvn = "R", "S"
    for col, v in heads.items():
        ws[col + "31"] = v
    r = 34
    for desc, q, p, hs, vn in items:
        ws["D%d" % r], ws["F%d" % r], ws["H%d" % r], ws["I%d" % r] = desc, q, p, "=F%d*H%d" % (r, r)
        if hs is not None:
            ws["%s%d" % (chs, r)] = hs
        if vn is not None:
            ws["%s%d" % (cvn, r)] = vn
        r += 1
    ws["A%d" % (r + 1)] = "TOTAL"
    wb.create_sheet("PACKING LIST")["A1"] = "PACKING LIST"
    return _save(wb)


def make_old(kind, items, ino, date=None):
    """예전 양식(수출자료 폴더 실물 모양) — 31줄 품명·수량·단가·금액 · HS 머리글은 32~33줄.
    '일반': Q32 'HS CODE CHECK' · Q33 'HS code' · R33 'Tên hàng' (260921 모양)
    '수책': N32 'HS CODE CHECK' · M33 'CATEGORY' · N33 'HS code' · O33 'Tên hàng' (260312 모양)
    금액 칸은 숫자(엑셀로 저장된 파일처럼 계산 값이 있음)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INVOICE"
    _top(ws, ino, date)
    for col, v in {"A": "12) Marks and Number of PKGS", "D": "13) Description of Goods", "F": "14) Q'ty /Unit",
                   "H": "15) Unit Price", "I": "16) Amount"}.items():
        ws[col + "31"] = v
    if kind == "일반":
        ws["L32"], ws["Q32"], ws["S32"], ws["Q33"], ws["R33"] = "Origin", "HS CODE CHECK", "Duty", "HS code", "Tên hàng"
        chs, cvn = "Q", "R"
    else:
        ws["K32"], ws["N32"], ws["S32"], ws["M33"], ws["N33"], ws["O33"] = "원산지", "HS CODE CHECK", "Duty AMOUNT", "CATEGORY", "HS code", "Tên hàng"
        chs, cvn = "N", "O"
    r = 34
    for desc, q, p, hs, vn in items:
        ws["D%d" % r], ws["F%d" % r], ws["H%d" % r], ws["I%d" % r] = desc, q, p, round(q * p, 4)
        ws["S%d" % r] = 9.99          # 오른쪽 'Duty' 칸 숫자 — 금액으로 읽으면 안 된다
        if hs is not None:
            ws["%s%d" % (chs, r)] = hs
        if vn is not None:
            ws["%s%d" % (cvn, r)] = vn
        r += 1
    ws["A%d" % (r + 1)] = "TOTAL"
    return _save(wb)


print("=" * 78)
print("§1 모듈")
print("=" * 78)
g = H.scan(make_old("일반", [("PARTS OF CONNECTORS SP100F", 400, 2, 85389019, "Chân pin"), ("XYZ-1", 3, 1.25, "8536201", None)],
                    "KNK-260921-1", datetime.datetime(2026, 9, 21)))
sg = g["sheets"][0]
check("1 예전 일반통관 — HS·품명 칸(Q·R) · 수량·단가·금액 칸 F·H·I(오른쪽 Duty 아님) · 날짜 I4 · 번호",
      (sg["col_hs"], sg["col_vn"], sg["col_qty"], sg["col_price"], sg["col_amount"]) == (17, 18, 6, 8, 9)
      and [(i["qty"], i["price"], i["amount"]) for i in sg["items"]] == [(400.0, 2.0, 800.0), (3.0, 1.25, 3.75)]
      and g["date"] == "2026-09-21" and g["invoice_no"] == "KNK-260921-1" and sg["form"] == "예전 양식",
      (sg["col_hs"], sg["col_vn"], sg["col_qty"], sg["col_price"], sg["col_amount"], g["date"]))
s2 = H.scan(make_old("수책", [("BS5-Y2M", 8, 3.2, "90314990", "Cảm biến")], "KNK-260312-1", "2026-03-12"))
ss = s2["sheets"][0]
check("2 예전 수책통관 — HS·품명 칸(N·O) · 단가 3.2 · 날짜 글자 '2026-03-12'",
      (ss["col_hs"], ss["col_vn"], ss["col_price"]) == (14, 15, 8) and ss["items"][0]["price"] == 3.2 and s2["date"] == "2026-03-12",
      (ss["col_hs"], ss["col_vn"], ss["col_price"], s2["date"]))
f3 = H.scan(make_final("수책", [("DCP32DN 15A", 11, 15.16, None, None)], "KNK-261015-1"))
hl = H.hist_lines(f3)
check("3 최종 수책 양식 — 단가 칸 H · 금액 수식에 계산 값 없음 → 이력 금액 = 수량×단가 · 날짜 칸 없음 → 번호에서",
      f3["sheets"][0]["col_price"] == 8 and f3["sheets"][0]["items"][0]["amount"] is None and abs(hl[0]["amount"] - 166.76) < 1e-6
      and H.date_of(f3) == "2026-10-15" and f3["date"] == "", (hl[0], H.date_of(f3)))
lines_g = H.hist_lines(g)
check("4 hist_lines — 7자리 HS 는 앞에 0(08536201) · 규격 키(공백 빼고 대문자) · 품명",
      lines_g[1]["hs"] == "08536201" and lines_g[0]["key"] == "PARTSOFCONNECTORSSP100F" and lines_g[0]["vn"] == "Chân pin"
      and lines_g[0]["hs"] == "85389019", lines_g)
check("5 date_of — 번호가 날짜가 아니면 '' · kind_of — 최종 수책 → 수책통관 · 최종 소모품 → 일반통관 · 예전 양식 → 모름",
      H.date_of({"invoice_no": "KNK-261399-1", "date": ""}) == "" and H.kind_of(f3) == "수책통관"
      and H.kind_of(H.scan(make_final("소모품", [("A-1", 1, 1, None, None)], "KNK-1"))) == "일반통관" and H.kind_of(g) == "",
      (H.date_of({"invoice_no": "KNK-261399-1"}), H.kind_of(f3), H.kind_of(g)))
acts = [{"sheet": "INVOICE", "row": 34, "key": "AAA-11111", "desc": "AAA-11111", "write_hs": "85362011", "hs_written": ""},
        {"sheet": "INVOICE", "row": 35, "key": "BBB-22222", "desc": "BBB-22222", "write_hs": None, "hs_written": "39269099"},
        {"sheet": "INVOICE", "row": 36, "key": "CCC-33333", "desc": "CCC-33333", "write_hs": None, "hs_written": ""}]
prices = {("INVOICE", 34): 14.0, ("INVOICE", 35): 5.0, ("INVOICE", 36): 1.0}
past = {"AAA-11111": [{"id": 1, "invoice_no": "KNK-260101-1", "inv_date": "2026-01-01", "hs": "85362011", "price": 10.0},
                      {"id": 5, "invoice_no": "KNK-260301-1", "inv_date": "2026-03-01", "hs": "85362012", "price": 11.0},
                      {"id": 9, "invoice_no": "KNK-269999-1", "inv_date": "2026-12-31", "hs": "11111111", "price": 99.0}],
        "BBB-22222": [{"id": 2, "invoice_no": "KNK-260101-1", "inv_date": "2026-01-01", "hs": "39269099", "price": 5.5}]}
cm = H.compare(acts, prices, past, "KNK-269999-1")
check("6 compare — 단가 14 ↔ 마지막 11 = +27.3%(±20% 넘음) · 5 ↔ 5.5 = -9.1%(안 넘음) · 같은 인보이스(99) 빼고 봄",
      cm["matched"] == 2 and cm["checked"] == 2 and len(cm["price"]) == 1 and cm["price"][0]["diff"] == 27.3
      and cm["price"][0]["last"] == 11.0 and cm["price"][0]["last_inv"] == "KNK-260301-1" and cm["price"][0]["n"] == 2, cm["price"])
check("7 compare — HS: 채울 값 85362011 ↔ 마지막 수출 85362012 → HS 다름(횟수 함께) · 적힌 값 = 마지막 → 안 셈 · 이력 없는 줄은 안 봄",
      len(cm["hs"]) == 1 and cm["hs"][0]["row"] == 34 and cm["hs"][0]["last"] == "85362012"
      and cm["hs"][0]["counts"] == {"85362011": 1, "85362012": 1} and ("INVOICE", 36) not in cm["ev"], cm["hs"])
check("8 compare — 근거 글(몇 번 · HS · 마지막 인보이스·날짜)",
      cm["ev"][("INVOICE", 34)] == "지난 수출 2번 · HS 85362011×1 · 85362012×1 · 마지막 KNK-260301-1(2026-03-01)", cm["ev"])
cm2 = H.compare([acts[0]], prices, {"AAA-11111": [
    {"id": 1, "invoice_no": "KNK-260101-1", "inv_date": "2026-01-01", "hs": "85362011", "price": 10.0},
    {"id": 2, "invoice_no": "KNK-260101-1", "inv_date": "2026-01-01", "hs": "85362011", "price": 10.5}]}, "")
check("8b 한 인보이스에 같은 규격이 여러 줄이어도 한 번으로 셈(지난 수출 1번 · HS ×1 — 단위를 인보이스 수로 맞춤)",
      cm2["ev"][("INVOICE", 34)] == "지난 수출 1번 · HS 85362011×1 · 마지막 KNK-260101-1(2026-01-01)", cm2["ev"])
check("9 대표 결정 값 — 단가 비교 기준 ±20% · 통관 구분 둘(수책통관·일반통관) · 양식 → 구분(수책 → 수책통관 · 소모품 → 일반통관)",
      H.HIST_PRICE_TOL == 20.0 and H.HIST_KINDS == ("수책통관", "일반통관")
      and H.KIND_BY_FORM == {"수책": "수책통관", "소모품": "일반통관"})

print()
print("=" * 78)
print("§2 경로")
print("=" * 78)
TMP = tempfile.mkdtemp(prefix="knk_z1191p_")
_REAL = D.DB_PATH
D.DB_PATH = os.path.join(TMP, "t.db")
assert D.DB_PATH != _REAL
with contextlib.redirect_stdout(io.StringIO()):
    D.init_db()
with D.db_session() as conn:
    have = {r[1] for r in conn.execute("PRAGMA table_info(users)")}
    for col, typ in [("employee_no", "TEXT"), ("entity", "TEXT"), ("name_vi", "TEXT"), ("rank", "TEXT"), ("team_id", "INTEGER"),
                     ("password_version", "INTEGER DEFAULT 1")]:
        if col not in have:
            conn.execute("ALTER TABLE users ADD COLUMN %s %s" % (col, typ))
    U = {}
    for sb, fl in (("M901", (1, 1, 0, 0)), ("M902", (0, 0, 0, 0))):
        conn.execute("INSERT INTO users (name, role, employee_no, login_id, password, is_active, password_version, "
                     "can_view_sales, can_use_sales, can_view_logistics, can_use_logistics) VALUES (?,?,?,?,'x',1,1,?,?,?,?)",
                     ("가상" + sb, "member", sb, sb) + fl)
        U[sb] = dict(conn.execute("SELECT * FROM users WHERE employee_no=?", (sb,)).fetchone())
import app.main as M                                        # noqa: E402
from fastapi.testclient import TestClient                  # noqa: E402

M._HS_FILL_DIR = os.path.join(TMP, "hs_fill")
M._HS_HIST_IMG_DIR = os.path.join(TMP, "hs_hist_img")       # z1191r: 회신·최종 파일의 품목 사진도 임시 폴더에(작업 폴더 data/ 에 남기지 않게)
cl = TestClient(M.app, follow_redirects=False)


def as_user(sb):
    M.get_user = lambda request, _u=U[sb]: _u


def quiet(fn, *a, **k):
    with contextlib.redirect_stdout(io.StringIO()):
        return fn(*a, **k)


def reg(kind, data, name):
    r = quiet(cl.post, "/export/hs-fill/hist", data={"kind": kind}, files={"file": (name, data)})
    return r.status_code, r.json()


def drow(k):
    with D.db_session() as c:
        r = c.execute("SELECT * FROM export_hs_dict WHERE part_key=?", (k,)).fetchone()
        return dict(r) if r else None


def q1(sql, *a):
    with D.db_session() as c:
        return c.execute(sql, a).fetchone()[0]


as_user("M901")
with D.db_session() as c:
    M._hs_dict_ensure(c)
    c.execute("DELETE FROM export_hs_dict")
    for k, m, hs, vn, st in (("DCP32DN15A", "DCP32DN 15A", "85362011", "Cầu dao điện", "CONFIRMED"),
                             ("FLOATINGPIN-77", "FLOATING PIN-77", "84798939", "Chốt", "CONFIRMED"),
                             ("WYLG1C15Z4", "WYLG1C15Z4", "85364140", "Rơ le", "CONFIRMED"),
                             ("SP100F120H6.5TTBM", "SP100F120H6.5TTBM", "85389019", "Chân pin", "CONFIRMED")):
        c.execute("INSERT INTO export_hs_dict(part_key,model,vn_name,hs_code,status,seen_count,last_seen,updated_at) "
                  "VALUES(?,?,?,?,?,3,'2026-07-23','2026-07-23 10:00:00')", (k, m, vn, hs, st))
A0, C0 = drow("DCP32DN15A"), drow("FLOATINGPIN-77")

F1 = make_final("수책", [("DCP32DN 15A", 2, 10.0, "85362011", "Cầu dao điện"), ("NEW-PART-B1", 1, 5.0, "39269099", "Hộp nhựa B"),
                         ("FLOATING PIN-77", 3, 1.0, "84799030", "Chốt nổi")], "KNK-260101-1", datetime.datetime(2026, 1, 1))
c1, j1 = reg("수책통관", F1, "Invoice&packinglist_VN(260101).xlsx")
B1 = drow("NEW-PART-B1")
check("10 📚 등록(수책통관) → 200 · 3줄 · 원본↔저장 3줄 같음 · 날짜 · 새로",
      c1 == 200 and j1.get("ok") and j1["lines"] == 3 and j1["check"] == {"ok": True, "n": 3} and j1["inv_date"] == "2026-01-01"
      and j1["replaced"] is None and j1["kind"] == "수책통관" and j1["kind_found"] == "수책통관", (c1, j1))
check("11 사전 — 없던 규격(B1)은 확정으로 들어감(횟수 1 · 처음·마지막 날 · 표시) · 원래 있던 같은 HS(A)·다른 HS(C)는 그대로",
      j1["dict"] == {"added": 1, "pending": 0, "synced": 0, "differ": 1, "kept": 1, "junk": 0}     # z1191v: junk = 함수 값 규격 수
      and B1 and B1["status"] == "CONFIRMED" and B1["hs_code"] == "39269099" and B1["seen_count"] == 1
      and B1["first_seen"] == B1["last_seen"] == "2026-01-01" and B1["hist_reg_at"] and B1["updated_at"] == B1["hist_reg_at"]
      and B1["vn_name"] == "Hộp nhựa B" and (B1["note"] or "").startswith("지난 수출 등록")
      and drow("DCP32DN15A") == A0 and drow("FLOATINGPIN-77") == C0, (j1["dict"], B1))
c1b, j1b = reg("수책통관", F1, "Invoice&packinglist_VN(260101) (2).xlsx")
check("12 같은 번호 다시 → 바꿔 넣음(전 출처 📚 등록) · 이력 줄 3 그대로(두 번 안 셈) · 사전 B1 횟수 1 그대로",
      c1b == 200 and j1b["replaced"] == {"source": "register", "file": "Invoice&packinglist_VN(260101).xlsx"}
      and q1("SELECT COUNT(*) FROM export_hist_line") == 3 and q1("SELECT COUNT(*) FROM export_hist_inv") == 1
      and drow("NEW-PART-B1")["seen_count"] == 1 and j1b["dict"]["synced"] == 1, (j1b, q1("SELECT COUNT(*) FROM export_hist_line")))
F2 = make_old("일반", [("NEW-PART-B1", 1, 5.5, "39269099", "Hộp nhựa B"), ("NEW-PART-D4", 4, 2.0, "85389019", "Chân pin D")],
              "KNK-260201-1", datetime.datetime(2026, 2, 1))
c2, j2 = reg("일반통관", F2, "Invoice&packinglist_VN(260201).xlsx")
B2, D2 = drow("NEW-PART-B1"), drow("NEW-PART-D4")
check("13 예전 일반통관 등록 → 등록이 넣은 B1 은 이력에 맞춤(횟수 2 · 마지막 2026-02-01) · D4 새로 확정",
      c2 == 200 and j2["dict"]["synced"] == 1 and j2["dict"]["added"] == 1 and B2["seen_count"] == 2
      and B2["last_seen"] == "2026-02-01" and B2["first_seen"] == "2026-01-01" and B2["status"] == "CONFIRMED"
      and D2["status"] == "CONFIRMED" and D2["hs_code"] == "85389019", (j2["dict"], B2["seen_count"], B2["last_seen"]))
F3 = make_old("수책", [("NEW-PART-B1", 1, 5.6, "39259000", "Hộp nhựa B2")], "KNK-260301-1", "2026-03-01")
c3, j3 = reg("수책통관", F3, "Invoice&packinglist_VN(260301).xlsx")
B3 = drow("NEW-PART-B1")
check("14 B1 이 다른 HS 로 또 나옴 → 이력끼리 HS 갈림 → '확인 대기'(사유에 횟수 · HS 는 마지막 값)",
      c3 == 200 and B3["status"] == "PENDING" and B3["hs_code"] == "39259000" and "이력끼리 HS 갈림" in (B3["note"] or "")
      and "39269099×2" in B3["note"] and B3["seen_count"] == 3, (B3["status"], B3["note"]))
F4 = make_final("소모품", [("RELAY(WYLG1C15Z4)", 1, 3.0, "85364199", "Rơ le")], "KNK-260310-1")
c4, j4 = reg("일반통관", F4, "Invoice&packinglist_VN(260310).xlsx")
R4 = drow("RELAY(WYLG1C15Z4)")
check("15 같은 모델(WYLG1C15Z4)이 사전에 확정으로 다른 HS → 새 규격은 '확인 대기'(충돌 사유) · 사전 줄은 그대로",
      c4 == 200 and j4["dict"]["pending"] == 1 and R4["status"] == "PENDING" and "같은 모델" in (R4["note"] or "")
      and drow("WYLG1C15Z4")["hs_code"] == "85364140", (j4["dict"], R4 and R4["note"]))
c4b, j4b = reg("일반통관", F4, "Invoice&packinglist_VN(260310) (2).xlsx")
R4b = drow("RELAY(WYLG1C15Z4)")
check("15b 같은 모델 충돌로 넣은 줄은 같은 값으로 다시 올라와도 '확인 대기' 그대로(사람이 고를 때까지)",
      c4b == 200 and j4b["dict"]["synced"] == 1 and R4b["status"] == "PENDING" and "같은 모델" in (R4b["note"] or ""),
      (j4b["dict"], R4b and R4b["status"]))
with D.db_session() as c:
    c.execute("UPDATE export_hs_dict SET updated_at='2099-01-01 00:00:00' WHERE part_key='NEW-PART-D4'")   # 사람이 사전 화면에서 고친 것처럼
D_before = drow("NEW-PART-D4")
F5 = make_old("일반", [("NEW-PART-D4", 1, 2.0, "85389099", "Chân pin D")], "KNK-260320-1", datetime.datetime(2026, 3, 20))
c5, j5 = reg("일반통관", F5, "Invoice&packinglist_VN(260320).xlsx")
check("16 등록이 넣은 줄이라도 그 뒤 사람이 고쳤으면 손대지 않음 — HS 가 달라도 '사전과 다름'으로 세기만",
      c5 == 200 and drow("NEW-PART-D4") == D_before and j5["dict"]["differ"] == 1 and j5["dict"]["synced"] == 0, (j5["dict"], drow("NEW-PART-D4")))
n_line0, n_inv0 = q1("SELECT COUNT(*) FROM export_hist_line"), q1("SELECT COUNT(*) FROM export_hist_inv")
_same = M._hs_hist_same
M._hs_hist_same = lambda a, b: False                       # 저장한 값이 원본과 다르다고 꾸밈
c6, j6 = reg("수책통관", F1, "Invoice&packinglist_VN(260101) (3).xlsx")
F7 = make_final("수책", [("NEW-PART-E7", 1, 1.0, "85365096", "x")], "KNK-260330-1")
c7, j7 = reg("수책통관", F7, "Invoice&packinglist_VN(260330).xlsx")
M._hs_hist_same = _same
check("17 원본↔저장 대조가 다르면 409 · 그 인보이스는 되돌림(있던 번호는 예전 줄 그대로 · 새 번호는 안 들어감 · 사전도 안 바뀜)",
      c6 == 409 and "대조" in j6.get("error", "") and c7 == 409 and q1("SELECT COUNT(*) FROM export_hist_line") == n_line0
      and q1("SELECT COUNT(*) FROM export_hist_inv") == n_inv0 and q1("SELECT COUNT(*) FROM export_hist_line WHERE invoice_no='KNK-260101-1'") == 3
      and drow("NEW-PART-E7") is None, (c6, j6, c7, j7))
bad = {
    "구분 없음": quiet(cl.post, "/export/hs-fill/hist", data={"kind": ""}, files={"file": ("a.xlsx", F1)}).status_code,
    "구분 틀림": quiet(cl.post, "/export/hs-fill/hist", data={"kind": "소모품"}, files={"file": ("a.xlsx", F1)}).status_code,
    "파일 없음": quiet(cl.post, "/export/hs-fill/hist", data={"kind": "수책통관"}).status_code,
    ".xls": quiet(cl.post, "/export/hs-fill/hist", data={"kind": "수책통관"}, files={"file": ("a.xls", F1)}).status_code,
    "엑셀 아님": reg("수책통관", b"not a zip", "a.xlsx")[0],
    "번호 없음": reg("수책통관", make_final("수책", [("Q-1", 1, 1.0, "85365096", "x")], None), "a.xlsx"),
    "HS 빈 줄 많음": reg("수책통관", make_final("수책", [("Q-1", 1, 1.0, None, None), ("Q-2", 1, 1.0, "85365096", None)], "KNK-260331-1"), "b.xlsx"),
}
as_user("M902")
bad["권한 없음"] = quiet(cl.post, "/export/hs-fill/hist", data={"kind": "수책통관"}, files={"file": ("a.xlsx", F1)}).status_code
as_user("M901")
check("18 막힘: 구분 없음·틀림 400 · 파일 없음 400 · .xls 400 · 엑셀 아님 400 · 번호 없음 400 · HS 빈 줄 절반 이상 400 · 권한 없음 403",
      (bad["구분 없음"], bad["구분 틀림"], bad["파일 없음"], bad[".xls"], bad["엑셀 아님"]) == (400, 400, 400, 400, 400)
      and bad["번호 없음"][0] == 400 and "번호" in bad["번호 없음"][1]["error"]
      and bad["HS 빈 줄 많음"][0] == 400 and "HS 가 빈 줄" in bad["HS 빈 줄 많음"][1]["error"] and bad["권한 없음"] == 403,
      {k: (v if isinstance(v, int) else v[0]) for k, v in bad.items()})
pg = quiet(cl.get, "/export/hs-fill").text
# z1191r: 「사전과 HS 다른 규격」 목록 → 「📌 HS 사전을 최근 수출 값으로 맞출 규격」(단추 하나) — 같은 FLOATING PIN-77 이 그 표에 있어야 한다
check("19 화면 📚 칸 — 인보이스 수·구분별·품목 줄·기간 · 목록에 번호 · '최근 수출 값으로 맞출 규격'에 FLOATING PIN-77(사전 84798939 → 수출 84799030)",
      'id="hfHistCard"' in pg and "등록된 인보이스 <b>5</b>개" in pg and "수책통관 2 · 일반통관 3" in pg and "2026-01-01 ~ 2026-03-20" in pg
      and "KNK-260310-1" in pg and re.search(r"FLOATING PIN-77</button></td>\s*<td class=\"mono\">84798939", pg) is not None
      and "84799030" in pg and "📌 HS 사전을 최근 수출 값으로 맞출 규격" in pg,
      re.findall(r"등록된 인보이스[^\n]{0,160}", pg)[:1])

U1 = make_final("수책", [("DCP32DN 15A", 2, 14.0, None, None), ("NEW-PART-B1", 1, 5.7, None, None), ("FLOATING PIN-77", 3, 1.0, None, None)],
                "KNK-260401-1")
r1 = quiet(cl.post, "/export/hs-fill", files={"file": ("Invoice&packinglist_VN(260401).xlsx", U1)})
p1 = r1.json()
hi = p1.get("hist") or {}
att_b = next((a for a in p1.get("attention", []) if a["desc"] == "NEW-PART-B1"), {})
check("20 ① 결과 비교 — 이력 있는 줄 3/3 · 단가 14 ↔ 마지막 10(260101) +40% · B1 5.7 ↔ 5.6 은 안 넘음 · 기준 20",
      r1.status_code == 200 and hi.get("n_inv") == 5 and hi.get("matched") == 3 and hi.get("tol") == 20.0
      and [(x["desc"], x["diff"], x["last"], x["last_inv"]) for x in hi.get("price", [])] == [("DCP32DN 15A", 40.0, 10.0, "KNK-260101-1")], hi)
check("21 ① HS 다름 — FLOATING PIN-77 사전 값 84798939 로 채움 ↔ 마지막 수출 84799030",
      [(x["desc"], x["now"], x["last"]) for x in hi.get("hs", [])] == [("FLOATING PIN-77", "84798939", "84799030")], hi.get("hs"))
check("22 ① 노랑 줄 근거에 지난 수출(3번 · HS 둘 · 마지막) · 이음 방 글에는 단가가 없음",
      att_b.get("hist_ev", "").startswith("지난 수출 3번 · HS 39269099×2 · 39259000×1 · 마지막 KNK-260301-1")
      and "단가" not in p1.get("message", "") and "14" not in p1.get("message", ""), (att_b.get("hist_ev"), p1.get("message")))
U2 = make_final("수책", [("DCP32DN 15A", 2, 10.0, None, None)], "KNK-260501-1")
r2 = quiet(cl.post, "/export/hs-fill", files={"file": ("Invoice&packinglist_VN(260501).xlsx", U2)}).json()
r2b = quiet(cl.post, "/export/hs-fill", files={"file": ("Invoice&packinglist_VN(260501).xlsx", U2)}).json()
with D.db_session() as c:
    iv = dict(c.execute("SELECT * FROM export_hist_inv WHERE invoice_no='KNK-260501-1'").fetchone() or {})
check("23 ① 회신 필요 없음 → 받은 파일(채운 결과)로 이력에(① 회신 필요 없음 · HS 채워짐) · 다시 올리면 바꿔 넣음 · 이력 있는 다른 인보이스와만 비교",
      r2.get("no_reply") and r2.get("hist_saved", {}).get("lines") == 1 and iv.get("source") == "fill" and iv.get("kind") == "수책통관"
      and q1("SELECT hs_code FROM export_hist_line WHERE invoice_no='KNK-260501-1'") == "85362011"
      and r2b.get("hist_saved", {}).get("replaced", {}).get("source") == "fill"
      and q1("SELECT COUNT(*) FROM export_hist_line WHERE invoice_no='KNK-260501-1'") == 1
      and r2b["hist"]["matched"] == 1 and r2b["hist"]["price"] == [], (r2.get("hist_saved"), iv, r2b.get("hist")))
R1 = make_final("수책", [("DCP32DN 15A", 2, 10.0, "85362011", "Cầu dao điện")], "KNK-260601-1")
rr = quiet(cl.post, "/export/hs-fill/reply", files=[("files", ("Invoice&packinglist_VN(260601)-re.xlsx", R1))]).json()
x = (rr.get("results") or [{}])[0]
c8, j8 = reg("수책통관", R1, "Invoice&packinglist_VN(260601).xlsx")
check("24 ③ 회신 → 이력(③ 회신) · 같은 번호를 📚 로 다시 올리면 바꿔 넣음(전 출처 ③ 회신)",
      x.get("hist", {}).get("lines") == 1 and c8 == 200 and j8["replaced"]["source"] == "reply", (x.get("hist"), j8.get("replaced")))
DU = make_final("소모품", [("DCP32DN 15A", 1, 9.9, "85362011", "Cầu dao điện")], "KNK-260701-1")
rd = quiet(cl.post, "/export/hs-dict/upload", files=[("files", ("Invoice&packinglist_VN(260701).xlsx", DU))]).json()
xd = (rd.get("results") or [{}])[0]
check("25 HS 사전 화면에 올린 최종 인보이스도 이력에(HS 사전 올림 · 일반통관)",
      xd.get("hist", {}).get("lines") == 1 and q1("SELECT source FROM export_hist_inv WHERE invoice_no='KNK-260701-1'") == "dict"
      and q1("SELECT kind FROM export_hist_inv WHERE invoice_no='KNK-260701-1'") == "일반통관", xd)
U3 = make_final("수책", [("PARTS OF CONNECTORS SP100F120H6.5TTBM", 2, 1.5, None, None)], "KNK-260801-1")
p3 = quiet(cl.post, "/export/hs-fill", files={"file": ("Invoice&packinglist_VN(260801).xlsx", U3)}).json()
a3 = (p3.get("attention") or [{}])[0]
cf = quiet(cl.post, "/export/hs-fill/%d/confirm" % p3["log_id"], json={"sheet": a3.get("sheet"), "row": a3.get("row"), "hs": "85389019"})
jc = cf.json()
check("26 주황 줄 본사 확정 → 회신 필요 없음이 되면 그 파일도 이력에(① 회신 필요 없음)",
      p3.get("no_reply") is False and "hist_saved" not in p3 and cf.status_code == 200 and jc.get("no_reply")
      and jc.get("hist_saved", {}).get("lines") == 1 and q1("SELECT source FROM export_hist_inv WHERE invoice_no='KNK-260801-1'") == "fill",
      (p3.get("no_reply"), cf.status_code, jc.get("hist_saved")))
pg2 = quiet(cl.get, "/export/hs-fill").text
check("27 화면 목록에 들어온 곳(📚 등록 · ③ 회신 → 📚 등록으로 바뀜 · HS 사전 올림 · ① 회신 필요 없음)",
      "HS 사전 올림" in pg2 and "① 회신 필요 없음" in pg2 and "📚 등록" in pg2 and "등록된 인보이스 <b>9</b>개" in pg2,
      re.findall(r"등록된 인보이스[^\n]{0,120}", pg2)[:1])

print()
print("=" * 78)
print("§3 원문")
print("=" * 78)
tpl = open(os.path.join(ROOT, "app", "templates", "export_hs_fill.html"), encoding="utf-8").read()
check("28 화면 스크립트 — 📚 등록(hfHist · 통관 구분 필수) · 비교 칸(histBlock) · 근거(hist_ev) · 최종 파일 이력(savedNote) · 요약 새로 그림",
      all(s in tpl for s in ("window.hfHist", "hfHistKind", "/export/hs-fill/hist", "function histBlock", "a.hist_ev",
                             "savedNote(j.hist_saved)", "'hfHistSum'")))

print()
print("결과: %d개 중 실패 %d" % (OK + FAIL, FAIL))
sys.exit(1 if FAIL else 0)
