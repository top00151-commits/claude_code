#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""BOM 부품 사진 — 최종안 확인용 산출물 만들기 (2026-09-10)

대표 확정: "사진은 아주 작게. 필요한 사용자가 **엑셀칸을 키워서** 확인할 수 있게."

두 벌을 만든다.
  ① 기본  — 평소 보이는 모습. 사진이 아주 작아 **표가 길어지지 않는다.**
  ② 키운것 — 사용자가 사진 열과 줄을 늘렸을 때. 사진이 **함께 커진다**는 증거.
     (②는 사람이 엑셀에서 끄는 것과 같은 일을 미리 해 둔 것 — 도구 기능이 아니다.)

⛔ 운영 자료 아님 — 설계팀 실물 인벤터 파일(BUDS 단차 검사기)로만 만든다.

실행:  python _검증/make_photo_samples.py [내보낼폴더]
"""
import os
import sys
import tempfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
os.chdir(ROOT)

from openpyxl import load_workbook                       # noqa: E402

import app.bom_tools.inventor_to_partlist as B           # noqa: E402
import app.bom_tools.merge_to_master as M                # noqa: E402
import app.bom_tools.photos as PH                        # noqa: E402

_TAIL = ("참고자료", "설계팀", "BOM 업무 자동화_2026.08.04",
         "1. BUDS 단차 검사기", "INVENTOR DOWN")
_CANDS = [os.path.join(ROOT, "..", *_TAIL),
          os.path.join(ROOT, "..", "..", "..", "KNK업무시스템구축", *_TAIL)]
INV = next((os.path.abspath(c) for c in _CANDS if os.path.isdir(c)), None)

OUT = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.join(HERE, "_시안_BOM사진")
os.makedirs(OUT, exist_ok=True)

if not INV:
    print("  [중단] 인벤터 실물 폴더를 찾지 못했습니다.")
    sys.exit(1)

print("=" * 70)
print("  BOM 부품 사진 - 최종안 (아주 작게 + 칸을 키우면 커짐)")
print("=" * 70)

T = tempfile.mkdtemp(prefix="knk_photosample_")
rows, _exc, _pf = B.read_inventor_files([INV])
print("  인벤터 구매품 %d줄 · 사진 %d장" % (len(rows), sum(1 for x in rows if x.get("사진"))))

draft = os.path.join(OUT, "간이판_사진.xlsx")
dres = B.write_partlist(rows, "005M2606", "BUDS 단차 검사기", "한재운", draft)
d_rows, _rep = M.read_draft_files([draft])

base = os.path.join(OUT, "통일판_사진_기본.xlsx")
mres = M.write_master(d_rows, "005M2606", "BUDS 단차 검사기", "한재운", base, sets=2)

dws = load_workbook(draft).active
mws = load_workbook(base).active
print("  간이판  사진 %d장 · Q열 너비 %.1f · 8행 높이 %.2fpt · %.0f KB"
      % (dres["사진"], dws.column_dimensions["Q"].width or 0,
         dws.row_dimensions[8].height or 0, os.path.getsize(draft) / 1024.0))
print("  통일판  사진 %d장 · AB열 너비 %.1f · 9행 높이 %.2fpt · %.0f KB"
      % (mres["사진"], mws.column_dimensions["AB"].width or 0,
         mws.row_dimensions[9].height or 0, os.path.getsize(base) / 1024.0))

# ② 사용자가 칸을 키웠을 때 — 사진이 함께 커지는지 눈으로 보이는 사본
big = os.path.join(OUT, "통일판_사진_칸을키웠을때.xlsx")
wb = load_workbook(base)
ws = wb.active
BIG_PT = 96.0
ws.column_dimensions["AB"].width = PH.col_width_for(BIG_PT)
for r in range(9, 9 + len(d_rows)):
    ws.row_dimensions[r].height = BIG_PT
wb.save(big)
print("  키운것  AB열 너비 %.1f · 줄 높이 %.0fpt (사람이 엑셀에서 끄는 것과 같은 조작)"
      % (PH.col_width_for(BIG_PT), BIG_PT))

print("-" * 70)
print("  폴더: " + OUT)
print("  두 파일을 나란히 열어 보시면 차이가 바로 보입니다.")
