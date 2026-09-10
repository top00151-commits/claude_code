#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""단가 시험용 **비식별 표본 생성기** (2026-09-10)

왜 만드나
  `test_bom_fill_prices.py` 의 「협력사 미정(뼈대 단계)」 시험이 **개인 임시 폴더**의
  특정 파일을 찾고 있었다. 그 폴더는 언제든 지워지고, 실제로 사라져서 시험이 실패했다.
  🔴 그 실패를 「단가 기능 이상 없음」의 근거로 쓸 수 없다 — **해당 사례 검증 미완료**다.

⛔ 실제 거래자료를 저장소에 넣지 않는다
  고객명·실단가·협력사명·도면·장비 사진이 **하나도 들어가지 않는다.**
  대신 **언제 돌려도 똑같은 표본**을 코드로 만든다(결정적 생성).

무엇을 지키나 (기존 시험 목적 그대로)
  협력사가 정해지지 않은 줄은 **단가를 자동으로 채우지 않는다.**
  과거 실적은 「참고」로만 알려 준다 — 협력사 확정은 구매팀 비교견적 몫이다(인터뷰 확정).
"""
import datetime
import os

from openpyxl import Workbook

# ── 가상 자료 (실제 회사 자료 아님) ─────────────────────────────────
VENDORS = ("가업체", "나업체", "다업체")

# 수불부에 실적이 있는 형번 14개 — 뼈대에서 「참고」로 뜨는 것
KNOWN = [("시험부품 %02d" % i, "TST-A-%03d" % i, 1000 + i * 100) for i in range(1, 15)]
# 수불부에 없는 형번 6개 — 참고가 뜨면 안 되는 것
UNKNOWN = [("시험부품 %02d" % i, "TST-B-%03d" % i, None) for i in range(15, 21)]

REF_COUNT = len(KNOWN)          # 기대되는 「과거 실적 참고」 건수
ROW_COUNT = len(KNOWN) + len(UNKNOWN)


def make_ledger(path):
    """수불부(발주관리대장) 표본. 머리글 이름은 실물 규칙 그대로 둔다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "발주관리대장"
    ws.cell(3, 2, "발주일자")
    ws.cell(3, 3, "규격 및 형번")
    ws.cell(3, 4, "협력사")
    ws.cell(3, 5, "단가")
    r = 4
    for i, (_nm, spec, price) in enumerate(KNOWN):
        ws.cell(r, 2, datetime.datetime(2026, 6, 1) + datetime.timedelta(days=i))
        ws.cell(r, 3, spec)
        ws.cell(r, 4, VENDORS[i % len(VENDORS)])
        ws.cell(r, 5, price)
        r += 1
    wb.save(path)
    return path


def make_skeleton(path):
    """통일판 **뼈대** 표본 — 협력사 칸이 비어 있고 기존단가도 비어 있다.

    실물 통일판과 같은 자리를 쓴다: E(5) 품명 · F(6) 형번 · H(8) 협력사 · P(16) 기존단가.
    """
    wb = Workbook()
    ws = wb.active
    ws.cell(7, 5, "PRODUCT NAME 제품명")
    ws.cell(7, 6, "PRODUCT CODE 코드명")
    ws.cell(7, 8, "VENDOR 외주사")
    ws.cell(7, 16, "UNIT PRICE 기존단가")
    r = 9
    for nm, spec, _p in KNOWN + UNKNOWN:
        ws.cell(r, 5, nm)
        ws.cell(r, 6, spec)
        # H(협력사) · P(기존단가) 는 **비워 둔다** = 뼈대 단계
        r += 1
    wb.save(path)
    return path


def make_all(out_dir):
    """(뼈대, 수불부) 두 파일을 만들어 경로를 돌려준다."""
    os.makedirs(out_dir, exist_ok=True)
    return (make_skeleton(os.path.join(out_dir, "표본_통일판뼈대_협력사미정.xlsx")),
            make_ledger(os.path.join(out_dir, "표본_수불부.xlsx")))


if __name__ == "__main__":
    import sys
    import tempfile
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
    d = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp(prefix="knk_pricefix_")
    a, b = make_all(d)
    print("뼈대  : %s (%d줄 · 협력사 미정)" % (a, ROW_COUNT))
    print("수불부: %s (실적 %d건)" % (b, REF_COUNT))
    print("기대  : 채움 0건 · 과거 실적 참고 %d건" % REF_COUNT)
